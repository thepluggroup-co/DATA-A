"""Classeurs Excel générés depuis le dictionnaire.

- `livre_de_codes` : onglet DICTIONNAIRE de V11 (VARIABLE, QUESTION / LIBELLÉ, FORMAT SPSS, MESURE, CODE, MODALITÉ).
- `parametrage` : onglets `03_QUESTIONS`, `04_MODALITES`, `05_VARIABLES_SPS` au format de
  `EDM_ELECTRIQUE_PV_CAMEROUN_PARAMETRAGE_V001.xlsx` (décision D-004), étendus aux colonnes
  nécessaires pour relire le dictionnaire sans perte (code, variable, langue anglaise, alias…).
  `imports/excel.py` relit ce format : l'aller-retour est testé.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs.masque import Colonne
from plateforme.generateurs.tri_a_plat import TableTriAPlat

ENTETE_FOND = PatternFill("solid", fgColor="1F4E78")
ENTETE_POLICE = Font(bold=True, color="FFFFFF")

COLONNES_QUESTIONS = ["Q_ID", "SECTION", "QUESTION_COURTE", "TYPE", "OBLIGATOIRE", "CONDITION", "VARIABLE_PRINCIPALE",
                      "NUMERO", "MESURE", "SECTION_LIBELLE", "QUESTION_EN", "AIDE", "AIDE_EN", "ROLE", "PREREMPLI",
                      "AUTRE_CODE", "AUTRE_LIBELLE", "AUTRE_LIBELLE_EN", "AUTRE_VARIABLE", "SECTION_LIBELLE_EN"]
COLONNES_MODALITES = ["VARIABLE", "CODE", "LIBELLE", "MESURE", "CONDITION", "LIBELLE_EN", "ALIAS", "VARIABLE_BINAIRE", "LIGNE_GRILLE"]
COLONNES_VARIABLES = ["NOM_VARIABLE", "LIBELLE_SPSS", "TYPE", "VALEURS", "MANQUANTS", "MESURE", "ROLE"]
COLONNES_REGLES = ["CODE", "EXPRESSION", "DESCRIPTION"]


def _entete(ws: Worksheet, titres: list[str]) -> None:
    ws.append(titres)
    for cellule in ws[1]:
        cellule.fill, cellule.font = ENTETE_FOND, ENTETE_POLICE
    ws.freeze_panes = "A2"


def livre_de_codes(colonnes: list[Colonne], chemin: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "DICTIONNAIRE"
    _entete(ws, ["VARIABLE", "QUESTION / LIBELLÉ", "FORMAT SPSS", "MESURE", "CODE", "MODALITÉ"])
    for c in colonnes:
        if c.valeurs:
            for i, (code, lib) in enumerate(c.valeurs):
                ws.append([c.code, c.label, c.spss, c.mesure or ""] if i == 0 else ["", "", "", ""])
                ws.cell(ws.max_row, 5, code)
                ws.cell(ws.max_row, 6, lib)
        else:
            ws.append([c.code, c.label, c.spss, c.mesure or "", "", "(texte / valeur libre)"])
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["F"].width = 40
    wb.save(chemin)


def parametrage(dico: Dictionnaire, colonnes: list[Colonne], chemin: str | Path) -> None:
    wb = Workbook()
    q_ws = wb.active
    assert q_ws is not None
    q_ws.title = "03_QUESTIONS"
    _entete(q_ws, COLONNES_QUESTIONS)
    for q in dico.questions:
        s = dico.section(q.section)
        q_ws.append([
            f"Q{q.numero}", s.id, q.libelles.get("fr", ""), q.type.value, "Oui" if q.obligatoire else "Non", q.condition or "",
            q.code, q.numero, q.mesure.value if q.mesure else "", s.libelles.get("fr", ""), q.libelles.get("en", ""),
            q.aide.get("fr", ""), q.aide.get("en", ""), q.role or "", q.prerempli or "",
            q.autre.code if q.autre else "", q.autre.libelles.get("fr", "") if q.autre else "",
            q.autre.libelles.get("en", "") if q.autre else "", q.autre.variable if q.autre else "", s.libelles.get("en", ""),
        ])
    m_ws = wb.create_sheet("04_MODALITES")
    _entete(m_ws, COLONNES_MODALITES)
    for q in dico.questions:
        for m in q.modalites:
            m_ws.append([q.code, m.code, m.libelles.get("fr", ""), q.mesure.value if q.mesure else "NOMINAL", "",
                         m.libelles.get("en", ""), " | ".join(m.alias), m.variable or "", ""])
        for li in q.lignes:
            m_ws.append([q.code, "", li.libelles.get("fr", ""), "", "", li.libelles.get("en", ""), "", "", li.variable])
    v_ws = wb.create_sheet("05_VARIABLES_SPS")
    _entete(v_ws, COLONNES_VARIABLES)
    speciaux = "/".join(str(c) for c in sorted(dico.codes_speciaux.tous()))
    for c in colonnes:
        v_ws.append([c.code, c.label, "NUMERIC" if c.est_numerique() else "STRING", "Voir 04_MODALITES" if c.valeurs else "",
                     speciaux if c.valeurs else "", c.mesure or "", c.section])
    r_ws = wb.create_sheet("07_CONTROLES")
    _entete(r_ws, COLONNES_REGLES)
    for r in dico.regles_coherence:
        r_ws.append([r.code, r.expression, r.description.get("fr", "")])
    p_ws = wb.create_sheet("00_PARAMETRES")
    _entete(p_ws, ["PARAMETRE", "VALEUR"])
    for k, v in [("ETUDE", dico.etude), ("VERSION", dico.version), ("LANGUES", ",".join(dico.langues)),
                 ("LANGUE_REFERENCE", dico.langue_reference), ("TITRE", dico.titre.get("fr", "")),
                 ("TITRE_EN", dico.titre.get("en", ""))]:
        p_ws.append([k, v])
    wb.save(chemin)


def tri_a_plat(tables: list[TableTriAPlat], villes: list[str], chemin: str | Path) -> None:
    """Gabarit vide (effectifs à remplir en phase 3)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "TRI_A_PLAT"
    ws.append(["TRI À PLAT — fiches REJETE exclues. Ligne TOTAL = base ; ligne 99 = question laissée vide."])
    ws.append([])
    for t in tables:
        ws.append([t.titre])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append(["Code", "Modalité", *villes, "Total", "% base"])
        for code, lib in t.lignes:
            ws.append([code, lib])
        if t.ligne_total:
            ws.append(["", "TOTAL (= base)"])
        ws.append([])
    wb.save(chemin)
