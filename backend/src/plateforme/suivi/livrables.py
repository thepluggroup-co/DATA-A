"""Livrables du suivi de terrain : tableau de bord, rapport journalier, état des enquêteurs, effet enquêteur, rappels.

Fichiers produits (dossier de sortie) :
- TABLEAU_DE_BORD.xlsx : enquêteurs × jours, villes, quotas cible × ville, cibles par jour, récapitulatif par jour ;
- TRI_A_PLAT.xlsx : chaque variable par ville et par jour, ligne TOTAL = base ;
- RAPPORT_<jour>.xlsx : rapport journalier (provisoire si demandé avant la fin de la journée) ;
- ETAT_ENQUETEURS_<jour>.docx : état des lieux des postes (sans noms : ils ne sont pas dans le dépôt) ;
- EFFET_ENQUETEUR.xlsx : tous les tests, écarts signalés en tête ;
- RAPPELS_<jour>.xlsx : rappels de contrôle tirés au sort — CONTIENT DES TÉLÉPHONES (données personnelles).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from plateforme.codification.fiche import FicheCodee
from plateforme.codification.qualite import ParametresBase, base_analyse, quasi_doublons
from plateforme.codification.rejeu import HORS_COMPTAGE, DecisionQC
from plateforme.controle.effet_enqueteur import detecter
from plateforme.controle.rappels import tirer_rappels
from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs.masque import colonnes_masque
from plateforme.generateurs.tri_a_plat import gabarit
from plateforme.suivi.base import Enqueteur, Ville, base_suivi
from plateforme.suivi.rapport import ParametresRapport, Rapport, calculer_rapport
from plateforme.suivi.statut import exporter_word, rapport_statut
from plateforme.suivi.tableaux import ParametresSuivi, recap_jours, suivi_cibles, tableau_de_bord, tri_a_plat
from plateforme.terrain.plan import PlanTerrain, Role

TETE = Font(bold=True, color="FFFFFF")
FOND = PatternFill("solid", fgColor="1F4E78")


def charger_quotas(chemin: str | Path, plan: PlanTerrain) -> dict[int, tuple[str, list[int]]]:
    """Quotas du plan d'échantillonnage (format provisoire, modèle complet en phase 6 / M2)."""
    brut = yaml.safe_load(Path(chemin).read_text(encoding="utf-8"))
    return {int(code): (str(q["cible"]), [int(q[v.id]) for v in plan.villes]) for code, q in brut["quotas"].items()}


def _feuille(wb: Workbook, titre: str, premiere: bool = False) -> Worksheet:
    ws = wb.active if premiere else wb.create_sheet(titre)
    assert ws is not None
    ws.title = titre
    return ws


def _entete(ws: Worksheet, valeurs: list[Any]) -> None:
    ws.append(valeurs)
    for c in ws[ws.max_row]:
        c.font, c.fill = TETE, FOND


def ecrire_tableau_de_bord(base: list[dict[str, Any]], p: ParametresSuivi, chemin: Path) -> None:
    tb, sc, rj = tableau_de_bord(base, p), suivi_cibles(base, p), recap_jours(base, p)
    jours = [f"J{i + 1} — {j[8:10]}/{j[5:7]}" for i, j in enumerate(p.jours)]
    wb = Workbook()
    ws = _feuille(wb, "TABLEAU_DE_BORD", premiere=True)
    ws.append(["SUIVI TERRAIN — fiches REJETE exclues des comptes, fiches ANNULEE absentes"])
    _entete(ws, ["ENQ_ID", "ZONE", *jours, "AUJOURD'HUI", "CUMUL", "OBJECTIF", "% RÉALISÉ", "DÉLAIS COURTS", "À VÉRIFIER", "REJETÉES",
                 "STATUT", "RANG"])
    for li in tb.lignes:
        ws.append([li["id"], li["zone"], *li["jours"], li["aujourdhui"], li["cumul"], li["objectif"], li["realise"], li["delais_courts"],
                   li["a_verifier"], li["rejetees"], li["statut"], li["rang"]])
    ws.append(["Hors liste", "", *[""] * len(jours), "", tb.hors_liste])
    t = tb.total
    ws.append(["TOTAL", "", *t["jours"], t["aujourdhui"], t["cumul"], t["objectif"], t["realise"], t["delais_courts"], t["a_verifier"],
               t["rejetees"]])
    ws.append([])
    _entete(ws, ["VILLE", "ENQUÊTEURS", *jours, "CUMUL", "OBJECTIF (capacité)", "QUOTA", "% DU QUOTA"])
    for z in [*tb.zones, {**tb.total_zones, "id": "TOTAL"}]:
        ws.append([z["id"], z["enqueteurs"], *z["jours"], z["cumul"], z["capacite"], z["quota"], z["part_quota"]])
    ws.append([f"CONTRÔLE — fiches comptées sans réponse à la question de cible : {tb.sans_cible}"])
    ws = _feuille(wb, "QUOTAS_CIBLES")
    _entete(ws, ["CODE", "CIBLE", *[f"{v.id} {x}" for v in p.villes for x in ("quota", "réalisé", "%")], "TOTAL quota", "réalisé", "%"])
    for c in sc.quotas:
        ws.append([c["code"], c["cible"], *[x for v in c["villes"] for x in (v["quota"], v["realise"], v["part"])],
                   c["total"]["quota"], c["total"]["realise"], c["total"]["part"]])
    ws.append(["", "TOTAL DES CIBLES", *[x for v in sc.total["villes"] for x in (v["quota"], v["realise"], v["part"])],
               sc.total["total"]["quota"], sc.total["total"]["realise"], sc.total["total"]["part"]])
    ws.append(["", "TOTAL GÉNÉRAL (= tableau de bord)", *[x for v in sc.total_general[:-1] for x in ("", v, "")], "", sc.total_general[-1], ""])
    ws = _feuille(wb, "RECAP_JOURS")
    _entete(ws, ["INDICATEUR", *jours, "CUMUL"])
    for nom, valeurs in rj.items():
        ws.append([nom, *valeurs])
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)


def ecrire_rapport(r: Rapport, chemin: Path) -> None:
    wb = Workbook()
    ws = _feuille(wb, "RAPPORT", premiere=True)
    ws.append([f"RAPPORT JOURNALIER — {r.jour}" + (" — PROVISOIRE (journée en cours)" if r.partiel else "")])
    for titre, entete, lignes in (
        ("1. SYNTHÈSE DU JOUR", ["Indicateur", "Valeur", "Lecture"], r.synthese),
        ("2. RECOMMANDATIONS POUR LE JOUR SUIVANT", ["Thème", "Action recommandée"], r.recos),
        ("3. PRODUCTION PAR VILLE", ["Ville", "Fiches du jour", "Cumul", "Quota", "Reste", "Enquêteurs", "Requis / enquêteur / jour", "Situation"],
         [[v["z"].libelle, v["nJ"], v["cumul"], v["quota"], v["reste"], v["nEnq"], "—" if v["requis"] is None else round(v["requis"], 1),
           v["etat"]] for v in r.villes]),
        ("4. CLASSEMENT DES ENQUÊTEURS", ["Rang", "Enquêteur", "Fiches du jour", "Zone", "% objectif", "Cumul", "À vérifier", "Rejetées",
                                         "Délais courts", "Tél. en double", "% sans réponse", "Score /100", "Appréciation"],
         [[c["rang"], c["id"], c["n"], c["zone"], c["pObj"], c["cumul"], c["pb"], c["rej"], c["dc"], c["tel"], c["nr"], c["score"], c["avis"]]
          for c in r.classement]),
        ("5. CIBLES INTERROGÉES", ["Cible", "Fiches du jour", "% du jour", "Cumul", "Quota", "% du quota", "Reste"],
         [[c["lib"], c["nJ"], c["pJ"], c["cumul"], c["quota"], c["pQ"], c["reste"]] for c in r.cibles]),
        ("7. QUESTIONS LAISSÉES SANS RÉPONSE", ["Question", "Fiches sans réponse", "% des fiches du jour"],
         [[x["q"], x["n"], x["tx"]] for x in r.nonRep]),
        ("8. RYTHME DE COLLECTE PAR HEURE", ["Heure", "Fiches"], r.heures),
    ):
        ws.append([])
        ws.append([titre])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        _entete(ws, entete)
        for li in lignes:
            ws.append(li)
    ws.append([])
    ws.append(["Score = 50 points de production (fiches ÷ objectif, plafonné) + 30 points de fiabilité (part des fiches sans anomalie "
               "imputable ni rejet) + 20 points de complétude (0 point à partir de 20 % de questions sans réponse)."])
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)


def produire(dico: Dictionnaire, fiches: list[FicheCodee], plan: PlanTerrain, quotas: dict[int, tuple[str, list[int]]],
             decisions: dict[str, DecisionQC], jour: str, maintenant: datetime, sortie: Path, graine_rappels: int = 0) -> dict[str, Path]:
    """Produit tous les livrables du suivi pour le jour demandé ; renvoie les chemins écrits.

    Seules les fiches antérieures à `maintenant` sont prises en compte : un rapport rejoué après coup voit
    exactement ce qui existait à l'heure demandée.
    """
    tz = maintenant.tzinfo
    fiches = [f for f in fiches
              if datetime.fromisoformat(f"{f.valeurs['DATEENQU']}T{f.valeurs['HEUREENQ']}:00").replace(tzinfo=tz) <= maintenant]
    jours = [j.isoformat() for j in plan.jours]
    titulaires = [e for e in plan.enqueteurs if e.role is Role.TITULAIRE]
    villes = [Ville(v.code, v.id, v.libelle) for v in plan.villes]
    base = base_suivi(fiches)
    heure = maintenant.hour + maintenant.minute / 60
    ps = ParametresSuivi([Enqueteur(e.id, "", e.ville or "") for e in plan.enqueteurs], villes, jours, quotas, plan.objectif_jour,
                         plan.nb_jours, plan.seuil_volume_jour, maintenant.strftime("%Y-%m-%d"))
    pr = ParametresRapport(jour, jours, plan.nb_jours, [Enqueteur(e.id, "", e.ville or "") for e in plan.enqueteurs], villes, quotas,
                           plan.objectif_jour, plan.seuil_delai_min, plan.seuil_volume_jour, (maintenant.strftime("%Y-%m-%d"), heure))
    ecrits: dict[str, Path] = {}
    ecrits["tableau"] = sortie / "TABLEAU_DE_BORD.xlsx"
    ecrire_tableau_de_bord(base, ps, ecrits["tableau"])
    n, par_jour, tables = tri_a_plat(base, gabarit(dico, colonnes_masque(dico, plan.villes_masque())), ps)
    wb = Workbook()
    ws = _feuille(wb, "TRI_A_PLAT", premiere=True)
    ws.append([f"TRI À PLAT — base : {n} fiches valides ; base du jour : {par_jour}. Ligne TOTAL = base ; ligne 99 = question laissée vide."])
    for tt in tables:
        ws.append([])
        ws.append([tt.titre])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        _entete(ws, ["Code", "Modalité", *[v.libelle for v in villes], "Total", "% base", *[f"J{i + 1}" for i in range(len(jours))],
                     *[f"% J{i + 1}" for i in range(len(jours))]])
        for li in [*tt.lignes, *([tt.total] if tt.total else [])]:
            ws.append(li)
    ecrits["tri_a_plat"] = sortie / "TRI_A_PLAT.xlsx"
    wb.save(ecrits["tri_a_plat"])
    ecrits["rapport"] = sortie / f"RAPPORT_{jour}.xlsx"
    ecrire_rapport(calculer_rapport(dico, colonnes_masque(dico, plan.villes_masque()), base, pr), ecrits["rapport"])
    if plan.objectif_total and titulaires:
        ecrits["statut"] = exporter_word(rapport_statut(fiches, plan, jours, maintenant), sortie / f"ETAT_ENQUETEURS_{jour}.docx", dico.etude)
    retenues, _ = base_analyse(dico, fiches, ParametresBase(), decisions)
    wb = Workbook()
    ws = _feuille(wb, "EFFET_ENQUETEUR", premiere=True)
    ws.append(["Écarts à profil égal (khi-deux, Bonferroni, barrière de Tukey). Un signalement déclenche une vérification, "
               "ce n'est pas une preuve. Enquêteur et quartier peuvent se confondre."])
    _entete(ws, ["Signalé", "Groupe", "Enquêteur", "Variable", "Question", "n", "khi²", "ddl", "p corrigé", "w", "Modalité sur-représentée",
                 "Part chez l'enquêteur", "Part attendue à profil égal"])
    for x in sorted(detecter(dico, retenues), key=lambda x: (not x.signale, x.p)):
        ws.append(["OUI" if x.signale else "", str(x.groupe), x.enqueteur, x.variable, x.libelle, x.n, round(x.khi2, 2), x.ddl, x.p_corrige,
                   round(x.w, 3), x.modalite_surrepresentee, x.part_enqueteur, x.part_attendue])
    ecrits["effet"] = sortie / "EFFET_ENQUETEUR.xlsx"
    wb.save(ecrits["effet"])
    valides = [f for f in fiches if f.statut not in HORS_COMPTAGE]
    du_jour = [f for f in valides if f.valeurs.get("DATEENQU") == jour]
    wb = Workbook()
    ws = _feuille(wb, "RAPPELS", premiere=True)
    ws.append([f"Rappels de contrôle du {jour} — graine du tirage : {graine_rappels} (tirage reproductible). DONNÉES PERSONNELLES."])
    _entete(ws, ["N° fiche", "Enquêteur", "Priorité", "Motif", "Téléphone", "Résultat du rappel"])
    for r in tirer_rappels(du_jour, 3, graine_rappels, quasi_doublons(dico, valides, ParametresBase())):
        ws.append([r.numfiche, r.enqueteur, r.priorite, r.motif, r.telephone, ""])
    for c in ws["E"]:
        c.number_format = "@"
    ecrits["rappels"] = sortie / f"RAPPELS_{jour}.xlsx"
    wb.save(ecrits["rappels"])
    return ecrits
