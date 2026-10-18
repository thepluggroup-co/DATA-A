"""Rapport de statut des enquêteurs (« état des lieux ») à une heure donnée, exportable en Word.

Reprend l'état des lieux produit pendant l'étude EDM (jour 3, 07/10/2026 à 12 h 26) :
- base : objectif de l'équipe réparti sur les POSTES (titulaires ; les réserves ne comptent pas) — 500 / 8 = 62,5
  fiches par poste, affiché arrondi (63) ; pourcentages calculés sur la valeur exacte ; « reste minimum »
  arrondi à la fiche supérieure ;
- 1. titulaires présents ayant atteint leur objectif ; 2. présents ne l'ayant pas atteint (reste à remplir) ;
  3. titulaires sortis (fiches produites, moment de sortie) ; 4. remplaçants (démarrage, fiches par jour,
  cumul du poste) ; situation par ville ; vue d'ensemble.
Les fiches d'un poste sont attribuées à la personne qui l'occupait à l'heure de la fiche (affectations datées
du plan de terrain). Les noms et motifs, données personnelles, ne sont fournis qu'au moment de l'export.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from plateforme.codification.fiche import FicheCodee
from plateforme.codification.rejeu import HORS_COMPTAGE
from plateforme.suivi.base import arrondi
from plateforme.terrain.plan import Affectation, EnqueteurPlan, PlanTerrain, Role


@dataclass
class LigneStatut:
    poste: str
    personne: str
    ville: str
    jours: list[int | None]          # None : la personne n'occupait pas encore le poste ce jour-là
    fiches: int
    cumul_poste: int
    part: float                      # cumul du poste / objectif exact du poste
    reste: int                       # arrondi à la fiche supérieure
    sortie: str = ""
    demarrage: str = ""
    note: str = ""


@dataclass
class RapportStatut:
    coupure: datetime
    jours: list[str]
    postes: int
    objectif_total: int
    objectif_poste: float
    vue: dict[str, Any] = field(default_factory=dict)
    atteints: list[LigneStatut] = field(default_factory=list)
    non_atteints: list[LigneStatut] = field(default_factory=list)
    sortis: list[LigneStatut] = field(default_factory=list)
    remplacants: list[LigneStatut] = field(default_factory=list)
    villes: list[dict[str, Any]] = field(default_factory=list)


def _moment(f: FicheCodee, tz: ZoneInfo) -> datetime:
    return datetime.fromisoformat(f"{f.valeurs['DATEENQU']}T{f.valeurs['HEUREENQ']}:00").replace(tzinfo=tz)


def _dans(a: Affectation, t: datetime) -> bool:
    return (a.debut is None or t >= a.debut) and (a.fin is None or t < a.fin)


def rapport_statut(fiches: list[FicheCodee], plan: PlanTerrain, jours: list[str], coupure: datetime) -> RapportStatut:
    if plan.objectif_total is None:
        raise ValueError("objectif_total absent du plan de terrain")
    tz = ZoneInfo(plan.fuseau)
    postes = [e for e in plan.enqueteurs if e.role is Role.TITULAIRE]
    obj = plan.objectif_total / len(postes)
    jour_coupure = coupure.astimezone(tz).strftime("%Y-%m-%d")
    jours = [j for j in jours if j <= jour_coupure]
    libelles = {v.id: v.libelle for v in plan.villes}
    valides = [f for f in fiches if f.statut not in HORS_COMPTAGE and _moment(f, tz) <= coupure]
    r = RapportStatut(coupure, jours, len(postes), plan.objectif_total, obj)

    def ligne(e: EnqueteurPlan, a: Affectation | None, cumul_poste: int, remplacant: bool = False) -> LigneStatut:
        siennes = [f for f in valides if f.valeurs.get("CODEENQU") == e.id and (a is None or _dans(a, _moment(f, tz)))]
        par_jour: list[int | None] = []
        for j in jours:
            debut_jour = datetime.fromisoformat(f"{j}T23:59:59").replace(tzinfo=tz)
            commence = not remplacant or a is None or a.debut is None or a.debut <= debut_jour   # « – » : remplaçant pas encore arrivé
            par_jour.append(sum(1 for f in siennes if f.valeurs["DATEENQU"] == j) if commence else None)
        return LigneStatut(e.id, a.personne if a else "titulaire", libelles.get(e.ville or "", ""), par_jour, len(siennes), cumul_poste,
                           cumul_poste / obj, max(0, math.ceil(obj - cumul_poste)),
                           sortie=a.sortie if a else "", demarrage=a.demarrage if a else "")

    for e in postes:
        cumul = sum(1 for f in valides if f.valeurs.get("CODEENQU") == e.id)
        affectations = e.affectations or [Affectation(personne="titulaire")]
        titulaire = affectations[0]
        lt = ligne(e, titulaire, cumul)
        if titulaire.fin is not None and titulaire.fin <= coupure:
            r.sortis.append(lt)
        else:
            if titulaire.debut is not None:
                lt.note = f"{e.id} n'a démarré qu'au {titulaire.demarrage or titulaire.debut.date().isoformat()}."
            (r.atteints if cumul >= obj else r.non_atteints).append(lt)
        for a in affectations[1:]:
            if a.debut is not None and a.debut <= coupure:
                r.remplacants.append(ligne(e, a, cumul, remplacant=True))
    r.non_atteints.sort(key=lambda x: -x.cumul_poste)
    r.atteints.sort(key=lambda x: -x.cumul_poste)

    de_poste = [f for f in valides if f.valeurs.get("CODEENQU") in {e.id for e in postes}]
    cumul_total = len(de_poste)
    r.vue = {"jours": [sum(1 for f in de_poste if f.valeurs["DATEENQU"] == j) for j in jours], "cumul": cumul_total,
             "objectif": plan.objectif_total, "part": cumul_total / plan.objectif_total, "reste": max(0, math.ceil(plan.objectif_total - cumul_total))}
    for v in plan.villes:
        ids = {e.id for e in postes if e.ville == v.id}
        cumul = sum(1 for f in de_poste if f.valeurs.get("CODEENQU") in ids)
        objectif_ville = plan.objectif_total * len(ids) / len(postes)
        r.villes.append({"ville": v.libelle, "postes": len(ids), "cumul": cumul, "objectif": objectif_ville,
                         "part": cumul / objectif_ville if objectif_ville else 0, "reste": max(0, math.ceil(objectif_ville - cumul))})
    return r


def pourcent(x: float) -> str:
    return f"{arrondi(100 * x)} %"


def exporter_word(r: RapportStatut, chemin: str | Path, titre_etude: str, noms: dict[tuple[str, str], str] | None = None) -> Path:
    """Export Word. `noms` (poste, personne) -> nom : fourni au moment de l'export, jamais stocké dans le dépôt."""
    from docx import Document

    noms = noms or {}
    tz = r.coupure.tzinfo
    h = r.coupure.astimezone(tz)
    rang = len(r.jours)
    doc = Document()
    doc.add_heading(f"ÉTAT DES LIEUX DES ENQUÊTEURS – JOUR {rang:02d}", 0)
    doc.add_paragraph(titre_etude)
    doc.add_paragraph(f"Situation au {h.strftime('%d/%m/%Y')} à {h.hour} h {h.minute:02d}.")
    doc.add_paragraph(f"Base de calcul. Objectif de {r.objectif_total} fiches pour {r.postes} postes, soit {math.ceil(r.objectif_poste)} fiches "
                      "par poste. Le reste à remplir est arrondi à la fiche supérieure.")

    def tableau(entete: list[str], lignes: list[list[Any]]) -> None:
        t = doc.add_table(rows=1, cols=len(entete))
        t.style = "Table Grid"
        for c, texte in zip(t.rows[0].cells, entete, strict=True):
            c.text = texte
        for li in lignes:
            cellules = t.add_row().cells
            for c, v in zip(cellules, li, strict=True):
                c.text = "–" if v is None else str(v)

    def jours_tete(debut: int = 0) -> list[str]:
        return [f"J{i + 1}" for i in range(debut, len(r.jours))]

    doc.add_heading("Vue d'ensemble", 1)
    tableau(["Périmètre", *jours_tete(), "Cumul", "Objectif", "% réalisé", "Reste minimum de fiches"],
            [[f"{r.postes} postes", *r.vue["jours"], r.vue["cumul"], r.vue["objectif"], pourcent(r.vue["part"]), r.vue["reste"]]])
    for titre, lignes in (("1. Enquêteurs présents depuis le début et ayant atteint leur objectif", r.atteints),
                          ("2. Enquêteurs présents depuis le début et n'ayant pas atteint leur objectif", r.non_atteints)):
        doc.add_heading(titre, 1)
        if not lignes:
            doc.add_paragraph("Aucun enquêteur dans cette situation.")
            continue
        tableau(["ENQ_ID", "Nom", "Zone", *jours_tete(), "Cumul", "% réalisé", "Reste minimum de fiches"],
                [[x.poste, noms.get((x.poste, x.personne), ""), x.ville, *x.jours, x.cumul_poste, pourcent(x.part), x.reste] for x in lignes])
        for x in lignes:
            if x.note:
                doc.add_paragraph(x.note)
    doc.add_heading("3. Enquêteurs présents au début et sortis du dispositif", 1)
    tableau(["ENQ_ID", "Nom", "Zone", "Fiches produites", "Sortie"],
            [[x.poste, noms.get((x.poste, x.personne), ""), x.ville, x.fiches, x.sortie] for x in r.sortis])
    doc.add_heading("4. Enquêteurs remplaçants", 1)
    tableau(["ENQ_ID", "Nom", "Zone", "Démarrage", *jours_tete(1), "Fiches du remplaçant", "Cumul du poste", "% réalisé", "Reste minimum de fiches"],
            [[x.poste, noms.get((x.poste, x.personne), ""), x.ville, x.demarrage, *x.jours[1:], x.fiches, x.cumul_poste, pourcent(x.part), x.reste]
             for x in r.remplacants])
    doc.add_heading("Situation par ville", 1)
    tableau(["Zone", "Postes", "Cumul", "Objectif", "% réalisé", "Reste minimum de fiches"],
            [[v["ville"], v["postes"], v["cumul"], math.ceil(v["objectif"]), pourcent(v["part"]), v["reste"]] for v in r.villes]
            + [["TOTAL", r.postes, r.vue["cumul"], r.objectif_total, pourcent(r.vue["part"]), r.vue["reste"]]])
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(chemin))
    return chemin
