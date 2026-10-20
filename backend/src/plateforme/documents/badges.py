"""Badges des enquêteurs et superviseurs (M7) : format carte CR80 (85,6 × 54 mm), 8 par page A4, Word puis PDF (D-049).

Contenu (gabarit du badge EDM, CLAUDE.md §3.1) : logo, intitulé de l'étude, photo (ou cadre « photo »), nom, qualité
(« Agent enquêteur » / « Agent enquêtrice » / « Superviseur(e) de ville »), téléphone, ville, identifiant (ENQ3), validité
(dates de terrain), téléphone et e-mail du bureau, QR code vers l'adresse de vérification paramétrée (D-049) si elle existe.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from plateforme.documents.commun import BANDEAU_FICTIF, date_fr, enregistrer
from plateforme.projet.projet import Projet

LARGEUR_CM, HAUTEUR_CM = 8.56, 5.4
PAR_PAGE = 8


@dataclass(frozen=True)
class Badge:
    identifiant: str
    nom: str
    qualite: str
    telephone: str
    ville: str
    photo: Path | None
    qr: str | None


def url_badge(base: str, etude: str, identifiant: str) -> str:
    return f"{base.rstrip('/')}/{etude}/{identifiant}"


def badges(p: Projet) -> list[Badge]:
    """Un badge par poste pourvu (titulaires, réserves) et par superviseur de ville, dans l'ordre du plan de terrain."""
    t = p.exiger_terrain()
    out = []
    base = t.url_verification or p.cabinet.url_verification
    for m in t.equipe:
        if m.fonction.value != "superviseur" or not m.personne:
            continue
        pers = p.vivier.personne(m.personne)
        out.append(Badge(m.id, pers.nom_complet, "Superviseure de ville" if pers.genre == "F" else "Superviseur de ville",
                         pers.telephone, t.ville(m.ville).libelle if m.ville else "", _photo(p, pers.photo),
                         url_badge(base, t.etude, m.id) if base else None))
    for e in t.enqueteurs:
        pid = e.personne or next((a.personne_id for a in reversed(e.affectations) if a.personne_id), None)
        if not pid:
            continue
        pers = p.vivier.personne(pid)
        out.append(Badge(e.id, pers.nom_complet, "Agent enquêtrice" if pers.genre == "F" else "Agent enquêteur", pers.telephone,
                         t.ville(e.ville).libelle if e.ville else "", _photo(p, pers.photo),
                         url_badge(base, t.etude, e.id) if base else None))
    return out


def _photo(p: Projet, chemin: str | None) -> Path | None:
    if not chemin:
        return None
    f = p.chemin(chemin)
    return f if f.exists() else None


def _qr_png(texte: str) -> io.BytesIO:
    import qrcode

    tampon = io.BytesIO()
    qrcode.make(texte, border=1).save(tampon, format="PNG")
    tampon.seek(0)
    return tampon


def _cellule_badge(cellule: Any, b: Badge, p: Projet, validite: tuple[date, date] | None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    from plateforme.analyse.rapport.rendus import charte_par_defaut

    c = charte_par_defaut()
    bleu = RGBColor.from_string(c.couleur_titres)
    cellule.width = Cm(LARGEUR_CM)
    interne = cellule.add_table(rows=1, cols=2)
    gauche, droite = interne.rows[0].cells
    gauche.width, droite.width = Cm(2.4), Cm(5.9)
    logo = p.chemin(p.cabinet.logo) if p.cabinet.logo else c.logo
    if logo and Path(logo).exists():
        gauche.paragraphs[0].add_run().add_picture(str(logo), width=Cm(2.0))
    if b.photo:
        gauche.add_paragraph().add_run().add_picture(str(b.photo), width=Cm(1.8))
    else:
        r = gauche.add_paragraph().add_run("PHOTO\n3 × 4")
        r.font.size = Pt(6)
    if b.qr:
        gauche.add_paragraph().add_run().add_picture(_qr_png(b.qr), width=Cm(1.4))

    def ligne(texte: str, taille: float = 7, gras: bool = False, couleur: RGBColor | None = None) -> None:
        par = droite.add_paragraph() if droite.paragraphs[0].text or len(droite.paragraphs) > 1 else droite.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        r = par.add_run(texte)
        r.font.size, r.bold = Pt(taille), gras
        if couleur is not None:
            r.font.color.rgb = couleur

    ligne("BADGE OFFICIEL D'ENQUÊTE", 7.5, True, bleu)
    ligne(p.brief.titre, 6.5)
    ligne(b.nom, 9, True)
    ligne(b.qualite, 7.5, True, bleu)
    ligne(f"ID : {b.identifiant}    Ville : {b.ville}", 7)
    if b.telephone:
        ligne(f"Tél. : {b.telephone}", 7)
    if validite:
        ligne(f"Valable du {date_fr(validite[0])} au {date_fr(validite[1])}", 6.5)
    bureau = " · ".join(x for x in (p.cabinet.nom, p.cabinet.telephone, p.cabinet.email) if x)
    ligne(bureau, 6)
    if p.brief.fictif:
        ligne("FICTIF — démonstration", 6, True, RGBColor(0xC0, 0x00, 0x00))
    for par in droite.paragraphs:
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT


def produire(p: Projet, dossier: Path, en_pdf: bool = False) -> list[Path]:
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Cm, Pt

    t = p.exiger_terrain()
    jours = t.jours_ouvres()
    validite = (min(jours), max(jours)) if jours else None
    liste = badges(p)
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.4)
    s.top_margin = s.bottom_margin = Cm(1.2)
    doc.styles["Normal"].font.size = Pt(7)
    if p.brief.fictif:
        doc.add_paragraph(BANDEAU_FICTIF).runs[0].font.size = Pt(8)
    for debut in range(0, len(liste), PAR_PAGE):
        if debut:
            doc.add_page_break()  # type: ignore[no-untyped-call]
        page = liste[debut:debut + PAR_PAGE]
        table = doc.add_table(rows=(len(page) + 1) // 2, cols=2)
        table.style = "Table Grid"
        for ligne_t in table.rows:
            ligne_t.height, ligne_t.height_rule = Cm(HAUTEUR_CM), WD_ROW_HEIGHT_RULE.EXACTLY
        for k, b in enumerate(page):
            _cellule_badge(table.rows[k // 2].cells[k % 2], b, p, validite)
    return enregistrer(doc, dossier / "BADGES.docx", en_pdf)
