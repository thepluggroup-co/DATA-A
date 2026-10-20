"""Outils communs aux documents Word : en-tête (logo, titre), bandeau « projet fictif », tableaux à la charte, montants,
dates à la française, espaces insécables, conversion PDF (même chaîne que le rapport, phase 5)."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from plateforme.analyse.rapport.rendus import Charte, _fond, _insecables_docx, charte_par_defaut, pdf

MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
BANDEAU_FICTIF = "PROJET FICTIF DE DÉMONSTRATION — personnes, montants et client fictifs"
BANDEAU_JURIDIQUE = ("MODÈLE À FAIRE VALIDER PAR UN AVOCAT CAMEROUNAIS AVANT TOUT USAGE — document généré, sans valeur "
                     "juridique en l'état")


def date_fr(j: date, jour: bool = False) -> str:
    texte = f"{j.day} {MOIS[j.month - 1]} {j.year}"
    return f"{JOURS[j.weekday()]} {texte}" if jour else texte


def montant(x: float, devise: str = "FCFA") -> str:
    """12 500 FCFA : entier, espace fine insécable entre les milliers."""
    return f"{round(x):,}".replace(",", " ") + (f" {devise}" if devise else "")


def nouveau(titre: str, charte: Charte | None = None, logo: Path | None = None, fictif: bool = False,
            sous_titre: str = "", marges_cm: float = 2.0) -> Any:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    c = charte or charte_par_defaut()
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(marges_cm)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = c.police, Pt(10.5)
    style.font.color.rgb = RGBColor.from_string(c.couleur_texte)
    if fictif:
        p = doc.add_paragraph()
        r = p.add_run(BANDEAU_FICTIF)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xC0, 0x00, 0x00)
    image = logo or c.logo
    if image and Path(image).exists():
        doc.add_picture(str(image), width=Cm(3.5))
    if titre:
        h = doc.add_heading(titre, level=0)
        for run in h.runs:
            run.font.color.rgb = RGBColor.from_string(c.couleur_titres)
    if sous_titre:
        doc.add_paragraph(sous_titre).runs[0].italic = True
    return doc


def titre(doc: Any, texte: str, niveau: int = 1, charte: Charte | None = None) -> None:
    from docx.shared import RGBColor

    c = charte or charte_par_defaut()
    h = doc.add_heading(texte, level=niveau)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(c.couleur_titres)


def tableau(doc: Any, entete: list[str], lignes: list[list[Any]], charte: Charte | None = None, largeurs_cm: list[float] | None = None) -> Any:
    from docx.shared import Cm, Pt, RGBColor

    c = charte or charte_par_defaut()
    t = doc.add_table(rows=1, cols=len(entete))
    t.style = "Table Grid"
    for cellule, texte in zip(t.rows[0].cells, entete, strict=True):
        cellule.text = str(texte)
        _fond(cellule, c.fond_entete)
        for run in cellule.paragraphs[0].runs:
            run.font.bold, run.font.color.rgb, run.font.size = True, RGBColor(0xFF, 0xFF, 0xFF), Pt(9.5)
    for k, li in enumerate(lignes):
        cellules = t.add_row().cells
        for j, v in enumerate(li):
            cellules[j].text = "" if v is None else str(v)
            for run in cellules[j].paragraphs[0].runs:
                run.font.size = Pt(9.5)
            if k % 2:
                _fond(cellules[j], c.fond_alterne)
    if largeurs_cm:
        for ligne in t.rows:
            for cellule, w in zip(ligne.cells, largeurs_cm, strict=False):
                cellule.width = Cm(w)
    return t


def enregistrer(doc: Any, chemin: Path, en_pdf: bool = False) -> list[Path]:
    """Word (espaces insécables appliqués), puis PDF si demandé (Word ou LibreOffice, comme le rapport)."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    _insecables_docx(doc)
    doc.save(str(chemin))
    out = [chemin]
    if en_pdf:
        out.append(pdf(chemin, chemin.with_suffix(".pdf")))
    return out
