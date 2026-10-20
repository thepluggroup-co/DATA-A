"""Documents de l'enquêteur (M7) : charte (modèle PEINTURE, CLAUDE.md §3.2 : I. comportement, II. consignes de
remplissage, III. argumentaire d'introduction, IV. ordre de mission), ordres de mission individuels, contrats de
prestation et engagements de confidentialité.

Contrat et engagement : MODÈLES à faire valider par un avocat camerounais (statut du vacataire, rémunération,
confidentialité, protection des données, propriété des fiches). La loi n° 2024/017 du 23 décembre 2024 relative à la
protection des données à caractère personnel est citée par son intitulé ; les articles applicables restent à préciser par
l'avocat (texte officiel non encore lu par l'agent, CLAUDE.md §8).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from plateforme.documents.commun import BANDEAU_JURIDIQUE, date_fr, enregistrer, montant, nouveau, tableau, titre
from plateforme.finances.paiements import LignePaiement, fiche
from plateforme.projet.projet import Projet
from plateforme.terrain.plan import EnqueteurPlan, PlanTerrain

LOI_DONNEES = "loi n° 2024/017 du 23 décembre 2024 relative à la protection des données à caractère personnel"


@dataclass(frozen=True)
class Mission:
    """Une personne sur un poste : tout ce que les documents individuels affichent, lu dans les sources."""
    poste: str
    personne: str
    nom: str
    qualite: str
    ville: str
    zones: list[str]
    localites: list[str]
    superviseur: str
    superviseur_tel: str
    debut: date | None
    fin: date | None
    paiement: LignePaiement | None


def missions(p: Projet) -> list[Mission]:
    t = p.exiger_terrain()
    jours = t.jours_ouvres()
    # la mission court de la formation (ou du pré-test) à la fin du terrain
    bornes = [j for j in (t.calendrier.formation, t.calendrier.pre_test, *jours) if j is not None]
    lignes = {(li.poste, li.personne): li for li in fiche(t, p.tarifs, p.vivier).lignes}
    out = []
    for e in t.enqueteurs:
        for pid in _occupants(e):
            pers = p.vivier.personne(pid)
            sup_id = t.ville(e.ville).superviseur if e.ville else None
            sup = next((m for m in t.equipe if m.id == sup_id), None)
            sup_pers = p.vivier.personne(sup.personne) if sup and sup.personne else None
            out.append(Mission(
                poste=e.id, personne=pid, nom=pers.nom_complet,
                qualite=("Enquêtrice" if pers.genre == "F" else "Enquêteur") + (" de réserve" if e.role.value == "reserve" else ""),
                ville=t.ville(e.ville).libelle if e.ville else "", zones=[z.libelle for z in t.zones_de(e)],
                localites=[p.referentiel.localite(x).nom for x in t.localites_de(e)],
                superviseur=sup_pers.nom_complet if sup_pers else "", superviseur_tel=sup_pers.telephone if sup_pers else "",
                debut=min(bornes) if bornes else None, fin=max(bornes) if bornes else None, paiement=lignes.get((e.id, pid))))
    return out


def _occupants(e: EnqueteurPlan) -> list[str]:
    return ([e.personne] if e.personne else []) + [a.personne_id for a in e.affectations if a.personne_id]


def _saut(doc: Any) -> None:
    doc.add_page_break()


# ------------------------------------------------------------------------------------------------ charte
COMPORTEMENT_DOIT = [
    "présenter l'étude et son contexte avant toute question (argumentaire de la partie III) ;",
    "porter son badge de façon visible pendant toute la durée du terrain ;",
    "éclaircir une question mal comprise, en la relisant ou en la reformulant sans en changer le sens ;",
    "laisser l'enquêté répondre librement et noter sa réponse telle qu'il la donne ;",
    "respecter le refus de répondre et l'accord (ou le refus) de recontact, sans insister ;",
    "garder confidentielles les réponses et les coordonnées recueillies.",
]
COMPORTEMENT_NE_DOIT_PAS = [
    "orienter les réponses ou suggérer une modalité ;",
    "imposer l'entretien ou ses propres opinions ;",
    "interpréter une réponse à la place de l'enquêté ;",
    "juger l'enquêté ou ses réponses ;",
    "remplir lui-même une fiche ou recopier une fiche existante.",
]


def _consignes(p: Projet, t: PlanTerrain) -> list[str]:
    d = p.dictionnaire
    villes = ", ".join(v.libelle for v in t.villes)
    out = [
        f"Le questionnaire compte {len(d.questions)} questions réparties en {len(d.sections)} parties ; une seule réponse "
        f"par question, sauf mention « plusieurs réponses ».",
        f"Utiliser uniquement votre lien personnel : il remplit votre identifiant et votre ville ({villes}).",
        "Ne jamais envoyer une fiche incomplète : une fiche sans la question de profil est rejetée (anomalie Q1_VIDE).",
        "Saisir le quartier tel qu'il figure dans la liste de vos zones (ordre de mission).",
        f"Objectif : {round(t.objectif_jour)} fiches par jour et par enquêteur, sans sacrifier la qualité ; un délai trop court "
        f"entre deux fiches (moins de {t.seuil_delai_min} minutes) est contrôlé.",
        "Pendant le pré-test, utiliser le mode test : ces fiches sont exclues de la base.",
    ]
    q = d.question_cible()
    if q is not None:
        out.insert(1, f"La question « {d.texte(q.libelles)} » détermine les quotas : la poser en premier et ne jamais la sauter.")
    return out


def _argumentaire(p: Projet) -> str:
    villes = ", ".join(p.referentiel.localite(v).nom for v in p.brief.villes)
    duree = f" Le questionnaire dure environ {p.brief.duree_entretien} minutes." if p.brief.duree_entretien else ""
    return (f"« Bonjour, je suis enquêteur (enquêtrice) pour {p.cabinet.nom}. Nous réalisons une étude sur le thème : "
            f"{p.brief.titre}, à {villes}. Vos réponses sont confidentielles et analysées de façon anonyme ; elles ne "
            f"servent qu'à cette étude.{duree} Acceptez-vous d'y répondre ? »")


def charte(p: Projet, dossier: Path, en_pdf: bool = False) -> list[Path]:
    t = p.exiger_terrain()
    doc = nouveau("Charte de l'enquêteur", fictif=p.brief.fictif, sous_titre=p.brief.titre, logo=_logo(p))
    titre(doc, "I. Comportement")
    doc.add_paragraph("L'enquêteur doit :")
    for x in COMPORTEMENT_DOIT:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_paragraph("L'enquêteur ne doit pas :")
    for x in COMPORTEMENT_NE_DOIT_PAS:
        doc.add_paragraph(x, style="List Bullet")
    titre(doc, "II. Consignes de remplissage")
    for x in _consignes(p, t):
        doc.add_paragraph(x, style="List Number")
    titre(doc, "III. Argumentaire d'introduction de l'enquête")
    doc.add_paragraph(_argumentaire(p))
    titre(doc, "IV. Ordre de mission")
    doc.add_paragraph("L'ordre de mission individuel (nom, identifiant, ville, zones, dates, superviseur) est remis avec "
                      "la présente charte ; il est signé par le superviseur de ville.")
    doc.add_paragraph("\nJe soussigné(e) ______________________ déclare avoir pris connaissance de la présente charte et "
                      "m'engage à la respecter.\n\nFait à ____________, le ____________        Signature :")
    return enregistrer(doc, dossier / "CHARTE_ENQUETEUR.docx", en_pdf)


def _logo(p: Projet) -> Path | None:
    return p.chemin(p.cabinet.logo) if p.cabinet.logo else None


# ------------------------------------------------------------------------------------------------ ordres de mission
def ordres_de_mission(p: Projet, dossier: Path, en_pdf: bool = False) -> list[Path]:
    t = p.exiger_terrain()
    doc = nouveau("", fictif=p.brief.fictif, logo=_logo(p))
    for k, m in enumerate(missions(p)):
        if k:
            _saut(doc)
        titre(doc, f"Ordre de mission — {m.poste}", 1)
        doc.add_paragraph(f"Étude : {p.brief.titre}")
        tableau(doc, ["Rubrique", "Détail"], [
            ["Nom", m.nom], ["Identifiant", m.poste], ["Qualité", m.qualite], ["Ville", m.ville],
            ["Zones", ", ".join(m.zones) or "—"], ["Quartiers et points de collecte", ", ".join(m.localites) or "—"],
            ["Formation", date_fr(t.calendrier.formation, True) if t.calendrier.formation else "—"],
            ["Pré-test", date_fr(t.calendrier.pre_test, True) if t.calendrier.pre_test else "—"],
            ["Terrain", f"du {date_fr(min(t.jours_ouvres()), True)} au {date_fr(max(t.jours_ouvres()), True)}" if t.jours_ouvres() else "—"],
            ["Objectif", f"{round(t.objectif_jour)} fiches par jour"],
            ["Superviseur de ville", f"{m.superviseur} ({m.superviseur_tel})" if m.superviseur else "—"],
        ], largeurs_cm=[5, 11])
        doc.add_paragraph(f"\nLe porteur est mandaté par {p.cabinet.nom} pour conduire l'enquête ci-dessus, dans les zones "
                          "indiquées et aux dates prévues.\n\nFait à ____________, le ____________\n\n"
                          "Le superviseur de ville                              L'enquêteur (lu et approuvé)")
    return enregistrer(doc, dossier / "ORDRES_DE_MISSION.docx", en_pdf)


# ------------------------------------------------------------------------------------------------ contrats
def contrats(p: Projet, dossier: Path, en_pdf: bool = False) -> list[Path]:
    from docx.shared import Pt, RGBColor

    doc = nouveau("", fictif=p.brief.fictif, logo=_logo(p))
    devise = p.tarifs.devise
    for k, m in enumerate(missions(p)):
        if k:
            _saut(doc)
        r = doc.add_paragraph().add_run(BANDEAU_JURIDIQUE)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0xC0, 0x00, 0x00)
        titre(doc, f"Contrat de prestation de services — enquête de terrain ({m.poste})", 1)
        doc.add_paragraph(f"Entre {p.cabinet.nom}, ci-après « le cabinet », et {m.nom}, ci-après « le prestataire ».")
        articles = [
            ("Article 1 — Objet", f"Le prestataire réalise, pour l'étude « {p.brief.titre} », la collecte de questionnaires "
                                  f"auprès des cibles de l'étude, à {m.ville} ({', '.join(m.zones) or 'zones précisées par l\'ordre de mission'})."),
            ("Article 2 — Durée", f"Du {date_fr(m.debut)} au {date_fr(m.fin)}, formation et pré-test compris (dates de l'ordre de mission)."
             if m.debut and m.fin else "Selon l'ordre de mission."),
            ("Article 3 — Statut", "Le prestataire intervient comme vacataire indépendant, sans lien de subordination permanent "
                                   "[régime à confirmer par l'avocat]."),
            ("Article 4 — Engagements", "Le prestataire respecte la charte de l'enquêteur, porte son badge, n'utilise que son lien "
                                        "personnel et remet des fiches complètes et sincères."),
            ("Article 5 — Confidentialité et données personnelles",
             f"Le prestataire garde confidentielles les réponses et les coordonnées des enquêtés, ne les copie ni ne les "
             f"communique, et les traite conformément à la {LOI_DONNEES} [articles à préciser par l'avocat]. Un engagement de "
             f"confidentialité est signé séparément."),
            ("Article 6 — Propriété des fiches", "Les fiches et les données collectées appartiennent au cabinet et à son client ; "
                                                 "le prestataire n'en conserve aucune copie."),
            ("Article 7 — Fin du contrat", "En cas de départ, de manquement à la charte ou de fiches rejetées au contrôle qualité, "
                                           "le cabinet peut mettre fin à la mission ; les jours réalisés restent dus [modalités à préciser]."),
        ]
        for intitule, texte in articles:
            doc.add_paragraph().add_run(intitule).bold = True
            doc.add_paragraph(texte)
        doc.add_paragraph().add_run("Article 8 — Rémunération (fiche profil des paiements)").bold = True
        li = m.paiement
        if li is not None:
            lignes = [[n, montant(r_.pu, ""), f"{r_.qte:g}", montant(r_.montant, devise)]
                      for n, r_ in (("Formation", li.formation), ("Collecte (jours)", li.collecte), ("Transport", li.transport),
                                    ("Hébergement", li.hebergement), ("Dépouillement", li.depouillement)) if r_.montant]
            if li.honoraires:
                lignes.insert(0, ["Honoraires", "", "", montant(li.honoraires, devise)])
            lignes += [["Total", "", "", montant(li.total, devise)], ["Avance", "", "", montant(li.avance, devise)],
                       ["Reste à payer", "", "", montant(li.reste, devise)]]
            tableau(doc, ["Rubrique", "Prix unitaire", "Quantité", "Montant"], lignes)
        doc.add_paragraph("\nFait en deux exemplaires à ____________, le ____________\n\n"
                          "Pour le cabinet                                       Le prestataire (lu et approuvé)")
    return enregistrer(doc, dossier / "CONTRATS_PRESTATION.docx", en_pdf)


def engagements(p: Projet, dossier: Path, en_pdf: bool = False) -> list[Path]:
    from docx.shared import Pt, RGBColor

    doc = nouveau("", fictif=p.brief.fictif, logo=_logo(p))
    for k, m in enumerate(missions(p)):
        if k:
            _saut(doc)
        r = doc.add_paragraph().add_run(BANDEAU_JURIDIQUE)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0xC0, 0x00, 0x00)
        titre(doc, f"Engagement de confidentialité ({m.poste})", 1)
        doc.add_paragraph(
            f"Je soussigné(e) {m.nom}, {m.qualite.lower()} pour l'étude « {p.brief.titre} » conduite par {p.cabinet.nom}, "
            f"m'engage à :")
        for x in ("ne divulguer à quiconque les réponses, noms, téléphones ou adresses des personnes interrogées ;",
                  "ne conserver aucune copie (photo, capture, papier) des fiches après la fin de ma mission ;",
                  "ne pas utiliser ces informations à d'autres fins que l'étude ;",
                  f"respecter les droits des personnes interrogées prévus par la {LOI_DONNEES} [articles à préciser par l'avocat] ;",
                  "signaler sans délai au superviseur toute perte de téléphone ou fuite de données."):
            doc.add_paragraph(x, style="List Bullet")
        doc.add_paragraph("Cet engagement reste valable après la fin de la mission.\n\nFait à ____________, le ____________"
                          "\n\nSignature (précédée de « lu et approuvé ») :")
    return enregistrer(doc, dossier / "ENGAGEMENTS_CONFIDENTIALITE.docx", en_pdf)
