"""Export des résultats d'analyse : classeur Excel (un tableau par résultat, avec base, pondération et avertissements)
et syntaxe SPSS équivalente (.sps), pour que l'analyste refasse chaque calcul dans SPSS Statistics 27.

Chaque tableau de fréquences porte une ligne TOTAL dont l'effectif est CONTRÔLÉ : il doit égaler la somme des lignes,
sinon l'export est refusé (règle « total général toujours égal à la somme des lignes »).
"""

from __future__ import annotations

import math
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from plateforme.analyse.donnees import JeuDonnees
from plateforme.analyse.procedures import (
    Anova,
    Correspondances,
    Frequences,
    Logistique,
    Moyenne,
    RegressionLineaire,
    ReponsesMultiples,
    Resultat,
    TestStatistique,
    TriCroise,
)
from plateforme.exports.tableaux import verifier_nom
from plateforme.generateurs.excel import ENTETE_FOND, ENTETE_POLICE

GRAS = Font(bold=True)


class TotalIncoherent(ValueError):
    """La ligne TOTAL d'un tableau ne correspond pas à la somme de ses lignes."""


@dataclass
class ResultatTitre:
    titre: str
    resultat: Resultat


def _n(x: float, decimales: int = 1) -> float | None:
    """Nombre arrondi pour la cellule (vide si manquant)."""
    return None if x is None or (isinstance(x, float) and math.isnan(x)) else round(float(x), decimales)


def _entete(ws: Worksheet, titres: Sequence[str]) -> None:
    ws.append(list(titres))
    for c in ws[ws.max_row]:
        c.fill, c.font = ENTETE_FOND, ENTETE_POLICE


def _lignes(r: Resultat, jd: JeuDonnees) -> tuple[list[str], list[list[Any]], list[list[Any]]]:
    """(en-tête, lignes du tableau, lignes de statistiques) pour un résultat."""
    if isinstance(r, Frequences):
        lignes: list[list[Any]] = [[li.code, li.libelle, _n(li.effectif), _n(li.pourcentage), _n(li.ic95[0]), _n(li.ic95[1])] for li in r.lignes]
        somme = sum(li.effectif for li in r.lignes)
        if abs(somme - r.total_pondere) > 1e-6:
            raise TotalIncoherent(f"{r.variable} : TOTAL {r.total_pondere} ≠ somme des lignes {somme}")
        lignes.append(["TOTAL", "", _n(r.total_pondere), 100.0, None, None])
        return ["Code", "Modalité", "Effectif", "% valide", "IC 95 % bas", "IC 95 % haut"], lignes, [["Manquants", r.manquants]]
    if isinstance(r, TriCroise):
        cols = [jd.libelle(r.colonne, c) for c in r.codes_colonnes]
        lignes = []
        for i, code in enumerate(r.codes_lignes):
            lignes.append([code, jd.libelle(r.ligne, code), "Effectif", *[_n(v, 0) for v in r.effectifs[i]]])
            lignes.append(["", "", "% colonne", *[_n(v) for v in r.pct_colonne[i]]])
            if r.residus_ajustes:
                lignes.append(["", "", "Résidu ajusté", *[_n(v, 2) for v in r.residus_ajustes[i]]])
        stats: list[list[Any]] = [["Khi-deux", _n(r.khi2, 3)], ["ddl", r.ddl], ["p", _n(r.p, 4)], ["V de Cramér", _n(r.v_cramer, 3)],
                 ["Part des effectifs théoriques < 5 (%)", _n(100 * r.part_theoriques_faibles)]]
        return ["Code", r.ligne, "", *cols], lignes, stats
    if isinstance(r, Moyenne):
        lignes = [["Ensemble", "", _n(r.moyenne, 2), _n(r.mediane, 2), _n(r.effectif, 0)]]
        lignes += [[code, "", _n(m, 2), _n(med, 2), _n(n, 0)] for code, (m, med, n) in r.groupes.items()]
        return ["Groupe", "Libellé", "Moyenne", "Médiane", "N"], lignes, []
    if isinstance(r, ReponsesMultiples):
        lignes = [[v, jd.variables[v].label if v in jd.variables else v, _n(r.pct_observations[v]), _n(r.pct_reponses[v])] for v in r.variables]
        return ["Variable", "Libellé", "% des observations", "% des réponses"], lignes, []
    if isinstance(r, TestStatistique):
        return ["Statistique", "ddl", "p"], [[_n(r.statistique, 3), r.ddl, _n(r.p, 4)]], [[k, v] for k, v in r.detail.items()]
    if isinstance(r, Correspondances):
        return ["Axe", "% d'inertie"], [[k, _n(v)] for k, v in enumerate(r.part_axes, start=1)], [["Inertie totale", _n(r.inertie_totale, 3)]]
    if isinstance(r, Logistique):
        lignes = [[terme, _n(e, 2), _n(b, 2), _n(h, 2), _n(p, 4)] for terme, (e, b, h, p) in r.coefficients.items()]
        return (["Terme", "Exp(B)", "IC 95 % bas", "IC 95 % haut", "p"], lignes,
                [["R² de Cox-Snell", _n(r.r2_cox_snell, 3)], ["R² de Nagelkerke", _n(r.r2_nagelkerke, 3)]])
    if isinstance(r, RegressionLineaire):
        lignes = [[terme, _n(b, 6), _n(e, 6), _n(p, 4)] for terme, (b, e, p) in r.coefficients.items()]
        return (["Terme", "B", "Erreur type", "p"], lignes,
                [["R²", _n(r.r2, 4)], ["R² ajusté", _n(r.r2_ajuste, 4)], ["Écart type résiduel", _n(r.ecart_type_residuel, 4)],
                 ["F", _n(r.f, 3)], ["ddl", f"{r.ddl[0]} ; {r.ddl[1]}"], ["p", _n(r.p, 4)]])
    if isinstance(r, Anova):
        lignes = [[code, jd.libelle(r.groupe, code), _n(m, 2), n] for code, (m, n) in r.moyennes.items()]
        return (["Groupe", "Libellé", "Moyenne", "N"], lignes,
                [["Somme des carrés inter", _n(r.sc_inter, 4)], ["Somme des carrés intra", _n(r.sc_intra, 4)],
                 ["ddl", f"{r.ddl[0]} ; {r.ddl[1]}"], ["F", _n(r.f, 3)], ["p", _n(r.p, 4)], ["R² (êta²)", _n(r.r2, 4)]])
    raise TypeError(f"résultat non exportable : {type(r).__name__}")


def exporter_excel(resultats: Sequence[ResultatTitre], jd: JeuDonnees, chemin: str | Path) -> Path:
    """Classeur : onglet « Résultats » (un bloc par résultat) et onglet « Syntaxe SPSS »."""
    chemin = Path(chemin)
    verifier_nom(chemin, ".xlsx")
    tableaux = [(rt, _lignes(rt.resultat, jd)) for rt in resultats]        # contrôles AVANT d'écrire quoi que ce soit
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Résultats"
    for k, (rt, (entete, lignes, stats)) in enumerate(tableaux, start=1):
        r = rt.resultat
        ws.append([f"Tableau {k}. {rt.titre}"])
        ws.cell(ws.max_row, 1).font = GRAS
        ws.append(["Procédure", r.procedure, "Base (fiches)", r.base, "Pondération", r.ponderation or "aucune"])
        _entete(ws, entete)
        for li in lignes:
            ws.append(li)
        for li in stats:
            ws.append(li)
        for a in r.avertissements:
            ws.append([f"Attention : {a}"])
        ws.append([])
    syn = wb.create_sheet("Syntaxe SPSS")
    _entete(syn, ["Tableau", "Titre", "Syntaxe SPSS"])
    for k, rt in enumerate(resultats, start=1):
        syn.append([k, rt.titre, rt.resultat.syntaxe_spss])
    wb.save(chemin)
    return chemin


def syntaxe(resultats: Sequence[ResultatTitre], base_attendue: int | None = None, nom_base: str = "") -> str:
    """Syntaxe SPSS complète (UTF-8) reproduisant les résultats, dans l'ordre des tableaux. AUCUN chemin de fichier
    (CLAUDE.md §5) : elle s'exécute sur la base ouverte, dont l'effectif est contrôlé en tête (même contrôle que
    N08_B_COLLECTE_446.sps:19-27)."""
    lignes = ["* Encoding: UTF-8.", "* Syntaxe générée par la plateforme : un bloc par tableau du classeur de résultats.",
              f"* Ouvrir d'abord {nom_base or 'la base d’analyse'} et la mettre au premier plan, puis Exécuter > Tout.",
              "SET TVARS=LABELS TNUMBERS=LABELS OVARS=LABELS ONUMBERS=LABELS.", "WEIGHT OFF.", "FILTER OFF.", "USE ALL."]
    if base_attendue is not None:
        lignes += ["COMPUTE CTRL_UN = 1.", "AGGREGATE /OUTFILE=* MODE=ADDVARIABLES /BREAK=CTRL_UN /CTRL_N=N.",
                   f"COMPUTE CTRL_BASE = (CTRL_N = {base_attendue}).", 'VARIABLE LABELS CTRL_BASE "Contrôle de la base active".',
                   f'VALUE LABELS CTRL_BASE 1 "OK : {base_attendue} fiches" 0 "ERREUR : la base active n’a pas {base_attendue} fiches, '
                   'résultats non valides".',
                   f'TITLE "Contrôle - la base active doit compter {base_attendue} fiches".', "FREQUENCIES VARIABLES=CTRL_BASE.",
                   "DELETE VARIABLES CTRL_UN CTRL_N CTRL_BASE."]
    for k, rt in enumerate(resultats, start=1):
        titre = " ".join(f"Tableau {k}. {rt.titre}".replace('"', "'").split())
        # TITLE est limité à 60 caractères dans SPSS : titre complet en commentaire, TITLE abrégé
        lignes += ["", f"* {titre}.", f'TITLE "{titre[:60]}".', "WEIGHT OFF.", rt.resultat.syntaxe_spss.strip()]
    return "\n".join(lignes) + "\n"


def exporter_syntaxe(resultats: Sequence[ResultatTitre], chemin: str | Path, base_attendue: int | None = None, nom_base: str = "") -> Path:
    chemin = Path(chemin)
    verifier_nom(chemin, ".sps")
    chemin.write_text(syntaxe(resultats, base_attendue, nom_base), encoding="utf-8")
    return chemin


# ------------------------------------------------------------------------------------------------ .sav du jeu d'analyse
class SavInvalide(RuntimeError):
    """Le .sav relu ne correspond pas au jeu d'analyse qui devait être écrit."""


def _format_numerique(serie: Any) -> str:
    """F<largeur>.0 si toutes les valeurs sont entières, sinon F<largeur>.2 (largeur suffisante pour la plus grande)."""
    s = serie.dropna()
    entiers = bool(len(s) == 0 or (s == s.round()).all())
    chiffres = len(f"{abs(s).max():.0f}") if len(s) else 1
    return f"F{max(chiffres + (s < 0).any(), 1)}.0" if entiers else f"F{chiffres + (s < 0).any() + 3}.2"


def exporter_sav_jeu(jd: JeuDonnees, chemin: str | Path, formats: dict[str, str] | None = None,
                     libelle_fichier: str = "", note: str = "") -> Path:
    """Écrit le jeu d'analyse (codes, variables de contrôle, dérivées) en .sav natif : étiquettes de variables et de
    valeurs, valeurs manquantes PAR VARIABLE (3 au plus, limite SPSS), mesures, formats. Le fichier est relu et comparé ;
    tout écart lève SavInvalide. `formats` impose le format SPSS de certaines variables (ex. dates du masque : SDATE10)."""
    import datetime as dt

    import pandas as pd
    import pyreadstat

    from plateforme.exports.sav import _lignes_de_80_octets

    chemin = Path(chemin)
    verifier_nom(chemin, ".sav")
    formats = dict(formats or {})
    cols: dict[str, Any] = {}
    fmts: dict[str, str] = {}
    for nom in jd.table.columns:
        s = jd.table[nom]
        impose = formats.get(nom, "")
        if impose.startswith(("SDATE", "DATE", "EDATE", "ADATE")):
            cols[nom] = s.map(lambda v: dt.date.fromisoformat(str(v)) if isinstance(v, str) and v else (v if isinstance(v, dt.date) else None))
            fmts[nom] = impose
        elif pd.api.types.is_numeric_dtype(s):
            cols[nom] = s.astype("float64")
            fmts[nom] = impose if impose.startswith("F") else _format_numerique(cols[nom])
        else:
            cols[nom] = s.map(lambda v: "" if v is None or (isinstance(v, float) and math.isnan(v)) else str(v))
            fmts[nom] = f"A{max([1, *(len(v.encode('utf-8')) for v in cols[nom])])}"
    df = pd.DataFrame(cols, index=jd.table.index)
    variables = {nom: jd.variables.get(nom) for nom in df.columns}
    trop = [n for n, v in variables.items() if v is not None and len(v.manquants) > 3]
    if trop:
        raise SavInvalide(f"plus de 3 valeurs manquantes discrètes (limite SPSS) : {', '.join(trop)}")
    labels: dict[str, str | None] = {n: (v.label if v else "") for n, v in variables.items()}
    etiquettes = {n: dict(v.valeurs) for n, v in variables.items() if v and v.valeurs and fmts[n].startswith("F")}
    manquants: dict[str, list[Any]] = {n: sorted(v.manquants) for n, v in variables.items() if v and v.manquants and fmts[n].startswith("F")}
    mesures = {n: (v.mesure if v and v.mesure in ("nominal", "ordinal", "scale") else "nominal") for n, v in variables.items()}
    chemin.parent.mkdir(parents=True, exist_ok=True)
    pyreadstat.write_sav(df, str(chemin), file_label=libelle_fichier[:64], note=_lignes_de_80_octets(note) or None,
                         column_labels=labels, variable_value_labels=etiquettes, missing_ranges=manquants,
                         variable_measure=mesures, variable_format=fmts)
    # relecture : tout ce qui devait être écrit doit être relu à l'identique
    relu, meta = pyreadstat.read_sav(str(chemin), user_missing=True)
    pb: list[str] = []
    if list(relu.columns) != list(df.columns) or len(relu) != len(df):
        pb.append("liste des variables ou nombre de lignes")
    else:
        for n in df.columns:
            if (meta.column_names_to_labels.get(n) or "") != labels[n]:
                pb.append(f"{n} : étiquette")
            if n in etiquettes and {float(k): v for k, v in (meta.variable_value_labels.get(n) or {}).items()} != etiquettes[n]:
                pb.append(f"{n} : étiquettes de valeurs")
            relus = sorted(float(r["lo"]) if isinstance(r, dict) else float(str(r)) for r in meta.missing_ranges.get(n, []))
            if relus != [float(x) for x in manquants.get(n, [])]:
                pb.append(f"{n} : valeurs manquantes")
            a, b = df[n].tolist(), relu[n].tolist()
            for x, y in zip(a, b, strict=True):
                vx = x in ("", None) or (isinstance(x, float) and math.isnan(x))
                vy = y in ("", None) or (isinstance(y, float) and math.isnan(y))
                if not (vx and vy) and (vx != vy or (abs(x - y) > 1e-9 if isinstance(x, float) else x != y)):
                    pb.append(f"{n} : valeur modifiée à l'écriture")
                    break
    if pb:
        chemin.unlink(missing_ok=True)
        raise SavInvalide(f"{chemin.name} : " + " ; ".join(pb[:20]))
    return chemin
