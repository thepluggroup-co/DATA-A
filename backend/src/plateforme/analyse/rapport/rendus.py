"""Rendus du rapport : Word (python-docx), présentation de restitution (python-pptx), PDF (LibreOffice ou Word).

Le rendu n'est appelé qu'APRÈS les contrôles bloquants (`produire`). Les chiffres des tableaux sont ceux de l'export
Excel (analyse/exports._lignes) : une seule mise en forme des résultats. Les graphiques sont tracés par matplotlib à
partir des résultats du moteur, jamais à partir de valeurs saisies.

La charte (couleurs, police, logo) est un paramètre : `Charte()` est une charte sobre provisoire, en attendant la
charte de THE PLUG (D-036), qui se charge par `charger_charte`.
"""

from __future__ import annotations

import re
import shutil
import subprocess
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml

from plateforme.analyse import procedures as p
from plateforme.analyse.donnees import JeuDonnees
from plateforme.analyse.exports import _lignes
from plateforme.analyse.rapport.controles import Anomalie, bloquantes, controler
from plateforme.analyse.rapport.modele import Calcul, ElementNumerote, Figure, Paragraphe, Rapport, TableauRapport, Titre, nombre_fr, resoudre
from plateforme.exports.tableaux import verifier_nom


class RapportBloque(RuntimeError):
    """Le rapport porte des anomalies bloquantes : il n'est pas exporté."""

    def __init__(self, anomalies: list[Anomalie]) -> None:
        super().__init__("; ".join(f"{a.code} : {a.message}" for a in anomalies[:10]))
        self.anomalies = anomalies


@dataclass
class Charte:
    nom: str = "charte provisoire (sobre)"
    police: str = "Calibri"
    couleur_titres: str = "1F4E78"                          # hexadécimal RVB, sans « # »
    palette: list[str] = field(default_factory=lambda: ["1F4E78", "2E75B6", "9DC3E6", "C55A11", "F4B183", "7F7F7F"])
    logo: Path | None = None
    couleur_texte: str = "000000"
    fond_entete: str = "1F4E78"                             # fond de la ligne d'en-tête des tableaux (texte blanc)
    fond_alterne: str = "F2F2F2"                            # une ligne sur deux


_CHAMPS_CHARTE = {"nom", "police", "couleur_titres", "palette", "logo", "couleur_texte", "fond_entete", "fond_alterne"}
_HEX = re.compile(r"^[0-9A-Fa-f]{6}$")


# Charte du produit DATA'A (D-042), utilisée quand l'étude ne déclare pas la sienne.
CHARTE_DATA_A = Path(__file__).resolve().parents[5] / "reference" / "charte" / "data_a" / "charte.yaml"


def charte_par_defaut() -> Charte:
    """Charte DATA'A si elle est présente dans le dépôt, sinon la charte sobre intégrée au code."""
    return charger_charte(CHARTE_DATA_A) if CHARTE_DATA_A.exists() else Charte()


def charger_charte(chemin: str | Path) -> Charte:
    brut: dict[str, Any] = yaml.safe_load(Path(chemin).read_text(encoding="utf-8")) or {}
    inconnus = set(brut) - _CHAMPS_CHARTE
    if inconnus:
        raise ValueError(f"charte : champs inconnus {sorted(inconnus)}")
    logo = Path(chemin).parent / brut["logo"] if brut.get("logo") else None
    if logo is not None and not logo.exists():
        raise FileNotFoundError(f"charte : logo introuvable {logo}")
    defaut = Charte()
    couleurs = {k: str(brut.get(k, getattr(defaut, k))) for k in ("couleur_titres", "couleur_texte", "fond_entete", "fond_alterne")}
    palette = [str(c) for c in brut.get("palette") or defaut.palette]
    mauvaises = [c for c in [*couleurs.values(), *palette] if not _HEX.match(c)]
    if mauvaises:
        raise ValueError(f"charte : couleurs invalides {mauvaises} (attendu : 6 chiffres hexadécimaux, sans « # »)")
    return Charte(str(brut.get("nom", "charte")), str(brut.get("police", defaut.police)), couleurs["couleur_titres"], palette, logo,
                  couleurs["couleur_texte"], couleurs["fond_entete"], couleurs["fond_alterne"])


def _fond(cellule: Any, couleur: str) -> None:
    """Fond d'une cellule Word (w:shd), sans dépendre d'un style de tableau du modèle."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    cellule._tc.get_or_add_tcPr().append(shd)


# ------------------------------------------------------------------------------------------------ graphiques
def graphique(r: p.Resultat, jd: JeuDonnees, charte: Charte, forme: str = "barres") -> bytes:
    """PNG du résultat aux couleurs de la charte. `forme` : « barres » (une variable, en %), « secteurs » (une variable),
    « groupees » ou « empilees » (tri croisé, % de colonne)."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    couleurs = ["#" + c for c in charte.palette]
    if isinstance(r, p.Batterie | p.Repartitions | p.Correspondances) or forme == "secteurs_colonnes":
        return _graphique_multiple(r, jd, charte, forme)
    fig, ax = plt.subplots(figsize=(6.3, 3.6), dpi=150)
    if forme == "secteurs":
        if not isinstance(r, p.Frequences):
            plt.close(fig)
            raise TypeError(f"secteurs : une seule variable attendue, pas {type(r).__name__}")
        libelles, valeurs = sans_prefixe_commun([li.libelle for li in r.lignes]), [li.pourcentage for li in r.lignes]
        etiquettes = arrondi_total_100(valeurs)
        secteurs, _ = ax.pie(valeurs, colors=[couleurs[k % len(couleurs)] for k in range(len(valeurs))], startangle=90,
                             counterclock=False)
        total = sum(valeurs) or 1
        for k, (secteur, v, etiquette) in enumerate(zip(secteurs, valeurs, etiquettes, strict=True)):
            angle = np.deg2rad((secteur.theta1 + secteur.theta2) / 2)
            etroit = v / total < 0.05                    # secteur étroit : étiquette à l'extérieur, lisible
            rayon = 1.18 if etroit else 0.72
            fond_clair = _luminance(charte.palette[k % len(charte.palette)]) > 0.6
            ax.text(rayon * np.cos(angle), rayon * np.sin(angle), f"{etiquette} %", ha="center", va="center", fontsize=8,
                    color="#" + charte.couleur_texte if etroit or fond_clair else "white")
        ax.legend(libelles, fontsize=8, frameon=False, loc="center left", bbox_to_anchor=(1, 0.5))
        ax.axis("equal")
    elif isinstance(r, p.TriCroise) and forme in ("groupees", "empilees", "barres"):
        lignes = sans_prefixe_commun([jd.libelle(r.ligne, c) for c in r.codes_lignes])
        colonnes = [jd.libelle(r.colonne, c) for c in r.codes_colonnes]
        if forme == "empilees":                         # une barre par colonne (ville, profil), segments = modalités
            gauche = np.zeros(len(colonnes))
            for i, lib in enumerate(lignes):
                segment = np.array([r.pct_colonne[i][j] for j in range(len(colonnes))])
                ax.barh(colonnes, segment, left=gauche, color=couleurs[i % len(couleurs)], label=lib)
                gauche += segment
            ax.invert_yaxis()
            ax.set_xlabel("% de la colonne")
            ax.legend(fontsize=7, frameon=False, ncol=min(len(lignes), 3), loc="upper center", bbox_to_anchor=(0.5, -0.15))
        else:
            largeur = 0.8 / max(len(colonnes), 1)
            for j, lib in enumerate(colonnes):
                ax.barh(np.arange(len(lignes)) + j * largeur, [r.pct_colonne[i][j] for i in range(len(lignes))], height=largeur,
                        color=couleurs[j % len(couleurs)], label=lib)
            ax.set_yticks(np.arange(len(lignes)) + largeur * (len(colonnes) - 1) / 2, lignes, fontsize=8)
            ax.invert_yaxis()
            ax.set_xlabel("% de la colonne")
            ax.legend(fontsize=7, frameon=False)
    elif forme != "barres":
        plt.close(fig)
        raise TypeError(f"forme « {forme} » non applicable à {type(r).__name__}")
    elif isinstance(r, p.Frequences):
        _barres(ax, [li.libelle for li in r.lignes], [li.pourcentage for li in r.lignes], couleurs[0])
    elif isinstance(r, p.ReponsesMultiples):
        libelles = [jd.variables[v].label if v in jd.variables else v for v in r.variables]
        _barres(ax, libelles, [r.pct_observations[v] for v in r.variables], couleurs[0])
    elif isinstance(r, p.Indicateurs):
        _barres(ax, [lib for _, lib, _, _ in r.lignes], [m for _, _, m, _ in r.lignes], couleurs[0])
    elif isinstance(r, p.Moyenne) and r.groupes:
        _barres(ax, [str(c) for c in r.groupes], [m for m, _, _ in r.groupes.values()], couleurs[0], unite="")
    else:
        plt.close(fig)
        raise TypeError(f"pas de graphique pour {type(r).__name__}")
    for cote in ("top", "right"):
        ax.spines[cote].set_visible(False)
    fig.tight_layout()
    tampon = BytesIO()
    fig.savefig(tampon, format="png", bbox_inches="tight")
    plt.close(fig)
    return tampon.getvalue()


def _secteurs(ax: Any, valeurs: list[float], charte: Charte, titre: str = "") -> None:
    """Un graphique en secteurs : étiquettes au plus fort reste (total 100), lisibles sur tout fond."""
    import numpy as np

    couleurs = ["#" + c for c in charte.palette]
    secteurs, _ = ax.pie(valeurs, colors=[couleurs[k % len(couleurs)] for k in range(len(valeurs))], startangle=90, counterclock=False)
    total = sum(valeurs) or 1
    for k, (secteur, v, etiquette) in enumerate(zip(secteurs, valeurs, arrondi_total_100(valeurs), strict=True)):
        angle = np.deg2rad((secteur.theta1 + secteur.theta2) / 2)
        etroit = v / total < 0.05
        rayon = 1.18 if etroit else 0.72
        fond_clair = _luminance(charte.palette[k % len(charte.palette)]) > 0.6
        ax.text(rayon * np.cos(angle), rayon * np.sin(angle), f"{etiquette} %", ha="center", va="center", fontsize=8,
                color="#" + charte.couleur_texte if etroit or fond_clair else "white")
    ax.axis("equal")
    if titre:
        ax.set_title(titre, fontsize=9)


def _graphique_multiple(r: p.Resultat, jd: JeuDonnees, charte: Charte, forme: str) -> bytes:
    """Figures à plusieurs parties :
    - batterie, forme « groupees » : une grappe par question (ex. gamme), une barre par modalité (ex. tranche de prix) ;
    - batterie, forme « panneaux » (ou défaut) : un panneau de barres par question ; forme « secteurs » : un secteur par question ;
    - tri croisé, forme « secteurs_colonnes » : un secteur par colonne (ex. ville), puis l'ensemble ;
    - répartitions : un secteur par sous-population, côte à côte ;
    - correspondances, forme « carte » : points lignes et colonnes sur les deux premiers axes (normalisation symétrique)."""
    import matplotlib
    matplotlib.use("Agg")
    import textwrap

    import matplotlib.pyplot as plt
    import numpy as np

    couleurs = ["#" + c for c in charte.palette]
    if isinstance(r, p.TriCroise) and forme == "secteurs_colonnes":
        n = len(r.codes_colonnes) + 1
        fig, axes = plt.subplots(1, n, figsize=(2.4 * n, 3.0), dpi=150)
        for j, code in enumerate(r.codes_colonnes):
            _secteurs(axes[j], [r.pct_colonne[i][j] for i in range(len(r.codes_lignes))], charte, jd.libelle(r.colonne, code))
        totaux = [sum(r.effectifs[i]) for i in range(len(r.codes_lignes))]
        _secteurs(axes[-1], totaux, charte, "Ensemble")
        fig.legend(sans_prefixe_commun([jd.libelle(r.ligne, c) for c in r.codes_lignes]), fontsize=7, frameon=False,
                   loc="lower center", ncol=min(len(r.codes_lignes), 4))
    elif isinstance(r, p.Repartitions):
        codes = sorted({li.code for f in r.parts.values() for li in f.lignes})
        n = len(r.parts)
        libelles = [jd.libelle(r.variable, c) for c in codes]
        fig, axes = plt.subplots(1, n, figsize=(3.2 * n, 3.2), dpi=150, squeeze=False)
        for j, (titre, f) in enumerate(r.parts.items()):
            par_code = {li.code: li.pourcentage for li in f.lignes}
            valeurs = [par_code.get(c, 0.0) for c in codes]
            if forme == "secteurs":
                _secteurs(axes[0][j], valeurs, charte, f"{titre} (n = {f.base})")
            else:                                           # un panneau de barres par sous-population, mêmes modalités
                _barres(axes[0][j], libelles, valeurs, couleurs[j % len(couleurs)])
                axes[0][j].set_title(f"{titre} (n = {f.base})", fontsize=9)
                if j:                                       # libellés des modalités sur le seul premier panneau
                    axes[0][j].set_yticklabels([])
                for cote in ("top", "right"):
                    axes[0][j].spines[cote].set_visible(False)
        if forme == "secteurs":
            fig.legend(sans_prefixe_commun(libelles), fontsize=7, frameon=False, loc="lower center", ncol=min(len(codes), 3),
                       bbox_to_anchor=(0.5, -0.08))
    elif isinstance(r, p.Correspondances):
        if forme != "carte":
            raise TypeError(f"forme « {forme} » non applicable à Correspondances (forme « carte »)")
        fig, ax = plt.subplots(figsize=(6.3, 4.6), dpi=150)
        for points, variable, couleur, marque in ((r.points_lignes, r.ligne, couleurs[1 % len(couleurs)], "o"),
                                                  (r.points_colonnes, r.colonne, couleurs[0], "s")):
            for code, xy in points.items():
                x, y = xy[0], (xy[1] if len(xy) > 1 else 0.0)        # tableau 2 × k : un seul axe
                ax.scatter(x, y, color=couleur, marker=marque, s=36)
                ax.annotate(jd.libelle(variable, code), (x, y), textcoords="offset points", xytext=(5, 4), fontsize=8, color=couleur)
        ax.axhline(0, color="#999999", lw=0.6)
        ax.axvline(0, color="#999999", lw=0.6)
        ax.set_xlabel(f"Dimension 1 ({nombre_fr(r.part_axes[0])} % de l'inertie)", fontsize=8)
        if len(r.part_axes) > 1:
            ax.set_ylabel(f"Dimension 2 ({nombre_fr(r.part_axes[1])} % de l'inertie)", fontsize=8)
        for cote in ("top", "right"):
            ax.spines[cote].set_visible(False)
    elif isinstance(r, p.Batterie) and forme == "secteurs":       # un secteur par question, côte à côte
        n = len(r.variables)
        fig, axes = plt.subplots(1, n, figsize=(3.2 * n, 3.6), dpi=150, squeeze=False)
        for k, v in enumerate(r.variables):
            f = r.frequences[v]
            ax = axes[0][k]
            _secteurs(ax, [li.pourcentage for li in f.lignes], charte, textwrap.fill(f"{jd.variables[v].label if v in jd.variables else v} (n = {f.base})", 34))
            ax.legend(sans_prefixe_commun([li.libelle for li in f.lignes]), fontsize=6, frameon=False, loc="upper center",
                      bbox_to_anchor=(0.5, 0.02), ncol=2)
    elif isinstance(r, p.Batterie) and forme == "groupees":
        codes = sorted({li.code for f in r.frequences.values() for li in f.lignes})
        noms = sans_prefixe_commun([jd.variables[v].label if v in jd.variables and jd.variables[v].label else v for v in r.variables])
        fig, ax = plt.subplots(figsize=(6.3, 3.6), dpi=150)
        largeur = 0.8 / max(len(codes), 1)
        for k, code in enumerate(codes):
            parts = [next((li.pourcentage for li in r.frequences[v].lignes if li.code == code), 0.0) for v in r.variables]
            ax.bar(np.arange(len(noms)) + k * largeur, parts, width=largeur, color=couleurs[k % len(couleurs)],
                   label=jd.libelle(r.variables[0], code))
        ax.set_xticks(np.arange(len(noms)) + largeur * (len(codes) - 1) / 2, [textwrap.fill(n, 32) for n in noms], fontsize=8)
        ax.set_ylabel("% des réponses valides")
        ax.legend(fontsize=7, frameon=False)
        for cote in ("top", "right"):
            ax.spines[cote].set_visible(False)
    elif isinstance(r, p.Batterie):
        n = len(r.variables)
        fig, axes = plt.subplots(n, 1, figsize=(6.3, 1.9 * n), dpi=150, squeeze=False)
        for k, v in enumerate(r.variables):
            f = r.frequences[v]
            ax = axes[k][0]
            _barres(ax, [li.libelle for li in f.lignes], [li.pourcentage for li in f.lignes], couleurs[k % len(couleurs)])
            ax.set_title(f"{jd.variables[v].label if v in jd.variables else v} (n = {f.base})", fontsize=8)
            for cote in ("top", "right"):
                ax.spines[cote].set_visible(False)
    else:
        raise TypeError(f"forme « {forme} » non applicable à {type(r).__name__}")
    fig.tight_layout()
    tampon = BytesIO()
    fig.savefig(tampon, format="png", bbox_inches="tight")
    plt.close(fig)
    return tampon.getvalue()


def _luminance(hexa: str) -> float:
    """Luminance relative (0 = noir, 1 = blanc) d'une couleur « RRGGBB » : texte sombre sur fond clair."""
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def arrondi_total_100(pourcentages: list[float]) -> list[int]:
    """Étiquettes entières d'un graphique en secteurs dont la somme fait exactement 100 (méthode du plus fort reste) :
    c'est l'arrondi du rapport EDM livré (V14, graphiques 1 à 8), qui évite un total affiché de 99 ou 101 %."""
    import math as m

    total = sum(pourcentages)
    if total <= 0:
        return [0 for _ in pourcentages]
    parts = [100 * x / total for x in pourcentages]
    entiers = [m.floor(x) for x in parts]
    ordre = sorted(range(len(parts)), key=lambda i: parts[i] - entiers[i], reverse=True)
    for i in ordre[: 100 - sum(entiers)]:
        entiers[i] += 1
    return entiers


def sans_prefixe_commun(libelles: list[str], separateur: str = " - ") -> list[str]:
    """« Q9. Produits achetés - Câbles/fils », « Q9. Produits achetés - Luminaires » → « Câbles/fils », « Luminaires »
    (la question est dans le titre du graphique)."""
    if len(libelles) > 1 and all(separateur in lib for lib in libelles):
        tetes = {lib.split(separateur, 1)[0] for lib in libelles}
        if len(tetes) == 1:
            return [lib.split(separateur, 1)[1] for lib in libelles]
    return libelles


def _barres(ax: Any, libelles: list[str], valeurs: list[float], couleur: str, unite: str = " %") -> None:
    barres = ax.barh(range(len(libelles)), valeurs, color=couleur)
    courts = [re.sub(r"\s*\(%\)$", "", lib) for lib in sans_prefixe_commun(libelles)]      # l'unité est sur les barres
    ax.set_yticks(range(len(libelles)), [lib if len(lib) <= 60 else lib[:59] + "…" for lib in courts], fontsize=8)
    ax.invert_yaxis()
    for b, v in zip(barres, valeurs, strict=True):
        ax.text(b.get_width(), b.get_y() + b.get_height() / 2, f" {nombre_fr(v)}" + unite, va="center", fontsize=7)


# ------------------------------------------------------------------------------------------------ Word
def _cellule(v: Any) -> str:
    """Nombre à la française (virgule décimale) ; les entiers (effectifs, bases) restent entiers."""
    if v is None:
        return ""
    return f"{v}".replace(".", ",") if isinstance(v, float) else str(v)


def _tableau(e: TableauRapport, jd: JeuDonnees, resultats: dict[str, p.Resultat]) -> tuple[list[Any], list[list[Any]], list[list[Any]], int]:
    """(en-tête, lignes, statistiques, première colonne affichée) d'un tableau du rapport. Rapport client : ni codes de
    variables ni codes de modalités (ils restent dans l'export Excel de l'analyste)."""
    if e.contenu:                                      # tableau statique : les références [[…]] des cellules sont résolues
        lignes_s = [[resoudre(c, resultats) for c in li] for li in e.contenu]
        return lignes_s[0], lignes_s[1:], [], 0
    r = e.resultat
    assert r is not None
    if isinstance(r, p.TriCroise) and e.affichage in ("effectifs", "pourcentages"):
        colonnes = [jd.libelle(r.colonne, c) for c in r.codes_colonnes]
        titre = libelle_lisible(jd.variables[r.ligne].label if r.ligne in jd.variables and jd.variables[r.ligne].label else r.ligne)
        if e.affichage == "effectifs":
            lignes = [[jd.libelle(r.ligne, c), *[int(v) for v in r.effectifs[i]], int(sum(r.effectifs[i]))]
                      for i, c in enumerate(r.codes_lignes)]
            totaux = [int(sum(r.effectifs[i][j] for i in range(len(r.codes_lignes)))) for j in range(len(colonnes))]
            return [titre, *colonnes, "Total"], [*lignes, ["Total", *totaux, sum(totaux)]], [], 0
        lignes = [[jd.libelle(r.ligne, c), *[round(v, 1) for v in r.pct_colonne[i]]] for i, c in enumerate(r.codes_lignes)]
        return [titre, *[f"{c} (%)" for c in colonnes]], lignes, [], 0
    entete, lignes, stats = _lignes(r, jd)
    debut = 1 if len(entete) > 1 and entete[0] in ("Variable", "Code") else 0
    entete = [libelle_lisible(jd.variables[h].label) if isinstance(h, str) and h in jd.variables and jd.variables[h].label else h
              for h in entete[debut:]]
    return entete, lignes, stats, debut


def libelle_lisible(libelle: str) -> str:
    """Libellé pour le lecteur : sans le préfixe technique des variables de contrôle (« [C] Enquêteur » → « Enquêteur »)."""
    return re.sub(r"^\[[A-Z]\]\s*", "", libelle)


def _legende(e: ElementNumerote) -> str:
    return f"{'Tableau' if isinstance(e, TableauRapport) else 'Graphique'} {e.numero}. {e.titre}"


_AVANT_PONCTUATION = re.compile(r" ([%:;!?»])")


def insecables(texte: str) -> str:
    """Typographie française : espace insécable avant % : ; ! ? » et après « (jamais « 69,3 » en fin de ligne et « % »
    au début de la suivante)."""
    return _AVANT_PONCTUATION.sub("\u00a0\\1", texte).replace("« ", "«\u00a0")


def _insecables_docx(doc: Any) -> None:
    paragraphes = list(doc.paragraphs)
    for table in doc.tables:
        paragraphes += [par for ligne in table.rows for cellule in ligne.cells for par in cellule.paragraphs]
    for par in paragraphes:
        for run in par.runs:
            if run.text:
                run.text = insecables(run.text)


def _insecables_pptx(prs: Any) -> None:
    for diapo in prs.slides:
        for forme in diapo.shapes:
            cadres = [forme.text_frame] if forme.has_text_frame else []
            if getattr(forme, "has_table", False) and forme.has_table:
                cadres += [c.text_frame for ligne in forme.table.rows for c in ligne.cells]
            for cadre in cadres:
                for par in cadre.paragraphs:
                    for run in par.runs:
                        if run.text:
                            run.text = insecables(run.text)


def word(rapport: Rapport, jd: JeuDonnees, chemin: Path, charte: Charte) -> Path:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = charte.police, Pt(10.5)
    style.font.color.rgb = RGBColor.from_string(charte.couleur_texte)
    couleur = RGBColor.from_string(charte.couleur_titres)
    if charte.logo:
        doc.add_picture(str(charte.logo), width=Cm(4))
    titre = doc.add_heading(rapport.titre, level=0)
    for run in titre.runs:
        run.font.color.rgb = couleur
    doc.add_paragraph(rapport.note_base)
    resultats = rapport.resultats()
    for e in rapport.elements:
        if isinstance(e, Titre):
            h = doc.add_heading(e.texte, level=min(e.niveau, 4))
            for run in h.runs:
                run.font.color.rgb = couleur
        elif isinstance(e, Paragraphe):
            doc.add_paragraph(resoudre(e.texte, resultats))
        elif isinstance(e, TableauRapport | Figure):
            doc.add_paragraph(_legende(e), style="Caption").paragraph_format.keep_with_next = True   # jamais seule en bas de page
            if isinstance(e, TableauRapport) and (e.resultat is not None or e.contenu):
                entete, lignes, stats, debut = _tableau(e, jd, resultats)
                table = doc.add_table(rows=1, cols=len(entete))
                table.style = "Table Grid"
                for cellule, texte in zip(table.rows[0].cells, entete, strict=True):
                    cellule.text = str(texte)
                    _fond(cellule, charte.fond_entete)
                    for run in cellule.paragraphs[0].runs:
                        run.font.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                for k, li in enumerate(lignes):
                    cellules = table.add_row().cells
                    for j, v in enumerate(li[debut: debut + len(entete)]):
                        cellules[j].text = _cellule(v)
                        if k % 2:
                            _fond(cellules[j], charte.fond_alterne)
                for li in stats:                                    # statistiques : libellé + valeur, sans décalage de colonnes
                    cellules = table.add_row().cells
                    cellules[0].text, cellules[min(1, len(cellules) - 1)].text = str(li[0]), _cellule(li[1] if len(li) > 1 else None)
            elif isinstance(e, Figure) and e.resultat is not None:
                doc.add_picture(BytesIO(graphique(e.resultat, jd, charte, e.forme)), width=Cm(15))
                if e.legende:
                    doc.add_paragraph(e.legende)
            if e.note_base:
                doc.add_paragraph(e.note_base).runs[0].font.size = Pt(8.5)
            doc.add_paragraph(f"Source : {rapport.source_affichee(e)}").runs[0].font.size = Pt(8.5)
    calculs = [e for e in rapport.elements if isinstance(e, Calcul)]
    if calculs or rapport.numerotes():   # traçabilité : la requête d'origine de chaque tableau, graphique et chiffre cité
        h = doc.add_heading("Sources des tableaux, graphiques et chiffres cités", level=2)
        for run in h.runs:
            run.font.color.rgb = couleur
        for n in rapport.numerotes():
            if n.source:
                genre = "Tableau" if isinstance(n, TableauRapport) else "Graphique"
                doc.add_paragraph(f"{genre} {n.numero} : {n.source}.").runs[0].font.size = Pt(8.5)
        for c in calculs:
            sous = f" ; sous-population : {c.filtre}" if c.filtre else ""
            doc.add_paragraph(f"{c.identifiant} : {c.source} ; base : {c.resultat.base} ({c.base}){sous}.").runs[0].font.size = Pt(8.5)
    _insecables_docx(doc)
    doc.save(str(chemin))
    return chemin


# ------------------------------------------------------------------------------------------------ PowerPoint
def presentation(rapport: Rapport, jd: JeuDonnees, chemin: Path, charte: Charte) -> Path:
    """Présentation de restitution DISTINCTE du rapport : une diapositive par graphique, par tableau d'indicateurs et
    par section ; les textes sont résolus depuis les mêmes références."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Cm, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    couleur = RGBColor.from_string(charte.couleur_titres)  # type: ignore[no-untyped-call]

    def diapo(titre: str) -> Any:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = titre
        for par in s.shapes.title.text_frame.paragraphs:
            for run in par.runs:
                run.font.color.rgb, run.font.name, run.font.size = couleur, charte.police, Pt(26)
        if charte.logo:
            s.shapes.add_picture(str(charte.logo), Cm(29.5), Cm(0.5), width=Cm(3.5))
        return s

    couverture = diapo(rapport.titre)
    couverture.shapes.add_textbox(Cm(2), Cm(8), Cm(30), Cm(4)).text_frame.text = rapport.note_base
    for e in rapport.elements:
        if isinstance(e, Titre) and e.niveau == 1:
            diapo(e.texte)
        elif isinstance(e, Figure | TableauRapport) and e.resultat is not None:
            try:
                image = graphique(e.resultat, jd, charte, e.forme if isinstance(e, Figure) else "barres")
            except TypeError:
                continue                                       # tableau sans représentation graphique : reste dans le rapport
            s = diapo(_legende(e))
            s.shapes.add_picture(BytesIO(image), Cm(3), Cm(3.5), height=Cm(12.5))
            note = s.shapes.add_textbox(Cm(3), Cm(16.3), Cm(28), Cm(2)).text_frame
            note.text = f"{e.note_base} Source : {rapport.source_affichee(e)}"
            for par in note.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(11)
    _insecables_pptx(prs)
    prs.save(str(chemin))
    return chemin


# ------------------------------------------------------------------------------------------------ PDF
def pdf(docx: Path, chemin: Path) -> Path:
    """Conversion du .docx en PDF : LibreOffice (serveur) si présent, sinon Microsoft Word (poste Windows)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(chemin.parent), str(docx)], check=True, timeout=300)
        produit = chemin.parent / (docx.stem + ".pdf")
        if produit != chemin:
            produit.replace(chemin)
        return chemin
    if shutil.which("powershell"):
        def ps(chemin_: Path) -> str:                    # chaîne PowerShell entre apostrophes (« DATA'A » → « DATA''A »)
            return "'" + str(chemin_.resolve()).replace("'", "''") + "'"
        script = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
                  f"try {{ $d = $w.Documents.Open({ps(docx)}); $d.SaveAs([ref]{ps(chemin)}, [ref]17); $d.Close() }} "
                  "finally { $w.Quit() }")
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script], timeout=300, capture_output=True,
                           text=True, errors="replace")
        if r.returncode == 0 and chemin.exists():
            return chemin
        raise RuntimeError(f"conversion PDF par Word échouée (fichier ouvert ailleurs ?) : {(r.stderr or r.stdout).strip()[:300]}")
    raise RuntimeError("conversion PDF impossible : ni LibreOffice (soffice) ni Microsoft Word disponibles")


# ------------------------------------------------------------------------------------------------ chaîne complète
def produire(rapport: Rapport, jeux: dict[str, JeuDonnees], interdits: list[str], dossier: Path, nom: str, formats: list[str],
             charte: Charte | None = None) -> tuple[dict[str, Path], list[Anomalie]]:
    """Contrôles bloquants, puis rendus demandés (« docx », « pptx », « pdf »). Rien n'est écrit si un contrôle bloque.
    `jeux` : bases du projet ({"analyse": …, "collecte": …}) ; les libellés sont lus dans la base d'analyse."""
    inconnus = set(formats) - {"docx", "pptx", "pdf"}
    if inconnus:
        raise ValueError(f"formats inconnus : {sorted(inconnus)}")
    for extension, fichier in ((".docx", f"{nom}.docx"), (".pptx", f"{nom}_RESTITUTION.pptx"), (".pdf", f"{nom}.pdf")):
        verifier_nom(dossier / fichier, extension)                 # refuse « RAPPORT.docx.docx » avant d'écrire
    anomalies = controler(rapport, jeux, interdits)
    if bloquantes(anomalies):
        raise RapportBloque(bloquantes(anomalies))
    charte = charte or Charte()
    jd = jeux["analyse"]
    dossier.mkdir(parents=True, exist_ok=True)
    ecrits: dict[str, Path] = {}
    if "docx" in formats or "pdf" in formats:
        ecrits["docx"] = word(rapport, jd, dossier / f"{nom}.docx", charte)
    if "pptx" in formats:
        ecrits["pptx"] = presentation(rapport, jd, dossier / f"{nom}_RESTITUTION.pptx", charte)
    if "pdf" in formats:
        ecrits["pdf"] = pdf(ecrits["docx"], dossier / f"{nom}.pdf")
        if "docx" not in formats:
            ecrits.pop("docx").unlink()
    return ecrits, anomalies
