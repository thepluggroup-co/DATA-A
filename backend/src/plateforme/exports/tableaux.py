"""Exports tabulaires : base codée en CSV (UTF-8, en-tête = codes de variables) et fichier CONTACTS séparé.

Le CSV reproduit l'onglet EXPORT_SPSS de V11 (séparateur virgule, comme `GET DATA /DELIMITERS=","` des
syntaxes de référence). Le fichier CONTACTS (coordonnées des répondants) est un fichier À PART, jamais
joint à la base d'analyse : il relève de la protection des données personnelles (loi n° 2024/017).
"""

from __future__ import annotations

import csv
from pathlib import Path

from plateforme.codification.fiche import ENTETE_CONTACTS, FicheCodee
from plateforme.generateurs.masque import Colonne


class NomDeFichierInvalide(ValueError):
    """Extension absente, multiple ou inattendue (cas réel : « EXPORT_SPSS.csv.csv »)."""


EXTENSIONS_CONNUES = {".csv", ".sav", ".zsav", ".sps", ".xlsx", ".xls", ".txt", ".docx", ".pdf"}


def verifier_nom(chemin: Path, extension: str) -> None:
    """Une seule extension, la bonne : refuse « EXPORT_SPSS.csv.csv », « EDM.csv.sav », « EDM » (sans extension)."""
    if (chemin.suffix.lower() != extension or chemin.name.lower().count(extension) != 1
            or Path(chemin.stem).suffix.lower() in EXTENSIONS_CONNUES):
        raise NomDeFichierInvalide(f"nom de fichier refusé : {chemin.name} (une seule extension {extension} attendue)")


def exporter_csv(colonnes: list[Colonne], fiches: list[FicheCodee], chemin: str | Path) -> Path:
    chemin = Path(chemin)
    verifier_nom(chemin, ".csv")
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with chemin.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=",", quoting=csv.QUOTE_MINIMAL)
        w.writerow([c.code for c in colonnes])
        for fiche in fiches:
            w.writerow([str(v).replace("\r", " ").replace("\n", " ") for v in fiche.ligne(colonnes)])
    return chemin


def exporter_masque(colonnes: list[Colonne], fiches: list[FicheCodee], chemin: str | Path) -> Path:
    """Masque de dépouillement (.xlsx) au format V11 : lignes 1 à 3 = section, code, libellé ; puis une fiche par ligne,
    avec l'identifiant de la réponse en dernière colonne (technique, non exporté vers SPSS)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    chemin = Path(chemin)
    verifier_nom(chemin, ".xlsx")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "MASQUE_DEPOUILLEMENT"
    ws.append([c.section for c in colonnes] + ["Technique"])
    ws.append([c.code for c in colonnes] + ["IDREPONSE"])
    ws.append([c.label for c in colonnes] + ["Identifiant de la réponse — ne pas modifier"])
    for cellule in ws[2]:
        cellule.font = Font(bold=True)
    for fiche in fiches:
        ws.append([*fiche.ligne(colonnes), fiche.id_reponse])
    ws.freeze_panes = "C4"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    return chemin


def exporter_contacts(fiches: list[FicheCodee], chemin: str | Path) -> Path:
    """Fichier CONTACTS (.xlsx) : une ligne par fiche, coordonnées et accord de recontact."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    chemin = Path(chemin)
    verifier_nom(chemin, ".xlsx")
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CONTACTS"
    ws.append(list(ENTETE_CONTACTS))
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    for fiche in fiches:
        ws.append(fiche.contact)
    for ligne in ws.iter_rows(min_row=2, min_col=10, max_col=10):      # téléphones en texte (pas de 6,94E+08)
        for cellule in ligne:
            cellule.number_format = "@"
    ws.freeze_panes = "A2"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)
    return chemin
