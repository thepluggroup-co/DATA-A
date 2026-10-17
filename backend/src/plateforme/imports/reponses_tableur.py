"""Lecture des réponses brutes d'un tableur de réponses Google Forms (onglet « Form Responses 1 », export .xlsx).

Pièges constatés sur l'étude EDM et gérés ici :
- l'en-tête peut ne pas être en première ligne (une réponse déplacée au-dessus par un tri) : la ligne
  d'en-tête est repérée par ses titres « Qn. … », puis toutes les autres lignes sont des réponses ;
- une grille occupe une colonne par ligne, titrée « Titre [Ligne] » ; une même grille recréée dans le
  formulaire donne un second bloc de colonnes, que Google Sheets suffixe « [Ligne] 2 », traité comme un
  second élément ;
- les cases à cocher sont jointes par « , » : elles sont redécoupées d'après les libellés du dictionnaire
  (un libellé peut lui-même contenir une virgule) ;
- un numéro de téléphone peut être stocké comme nombre (694388296.0) : il est remis en texte (694388296) ;
- l'horodatage est exprimé dans le fuseau du classeur (Africa/Douala pour l'étude EDM).

Le tableur ne contient pas l'identifiant Google de la réponse : l'identifiant est dérivé de l'horodatage
(précis à la milliseconde) et d'une empreinte du contenu de la ligne, donc stable si l'onglet est trié.
"""

from __future__ import annotations

import hashlib
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from plateforme.codification.reponses import ItemReponse, ReponseBrute, Valeur
from plateforme.dictionnaire.modele import Dictionnaire, TypeQuestion
from plateforme.noyau.textes import cle, normaliser, numero_question

TITRES_HORODATAGE = {"timestamp", "horodateur", "horodatage", "marquehorodateur"}
# « Titre [Ligne] », ou « Titre [Ligne] 2 » quand Google Sheets numérote un titre de colonne en double
RE_LIGNE_GRILLE = re.compile(r"^(.*?)\s*\[([^\]]*)\](?:\s+\d+)?\s*$")


class TableurInvalide(ValueError):
    """Structure de tableur non reconnue."""


def lire_onglet(chemin: str | Path, onglet: str) -> list[list[object]]:
    """Lit UN onglet (les autres ne sont pas chargés en mémoire)."""
    from openpyxl import load_workbook
    wb = load_workbook(chemin, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[onglet].iter_rows(values_only=True)]
    finally:
        wb.close()


def _texte(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _decouper_cases(texte: str, libelles: list[str]) -> list[str]:
    """« Câbles/fils, Luminaires » → ["Câbles/fils", "Luminaires"], en respectant les libellés qui contiennent « , »."""
    connus = {cle(x) for x in libelles}
    morceaux = [m for m in texte.split(", ")]
    out: list[str] = []
    i = 0
    while i < len(morceaux):
        for j in range(len(morceaux), i, -1):                  # plus long assemblage reconnu d'abord
            candidat = ", ".join(morceaux[i:j])
            if cle(candidat) in connus:
                out.append(candidat)
                i = j
                break
        else:
            out.append(morceaux[i])
            i += 1
    return [x for x in out if normaliser(x)]


def lire_reponses(lignes: list[list[object]], dico: Dictionnaire, fuseau: str = "Africa/Douala",
                  source: str = "") -> list[ReponseBrute]:
    i_entete = next((i for i, li in enumerate(lignes) if sum(numero_question(_texte(v)) is not None for v in li) >= 3), None)
    if i_entete is None:
        raise TableurInvalide("aucune ligne d'en-tête « Qn. … » trouvée")
    entete = [_texte(v) for v in lignes[i_entete]]
    i_horo = next((i for i, t in enumerate(entete) if cle(t) in TITRES_HORODATAGE), None)
    if i_horo is None:
        raise TableurInvalide("colonne d'horodatage introuvable")

    # Blocs de colonnes : (titre, [indices], [libellés de lignes de grille] ou None)
    blocs: list[tuple[str, list[int], list[str] | None]] = []
    for i, t in enumerate(entete):
        if i == i_horo or not t:
            continue
        m = RE_LIGNE_GRILLE.match(t)
        if m:
            base, ligne = normaliser(m.group(1)), normaliser(m.group(2))
            dernier = blocs[-1] if blocs else None
            if dernier and dernier[2] is not None and dernier[0] == base and dernier[1][-1] == i - 1 and ligne not in dernier[2]:
                dernier[1].append(i)
                dernier[2].append(ligne)
            else:
                blocs.append((base, [i], [ligne]))
        else:
            blocs.append((t, [i], None))

    libelles_multiples = {q.numero: [t for m in q.modalites for t in m.libelles.values()]
                          for q in dico.questions if q.type is TypeQuestion.CHOIX_MULTIPLE}
    tz = ZoneInfo(fuseau)
    reponses = []
    for n_ligne, li in enumerate(lignes):
        if n_ligne == i_entete or all(v in (None, "") for v in li):
            continue
        horo = li[i_horo] if i_horo < len(li) else None
        if isinstance(horo, str):
            horo = datetime.fromisoformat(horo.strip())
        if not isinstance(horo, datetime):
            raise TableurInvalide(f"ligne {n_ligne + 1} : horodatage illisible")
        horo = horo.replace(tzinfo=tz) if horo.tzinfo is None else horo
        items = []
        for titre, cols, lignes_grille in blocs:
            valeurs = [_texte(li[c]) if c < len(li) else "" for c in cols]
            valeur: Valeur
            if lignes_grille is not None:
                valeur = valeurs
            else:
                n = numero_question(titre)
                valeur = _decouper_cases(valeurs[0], libelles_multiples[n]) if n in libelles_multiples else valeurs[0]
            items.append(ItemReponse(titre, valeur, tuple(lignes_grille) if lignes_grille is not None else None))
        empreinte = hashlib.sha1("\x1f".join(_texte(v) for v in li).encode("utf-8")).hexdigest()[:12]
        rid = f"tableur:{horo.isoformat(timespec='milliseconds')}:{empreinte}"     # stable même si l'onglet est trié
        reponses.append(ReponseBrute(id=rid, horodatage=horo, items=tuple(items), source=source))
    return reponses
