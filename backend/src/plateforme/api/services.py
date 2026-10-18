"""Services de la console « usine » : les mêmes fonctions que la ligne de commande, renvoyées en JSON.

Règles :
- aucune donnée personnelle de répondant (nom, téléphone, e-mail) dans une réponse JSON ; elles ne sortent que
  dans les archives téléchargées (CONTACTS, RAPPELS), réservées à l'usine ;
- le tableur de réponses reçu n'est jamais conservé : il est lu en mémoire puis oublié ;
- un projet = un dossier `projets/<id>/` contenant les trois sources de vérité (dictionnaire, plan d'échantillonnage,
  plan de terrain).
"""

from __future__ import annotations

import base64
import dataclasses
import io
import math
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, time
from enum import Enum
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import yaml

from plateforme.codification.chargement import coder
from plateforme.codification.fiche import FicheCodee
from plateforme.codification.qualite import ParametresBase, base_analyse, quasi_doublons
from plateforme.codification.rejeu import HORS_COMPTAGE
from plateforme.codification.reponses import ReponseBrute
from plateforme.controle.effet_enqueteur import detecter
from plateforme.controle.rappels import tirer_rappels
from plateforme.dictionnaire import diagnostic, diff, fichiers
from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs import formulaire
from plateforme.generateurs.masque import colonnes_masque
from plateforme.imports.reponses_tableur import lire_reponses
from plateforme.suivi.base import Enqueteur, Ville, base_suivi
from plateforme.suivi.livrables import charger_quotas
from plateforme.suivi.rapport import ParametresRapport, calculer_rapport
from plateforme.suivi.statut import rapport_statut
from plateforme.suivi.tableaux import ParametresSuivi, recap_jours, suivi_cibles, tableau_de_bord
from plateforme.terrain import plan as plan_terrain
from plateforme.terrain.plan import PlanTerrain, Role

RACINE = Path(__file__).resolve().parents[4]
RE_PROJET = re.compile(r"^[a-z0-9_]{1,64}$")
FICHIERS_PROJET = ("dictionnaire.yaml", "plan_terrain.yaml", "plan_echantillonnage.yaml")


class ProjetIntrouvable(LookupError):
    """Identifiant de projet inconnu ou mal formé."""


def dossier_projets() -> Path:
    return Path(os.environ.get("PLATEFORME_PROJETS", RACINE / "projets"))


def dossier(projet: str) -> Path:
    if not RE_PROJET.match(projet):
        raise ProjetIntrouvable(projet)
    d = dossier_projets() / projet
    if not all((d / f).is_file() for f in FICHIERS_PROJET):
        raise ProjetIntrouvable(projet)
    return d


def charger(projet: str) -> tuple[Dictionnaire, PlanTerrain, dict[int, tuple[str, list[int]]]]:
    d = dossier(projet)
    terrain = plan_terrain.charger(d / "plan_terrain.yaml")
    return fichiers.charger(d / "dictionnaire.yaml"), terrain, charger_quotas(d / "plan_echantillonnage.yaml", terrain)


def jsonable(x: Any) -> Any:
    """Convertit dataclasses, dates, énumérations et flottants non finis en JSON strict."""
    if dataclasses.is_dataclass(x) and not isinstance(x, type):
        return {f.name: jsonable(getattr(x, f.name)) for f in dataclasses.fields(x)}
    if isinstance(x, dict):
        return {str(k): jsonable(v) for k, v in x.items()}
    if isinstance(x, list | tuple | set | frozenset):
        return [jsonable(v) for v in x]
    if isinstance(x, datetime | date | time):
        return x.isoformat()
    if isinstance(x, Enum):
        return x.value
    if isinstance(x, float) and not math.isfinite(x):
        return None
    return x


# ---------------------------------------------------------------- projets, dictionnaire, terrain

def projets() -> list[dict[str, Any]]:
    out = []
    racine = dossier_projets()
    for d in sorted(p for p in racine.iterdir() if p.is_dir()) if racine.is_dir() else []:
        try:
            out.append(resume(d.name))
        except (ProjetIntrouvable, ValueError):
            continue
    return out


def resume(projet: str) -> dict[str, Any]:
    dico, terrain, quotas = charger(projet)
    titulaires = [e for e in terrain.enqueteurs if e.role is Role.TITULAIRE]
    return {
        "id": projet, "etude": dico.etude, "titre": dico.titre, "version_dictionnaire": dico.version, "statut_dictionnaire": dico.statut,
        "langues": dico.langues, "questions": len(dico.questions), "variables": sum(len(q.variables()) for q in dico.questions),
        "colonnes_masque": len(colonnes_masque(dico, terrain.villes_masque())),
        "villes": [{"id": v.id, "libelle": v.libelle, "prefixe": v.prefixe} for v in terrain.villes],
        "postes": len(titulaires), "reserves": len(terrain.enqueteurs) - len(titulaires),
        "objectif_total": terrain.objectif_total, "jours": [j.isoformat() for j in terrain.jours],
        "quota_total": sum(sum(q[1]) for q in quotas.values()),
    }


def dictionnaire(projet: str) -> dict[str, Any]:
    dico, _, _ = charger(projet)
    brut: dict[str, Any] = dico.model_dump(mode="json")
    brut["nb_variables"] = {q.code: len(q.variables()) for q in dico.questions}
    return brut


def diagnostic_formulaire(projet: str) -> dict[str, Any]:
    dico, terrain, _ = charger(projet)
    diag = diagnostic.diagnostiquer(dico, formulaire.plan_formulaire(dico, [e.id for e in terrain.enqueteurs]))
    return {"verdict": diag.verdict(), "blocages": diag.blocages, "alertes": diag.alertes,
            "chemin_sans_cible": diag.chemin_sans_cible, "lignes": jsonable(diag.lignes)}


def versions(projet: str) -> list[str]:
    d = dossier(projet) / "versions"
    return sorted(p.name for p in d.glob("*.yaml")) if d.is_dir() else []


def difference(projet: str, ancienne: str) -> dict[str, Any]:
    if ancienne not in versions(projet):
        raise ProjetIntrouvable(ancienne)
    d = dossier(projet)
    a, b = fichiers.charger(d / "versions" / ancienne), fichiers.charger(d / "dictionnaire.yaml")
    ecarts = diff.comparer(a, b)
    return {"ancienne": a.version, "nouvelle": b.version, "migration": sum(e.migration for e in ecarts),
            "ecarts": [{"type": e.type.value, "question": e.question, "detail": e.detail, "migration": e.migration} for e in ecarts],
            "texte": diff.rapport_texte(a, b, ecarts)}


def terrain(projet: str) -> dict[str, Any]:
    _, t, quotas = charger(projet)
    return {**t.model_dump(mode="json"),
            "quotas": [{"code": c, "cible": lib, "villes": dict(zip([v.id for v in t.villes], q, strict=True)), "total": sum(q)}
                       for c, (lib, q) in sorted(quotas.items())]}


def liens(projet: str, modele: str, enqueteur_modele: str, villes_modele: list[str]) -> list[dict[str, Any]]:
    from plateforme.deploiement.liens import analyser_modele, lien, qr_code

    _, t, _ = charger(projet)
    libelles = {v.id: v.libelle for v in t.villes}
    m = analyser_modele(modele, enqueteur_modele, villes_modele)
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        for e in t.enqueteurs:
            url = lien(m, e.id, libelles.get(e.ville or ""))
            png = qr_code(url, Path(tmp) / f"{e.id}.png")
            out.append({"id": e.id, "ville": e.ville, "role": e.role.value, "lien": url,
                        "qr_png": "data:image/png;base64," + base64.b64encode(png.read_bytes()).decode("ascii")})
    return out


# ---------------------------------------------------------------- réponses : lecture, codage, suivi

def lire_tableur(contenu: bytes, onglet: str, dico: Dictionnaire, t: PlanTerrain, source: str) -> list[ReponseBrute]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    try:
        if onglet not in wb.sheetnames:
            raise ValueError(f"onglet « {onglet} » absent (onglets : {', '.join(wb.sheetnames)})")
        lignes: list[list[object]] = [list(r) for r in wb[onglet].iter_rows(values_only=True)]
    finally:
        wb.close()
    return lire_reponses(lignes, dico, fuseau=t.fuseau, source=source)


def lire_decisions(contenu: bytes | None) -> dict[str, dict[str, Any]] | None:
    if not contenu:
        return None
    brut = yaml.safe_load(contenu.decode("utf-8")) or {}
    if not isinstance(brut, dict):
        raise ValueError("décisions : un dictionnaire {N° de fiche: {statut, motif}} est attendu")
    return brut


def _moment(f: FicheCodee, tz: ZoneInfo) -> datetime:
    return datetime.fromisoformat(f"{f.valeurs['DATEENQU']}T{f.valeurs['HEUREENQ']}:00").replace(tzinfo=tz)


def reference_temps(fiches: list[FicheCodee], t: PlanTerrain, jour: str | None, maintenant: str | None) -> tuple[str, datetime]:
    """Par défaut : le dernier jour où des fiches existent, vu à 23 h 59 (rapport définitif)."""
    tz = ZoneInfo(t.fuseau)
    if maintenant:
        m = datetime.fromisoformat(maintenant).replace(tzinfo=tz)
        return jour or m.strftime("%Y-%m-%d"), m
    if not jour:
        dates = sorted(str(f.valeurs.get("DATEENQU")) for f in fiches if f.valeurs.get("DATEENQU"))
        jour = dates[-1] if dates else datetime.now(tz).strftime("%Y-%m-%d")
    return jour, datetime.fromisoformat(f"{jour}T23:59:00").replace(tzinfo=tz)


def analyser(projet: str, reponses: list[ReponseBrute], brutes: dict[str, dict[str, Any]] | None,
             jour: str | None = None, maintenant: str | None = None, graine: int = 0) -> dict[str, Any]:
    """Tout le suivi d'un jour : tableau de bord, quotas, rapport, statut, qualité, effet enquêteur, note de base."""
    dico, t, quotas = charger(projet)
    toutes, decisions = coder(dico, t, reponses, brutes)
    jour, m = reference_temps(toutes, t, jour, maintenant)
    tz = ZoneInfo(t.fuseau)
    fiches = [f for f in toutes if _moment(f, tz) <= m]
    jours = [j.isoformat() for j in t.jours]
    villes = [Ville(v.code, v.id, v.libelle) for v in t.villes]
    enq = [Enqueteur(e.id, "", e.ville or "") for e in t.enqueteurs]
    base = base_suivi(fiches)
    ps = ParametresSuivi(enq, villes, jours, quotas, t.objectif_jour, t.nb_jours, t.seuil_volume_jour, m.strftime("%Y-%m-%d"))
    pr = ParametresRapport(jour, jours, t.nb_jours, enq, villes, quotas, t.objectif_jour, t.seuil_delai_min, t.seuil_volume_jour,
                           (m.strftime("%Y-%m-%d"), m.hour + m.minute / 60))
    cols = colonnes_masque(dico, t.villes_masque())
    titulaires = [e for e in t.enqueteurs if e.role is Role.TITULAIRE]
    statut = jsonable(rapport_statut(fiches, t, jours, m)) if t.objectif_total and titulaires else None
    retenues, note = base_analyse(dico, fiches, ParametresBase(), decisions)
    ecarts = sorted(detecter(dico, retenues), key=lambda x: (not x.signale, x.p))
    valides = [f for f in fiches if f.statut not in HORS_COMPTAGE]
    du_jour = [f for f in valides if f.valeurs.get("DATEENQU") == jour]
    rappels = tirer_rappels(du_jour, 3, graine, quasi_doublons(dico, valides, ParametresBase()))
    anomalies: dict[str, int] = {}
    for f in fiches:
        for a in f.anomalies:
            k = a.split(":")[0]
            anomalies[k] = anomalies.get(k, 0) + 1
    statuts: dict[str, int] = {}
    for f in fiches:
        statuts[f.statut] = statuts.get(f.statut, 0) + 1
    resultat: dict[str, Any] = jsonable({
        "projet": projet, "jour": jour, "maintenant": m, "reponses": len(reponses), "fiches": len(fiches),
        "note_de_base": {"texte": note.texte(), "collectees": note.collectees, "retenues": note.retenues, "valides": note.valides,
                         "ecartees": note.ecartees},
        "tableau_de_bord": tableau_de_bord(base, ps),
        "suivi_cibles": suivi_cibles(base, ps),
        "recap_jours": recap_jours(base, ps),
        "rapport": calculer_rapport(dico, cols, base, pr),
        "statut": statut,
        "qualite": {
            "statuts": statuts, "anomalies": dict(sorted(anomalies.items(), key=lambda kv: -kv[1])),
            "a_verifier": [{"numfiche": f.numfiche, "enqueteur": f.valeurs.get("CODEENQU", ""), "date": f.valeurs.get("DATEENQU", ""),
                            "heure": f.valeurs.get("HEUREENQ", ""), "anomalies": f.anomalies, "statut": f.statut}
                           for f in fiches if f.statut == "A_VERIFIER"],
            "rappels": [{"numfiche": r.numfiche, "enqueteur": r.enqueteur, "priorite": r.priorite, "motif": r.motif} for r in rappels],
            "graine_rappels": graine,
        },
        "effet_enqueteur": {
            "tests": len(ecarts),
            "signales": [{"groupe": str(x.groupe), "enqueteur": x.enqueteur, "variable": x.variable, "libelle": x.libelle, "n": x.n,
                          "khi2": x.khi2, "ddl": x.ddl, "p_corrige": x.p_corrige, "w": x.w,
                          "modalite": x.modalite_surrepresentee, "part_enqueteur": x.part_enqueteur, "part_attendue": x.part_attendue}
                         for x in ecarts if x.signale],
        },
    })
    return resultat


# ---------------------------------------------------------------- archives téléchargeables (usine uniquement)

def _zipper(dossier_source: Path) -> bytes:
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(dossier_source.rglob("*")):
            if p.is_file():
                z.write(p, p.relative_to(dossier_source).as_posix())
    return tampon.getvalue()


def archive_generations(projet: str) -> bytes:
    """Masque, syntaxe SPSS, livre de codes, paramétrage, gabarit de tri à plat, requêtes Google Forms, diagnostic."""
    from plateforme.cli import main

    d = dossier(projet)
    with tempfile.TemporaryDirectory() as tmp:
        main(["dictionnaire", "generer", str(d / "dictionnaire.yaml"), "--plan-terrain", str(d / "plan_terrain.yaml"), "--sortie", tmp])
        return _zipper(Path(tmp))


def archive_suivi(projet: str, reponses: list[ReponseBrute], brutes: dict[str, dict[str, Any]] | None,
                  jour: str | None, maintenant: str | None, graine: int) -> bytes:
    """Livrables Excel/Word du suivi. CONTIENT DES TÉLÉPHONES (RAPPELS)."""
    from plateforme.suivi.livrables import produire

    dico, t, quotas = charger(projet)
    fiches, decisions = coder(dico, t, reponses, brutes)
    jour, m = reference_temps(fiches, t, jour, maintenant)
    with tempfile.TemporaryDirectory() as tmp:
        produire(dico, fiches, t, quotas, decisions, jour, m, Path(tmp), graine)
        return _zipper(Path(tmp))


def archive_base(projet: str, reponses: list[ReponseBrute], brutes: dict[str, dict[str, Any]] | None) -> bytes:
    """Masque, base .sav vérifiée, CSV, note de base et CONTACTS (séparé). CONTIENT DES DONNÉES PERSONNELLES (CONTACTS)."""
    from plateforme.exports.sav import exporter_sav
    from plateforme.exports.tableaux import exporter_contacts, exporter_csv, exporter_masque

    dico, t, _ = charger(projet)
    fiches, decisions = coder(dico, t, reponses, brutes)
    retenues, note = base_analyse(dico, fiches, ParametresBase(), decisions)
    cols = colonnes_masque(dico, t.villes_masque())
    with tempfile.TemporaryDirectory() as tmp:
        s = Path(tmp)
        exporter_masque(cols, fiches, s / "MASQUE_DEPOUILLEMENT.xlsx")
        exporter_sav(dico, cols, retenues, s / "BASE_ANALYSE.sav", libelle_fichier=dico.etude, note=note.texte())
        exporter_csv(cols, retenues, s / "BASE_ANALYSE.csv")
        exporter_contacts(fiches, s / "CONTACTS" / "CONTACTS.xlsx")
        (s / "NOTE_DE_BASE.txt").write_text(note.texte() + "\n", encoding="utf-8")
        return _zipper(s)
