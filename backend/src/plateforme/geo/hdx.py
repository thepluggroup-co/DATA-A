"""Import des limites administratives « COD-AB » publiées sur HDX (répertoire tabulaire .xlsx : feuilles cmr_admin0 à
cmr_admin3). On lit les feuilles faisant foi (pas les variantes « _em », à la géométrie retouchée) ; noms français
(`admN_name1`) et anglais (`admN_name`). Licence et provenance : à reporter dans SOURCE.md du pays."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from plateforme.geo.referentiel import UniteAdmin


def lire_hdx(chemin: str | Path, iso3: str, niveau_max: int = 3) -> list[UniteAdmin]:
    wb = load_workbook(str(chemin), read_only=True)
    try:
        unites: list[UniteAdmin] = []
        for n in range(niveau_max + 1):
            nom_feuille = f"{iso3.lower()}_admin{n}"
            if nom_feuille not in wb.sheetnames:
                raise ValueError(f"feuille « {nom_feuille} » absente du fichier HDX")
            lignes = list(wb[nom_feuille].iter_rows(values_only=True))
            entete = [str(c) for c in lignes[0]]
            for brute in lignes[1:]:
                li = dict(zip(entete, brute, strict=False))
                if not li.get(f"adm{n}_pcode"):
                    continue
                nom_fr = li.get(f"adm{n}_name1") or li.get(f"adm{n}_name") or ""
                unites.append(UniteAdmin(niveau=n, pcode=str(li[f"adm{n}_pcode"]), nom=str(nom_fr).strip(),
                                         nom_en=str(li.get(f"adm{n}_name") or "").strip(),
                                         parent=str(li.get(f"adm{n - 1}_pcode") or "") if n else ""))
        return unites
    finally:
        wb.close()
