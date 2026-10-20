"""Phase 6 — amont et documents. Critère d'acceptation (CLAUDE.md §11) : un projet fictif « 3 villes » passe du brief au
badge, à la charte, au contrat, à l'ordre de mission, au chronogramme et à la fiche de paiement sans ressaisie."""

from __future__ import annotations

import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
import yaml
from conftest import RACINE
from pydantic import ValidationError

from plateforme.cli import main
from plateforme.documents import badges, formation
from plateforme.documents.produire import GENERATEURS, ProjetIncoherent, produire
from plateforme.finances.paiements import devis, fiche
from plateforme.projet import projet as pj
from plateforme.projet.brief import Brief
from plateforme.terrain.chronogramme import charger_gabarit, planifier
from plateforme.terrain.plan import Affectation
from plateforme.terrain.vivier import Personne, Statut, Vivier, charger, score

DEMO = RACINE / "projets" / "demo_3_villes"
GABARIT = RACINE / "reference" / "gabarits" / "chronogramme_etude.yaml"
SOURCES = ("brief.yaml", "tarifs.yaml", "cabinet.yaml", "vivier.yaml", "dictionnaire.yaml")


@pytest.fixture
def projet_vierge(tmp_path: Path) -> Path:
    """Copie du projet fictif SANS ses plans : seules les sources saisies (brief, tarifs, cabinet, vivier, dictionnaire)."""
    dossier = tmp_path / "demo"
    dossier.mkdir()
    for f in SOURCES:
        shutil.copy(DEMO / f, dossier / f)
    m = yaml.safe_load((DEMO / "projet.yaml").read_text(encoding="utf-8"))
    m["referentiel"] = str(RACINE / "reference" / "geo" / "CMR")
    m["gabarit_chronogramme"] = str(GABARIT)
    (dossier / "projet.yaml").write_text(yaml.safe_dump(m, allow_unicode=True), encoding="utf-8")
    cab = yaml.safe_load((dossier / "cabinet.yaml").read_text(encoding="utf-8"))
    cab["logo"] = str(RACINE / "reference" / "charte" / "data_a" / "logo_data_a.png")
    (dossier / "cabinet.yaml").write_text(yaml.safe_dump(cab, allow_unicode=True), encoding="utf-8")
    return dossier


# ------------------------------------------------------------------------------------------------ chronogramme (M5)
def test_gabarit_reprend_les_durees_du_chronogramme_edm() -> None:
    """Chronogramme EDM : 24 jours ouvrés, démarrage le lundi 12/10/2026 (tâches et durées du PDF)."""
    c = planifier(charger_gabarit(GABARIT), date(2026, 10, 12), "quanti")
    assert sum(t.jours for t in c.taches) == 24 and [t.tache.id for t in c.taches][:3] == ["DOC", "QUALEXP", "QUEST"]
    assert c.debut == date(2026, 10, 12) and all(t.debut.weekday() < 5 and t.fin.weekday() < 5 for t in c.taches)
    assert "RDV" not in [t.tache.id for t in c.taches]                                   # prise de rendez-vous : quali seulement
    assert "RDV" in [t.tache.id for t in planifier(charger_gabarit(GABARIT), date(2026, 10, 12), "quali").taches]


def test_ancre_avant_le_predecesseur_signalee_et_jours_non_ouvres() -> None:
    g = charger_gabarit(GABARIT)
    c = planifier(g, date(2026, 10, 12), ancres={"formation": (date(2026, 10, 13), date(2026, 10, 13))})
    assert any("Recrutement et formation" in a for a in c.anomalies)
    ferie = planifier(g, date(2026, 10, 12), non_ouvres=[date(2026, 10, 12)])
    assert ferie.debut == date(2026, 10, 13)


# ------------------------------------------------------------------------------------------------ vivier (M6)
def _pers(**kw: object) -> Personne:
    base: dict[str, object] = {"id": "P1", "nom": "X", "prenom": "Y", "ville": "DOUALA", "fictif": True,
                               "historique": [{"statut": "candidat", "date": "2026-10-01"}]}
    base.update(kw)
    return Personne.model_validate(base)


def test_pipeline_du_vivier() -> None:
    ok = _pers(historique=[{"statut": s, "date": f"2026-10-0{k + 1}"} for k, s in
                           enumerate(["candidat", "preselectionne", "convoque", "forme", "teste", "recrute", "reserve", "actif"])])
    assert ok.statut is Statut.ACTIF
    with pytest.raises(ValidationError, match="non permis"):
        _pers(historique=[{"statut": "candidat", "date": "2026-10-01"}, {"statut": "recrute", "date": "2026-10-02"}])
    with pytest.raises(ValidationError, match="motif"):
        _pers(historique=[{"statut": "candidat", "date": "2026-10-01"}, {"statut": "ecarte", "date": "2026-10-02"}])
    with pytest.raises(ValidationError, match="chronologique"):
        _pers(historique=[{"statut": "candidat", "date": "2026-10-05"}, {"statut": "preselectionne", "date": "2026-10-01"}])


def test_score_et_reprise_de_l_historique_de_performance(tmp_path: Path) -> None:
    neuf = _pers(experience=0)
    fiable = _pers(id="P2", experience=5, smartphone=True, piece_verifiee=True, langues=["fr", "en"],
                   performances=[{"etude": "EDM", "fiches": 75, "taux_rejet": 0.02}])
    biaise = _pers(id="P3", experience=5, smartphone=True, piece_verifiee=True, langues=["fr", "en"],
                   performances=[{"etude": "EDM", "fiches": 75, "taux_rejet": 0.02, "effet_enqueteur": True}])
    assert score(fiable) > score(biaise) > score(neuf)
    assert [p.id for p, _ in Vivier(personnes=[neuf, fiable, biaise]).classement()] == ["P2", "P3", "P1"]
    reel = tmp_path / "vivier.yaml"
    reel.write_text(yaml.safe_dump({"personnes": [_pers(fictif=False).model_dump(mode="json")]}), encoding="utf-8")
    with pytest.raises(ValueError, match="registre privé"):
        charger(reel, versionne=True)


# ------------------------------------------------------------------------------------------------ brief et paiements
def test_brief_repartition_et_validation() -> None:
    b = pj.charger(DEMO).brief
    rep = b.repartition()
    assert {v: sum(q.values()) for v, q in rep.items()} == {"DOUALA": 120, "YAOUNDE": 120, "BAFOUSSAM": 60}
    assert rep["DOUALA"] == {1: 60, 2: 30, 3: 30}
    with pytest.raises(ValidationError, match="échéance"):
        Brief.model_validate({**b.model_dump(), "echeance": b.debut})


def test_remplacement_partage_les_jours_et_la_formation_va_au_titulaire() -> None:
    p = pj.charger(DEMO)
    t = p.exiger_terrain()
    jours = t.jours_ouvres()
    bascule = datetime.combine(jours[1], datetime.min.time())
    enq = next(e for e in t.enqueteurs if e.id == "ENQ1")
    remplace = enq.model_copy(update={"personne": None, "affectations": [
        Affectation(personne="titulaire", personne_id=enq.personne, fin=bascule),
        Affectation(personne="remplaçant", personne_id="PER_D3", debut=bascule)]})
    autres = [e if e.id != "ENQ1" else remplace for e in t.enqueteurs if e.personne != "PER_D3"]
    f = fiche(t.model_copy(update={"enqueteurs": autres}), p.tarifs, p.vivier)
    lignes = [li for li in f.lignes if li.poste == "ENQ1"]
    assert [(li.statut, li.collecte.qte, li.formation.montant) for li in lignes] == [
        ("Enquêteur titulaire", 1, 5000), ("Enquêteur remplaçant", 2, 0)]
    assert f.anomalies == []


def test_devis_marge_sur_le_prix() -> None:
    p = pj.charger(DEMO)
    d = devis(fiche(p.exiger_terrain(), p.tarifs, p.vivier), p.tarifs)
    assert abs(d.marge / d.prix_ht - p.tarifs.marge) < 1e-9 and d.cout == sum(v for _, v in d.lignes)


# ------------------------------------------------------------------------------------------------ acceptation
def _textes_docx(chemin: Path) -> str:
    from docx import Document

    doc = Document(str(chemin))
    morceaux = [par.text for par in doc.paragraphs]
    for t in doc.tables:
        for ligne in t.rows:
            for c in ligne.cells:
                morceaux += [par.text for par in c.paragraphs]
                for t2 in c.tables:
                    morceaux += [par.text for li in t2.rows for c2 in li.cells for par in c2.paragraphs]
    return "\n".join(morceaux).replace(" ", " ").replace(" ", " ")


def test_acceptation_phase_6_du_brief_aux_documents_sans_ressaisie(projet_vierge: Path, tmp_path: Path) -> None:
    sortie = tmp_path / "documents"
    # 1. Brief seul → plans (aucune saisie de plan) → sources cohérentes
    assert main(["projet", "initialiser", str(projet_vierge)]) == 0
    p = pj.charger(projet_vierge)
    assert pj.verifier(p) == []
    # 2. Tous les documents, d'une commande
    ecrits, modifiees = produire(p, sortie)
    assert set(ecrits) == {nom for nom, _ in GENERATEURS} and modifiees == []
    for fichiers in ecrits.values():
        assert all(f.exists() and f.stat().st_size > 0 for f in fichiers)
    # 3. Mêmes faits partout : identifiants, noms, dates, montants
    f = fiche(p.exiger_terrain(), p.tarifs, p.vivier)
    textes = {n: _textes_docx(sortie / n) for n in ("BADGES.docx", "ORDRES_DE_MISSION.docx", "CONTRATS_PRESTATION.docx",
                                                     "ENGAGEMENTS_CONFIDENTIALITE.docx", "CHRONOGRAMME.docx", "ORGANISATION_TERRAIN.docx")}
    enqueteurs = [li for li in f.lignes if li.statut.startswith("Enquêteur")]
    assert len(enqueteurs) == 8 and len(badges.badges(p)) == 8 + 3                    # 8 enquêteurs + 3 superviseurs
    for li in enqueteurs:
        for n in ("BADGES.docx", "ORDRES_DE_MISSION.docx", "CONTRATS_PRESTATION.docx", "ENGAGEMENTS_CONFIDENTIALITE.docx",
                  "ORGANISATION_TERRAIN.docx"):
            assert li.nom in textes[n], (li.nom, n)
        assert f"Ordre de mission — {li.poste}" in textes["ORDRES_DE_MISSION.docx"]
        assert f"ID : {li.poste}" in textes["BADGES.docx"]
    for li in enqueteurs:
        if li.total:
            total = f"{round(li.total):,}".replace(",", " ")
            assert f"{total} FCFA" in textes["CONTRATS_PRESTATION.docx"]
    collecte = p.chronogramme().tache("COLLECTE")
    t = p.exiger_terrain()
    assert (collecte.debut, collecte.fin) == (min(t.jours), max(t.jours))
    assert "vendredi 13 novembre 2026 au mardi 17 novembre 2026" in textes["ORDRES_DE_MISSION.docx"]
    assert "13 novembre 2026" in textes["CHRONOGRAMME.docx"] and "Valable du 13 novembre 2026" in textes["BADGES.docx"]
    from openpyxl import load_workbook

    ws = load_workbook(sortie / "PROFIL_PAIEMENTS.xlsx")["PROFIL_PAIEMENTS"]
    valeurs = [[c.value for c in r] for r in ws.iter_rows()]
    assert next(r for r in valeurs if r[0] == "TOTAL GÉNÉRAL")[16] == f.total_general
    assert next(r for r in valeurs if r[0] == "Contrôle")[1] == "somme des lignes = total général"
    # 4. Une source change (taille du brief) : plans et documents suivent, le changement est signalé
    brief = yaml.safe_load((projet_vierge / "brief.yaml").read_text(encoding="utf-8"))
    brief["taille"] = 360
    (projet_vierge / "brief.yaml").write_text(yaml.safe_dump(brief, allow_unicode=True), encoding="utf-8")
    assert main(["projet", "initialiser", str(projet_vierge)]) == 1                     # plans retouchés : pas d'écrasement
    assert main(["projet", "initialiser", str(projet_vierge), "--ecraser"]) == 0
    p2 = pj.charger(projet_vierge)
    _, modifiees = produire(p2, sortie)
    assert {"brief", "plan_echantillonnage", "plan_terrain"} <= set(modifiees)
    assert "Cibles (360 fiches)" in _textes_docx(sortie / "ORGANISATION_TERRAIN.docx")
    # 5. Une incohérence entre sources bloque la production (cas des trois listes de cibles de l'EDM)
    pe = (projet_vierge / "plan_echantillonnage.yaml").read_text(encoding="utf-8")
    (projet_vierge / "plan_echantillonnage.yaml").write_text(pe.replace("cible: Enseignant", "cible: Professeur"), encoding="utf-8")
    with pytest.raises(ProjetIncoherent, match="cibles différentes"):
        produire(pj.charger(projet_vierge), sortie)
    assert main(["projet", "verifier", str(projet_vierge)]) == 1


def test_personne_inconnue_ou_ecartee_refusee(projet_vierge: Path) -> None:
    assert main(["projet", "initialiser", str(projet_vierge)]) == 0
    chemin = projet_vierge / "plan_terrain.yaml"
    texte = chemin.read_text(encoding="utf-8")
    chemin.write_text(texte.replace("personne: PER_B1", "personne: PER_B4"), encoding="utf-8")       # PER_B4 : écarté
    assert any("n'est pas mobilisable" in e for e in pj.verifier(pj.charger(projet_vierge)))
    chemin.write_text(texte.replace("personne: PER_B1", "personne: PER_INCONNU"), encoding="utf-8")
    assert any("personne inconnue" in e for e in pj.verifier(pj.charger(projet_vierge)))


def test_badges_qr_et_quiz() -> None:
    p = pj.charger(DEMO)
    b = badges.badges(p)
    assert b[0].qr == "https://verification.demo.invalid/badge/DEMO_FOURNITURES_SCOLAIRES/SUP_DOUALA"
    assert {x.qualite for x in b} == {"Superviseur de ville", "Superviseure de ville", "Agent enquêteur", "Agent enquêtrice"}
    sans_qr = p.exiger_terrain().model_copy(update={"url_verification": None})
    p.terrain = sans_qr
    p.cabinet = p.cabinet.model_copy(update={"url_verification": None})
    assert all(x.qr is None for x in badges.badges(p))
    questions = formation.quiz(p)
    assert formation.noter(questions, [q.bonne for q in questions]) == (len(questions), True)
    assert formation.noter(questions, [(q.bonne + 1) % len(q.choix) for q in questions])[1] is False


def test_plan_terrain_edm_reste_valide_et_compatible() -> None:
    """Le plan de terrain EDM (phases 2 à 5) se charge sans les nouveaux champs de la phase 6."""
    from plateforme.terrain import plan as plan_terrain

    t = plan_terrain.charger(RACINE / "projets" / "edm_cameroun" / "plan_terrain.yaml")
    assert t.zones == [] and t.baremes == {} and len(t.enqueteurs) == 11
    assert t.jours_ouvres() == t.jours and t.jours[0] + timedelta(days=3) == t.jours[-1]
