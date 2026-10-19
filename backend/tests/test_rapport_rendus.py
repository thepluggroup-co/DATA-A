"""Rapport déclaré → contrôles → Word, PowerPoint, PDF. Données synthétiques (aucune donnée réelle)."""

from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
import pytest

from plateforme.analyse.donnees import JeuDonnees, Variable
from plateforme.analyse.procedures import indicateurs, recapitulatif, synthese_khi2
from plateforme.analyse.rapport.chargeur import RapportInvalide, charger
from plateforme.analyse.rapport.controles import bloquantes, controler
from plateforme.analyse.rapport.rendus import Charte, RapportBloque, charger_charte, graphique, produire

RAPPORT = """
titre: "Étude test du matériel électrique"
etude: TEST
termes_interdits: [boissons]
elements:
  - {section: "Synthèse", niveau: 1}
  - {tableau: "S1", id: s1, titre: "Chiffres clés (en %)", source: "plan : S1", procedure: indicateurs, arguments: {variables: [P_SOL, P_DECU]}}
  - {texte: "[[s1.ligne(P_DECU)]] % des répondants ont déjà été déçus (tableau S1)."}
  - {section: "Résultats", niveau: 1}
  - {figure: "1", id: villes, titre: "Fiches collectées par ville", source: "plan : 1", procedure: frequences, arguments: {variable: VILLE}, base: collecte}
  - {figure: "2", id: marques, titre: "Marque la plus achetée", legende: "Une seule marque par répondant.", source: "plan : 2",
     procedure: frequences, arguments: {variable: MARQ}}
  - {tableau: "1", id: croise, titre: "Marque selon la ville", source: "plan : 3", procedure: tri_croise, arguments: {ligne: MARQ, colonne: VILLE}}
  - {tableau: "2", id: recap, titre: "Récapitulatif", source: "plan : 4", procedure: recapitulatif, arguments: {variables: [VILLE, MARQ], attendu: 2}}
  - {texte: "Schneider est citée par [[marques.pct(1)]] % des [[marques.base]] acheteurs (graphique 2)."}
  - {section: "Conclusions", niveau: 1, role: conclusions}
  - {texte: "Deux marques dominent le marché."}
  - {section: "Recommandations", niveau: 1, role: recommandations}
  - {texte: "Nous recommandons de renforcer la garantie."}
"""


def _jeu(n_bafoussam: int) -> JeuDonnees:
    villes = [1] * 5 + [2] * 5 + [3] * n_bafoussam
    marques = ([1, 1, 2, 1, 99] * 4)[: len(villes)]
    table = pd.DataFrame({"VILLE": villes, "MARQ": marques, "P_SOL": [100.0, 0.0] * (len(villes) // 2) + [0.0] * (len(villes) % 2),
                          "P_DECU": [100.0] * (len(villes) - 3) + [0.0] * 3})
    return JeuDonnees(table, {"VILLE": Variable("VILLE", "Ville", {1.0: "Douala", 2.0: "Yaoundé", 3.0: "Bafoussam"}),
                              "MARQ": Variable("MARQ", "Marque", {1.0: "Schneider", 2.0: "Legrand"}, {99.0}),
                              "P_SOL": Variable("P_SOL", "Achètent du solaire (%)", mesure="scale"),
                              "P_DECU": Variable("P_DECU", "Déjà déçus (%)", mesure="scale")})


@pytest.fixture
def jeux() -> dict[str, JeuDonnees]:
    return {"analyse": _jeu(4), "collecte": _jeu(6)}


@pytest.fixture
def declaration(tmp_path: Path) -> Path:
    chemin = tmp_path / "rapport.yaml"
    chemin.write_text(RAPPORT, encoding="utf-8")
    return chemin


def test_chargement_et_notes_de_base(declaration: Path, jeux: dict[str, JeuDonnees]) -> None:
    rapport, interdits = charger(declaration, jeux, "Base : 16 fiches reçues ; 14 retenues.")
    assert interdits == ["boissons"]
    notes = {e.identifiant: e.note_base for e in rapport.numerotes()}
    assert notes["villes"] == "Base : 16 répondants."                                 # base de collecte
    assert notes["marques"] == "Base : 12 répondants (2 sans réponse valide exclus)."
    assert notes["s1"].startswith("Base : 14 fiches retenues ; base de chaque indicateur de 14 à 14")
    assert controler(rapport, jeux, interdits) == []


def test_rendus(declaration: Path, jeux: dict[str, JeuDonnees], tmp_path: Path) -> None:
    from docx import Document
    from pptx import Presentation

    rapport, interdits = charger(declaration, jeux, "Base : 16 fiches reçues ; 14 retenues.")
    ecrits, _ = produire(rapport, jeux, interdits, tmp_path / "sortie", "RAPPORT_TEST", ["docx", "pptx"])
    paragraphes = Document(str(ecrits["docx"])).paragraphs
    textes = [par.text.replace("\u00a0", " ") for par in paragraphes]                           # espaces insécables du rendu
    legendes = [par.text.replace("\u00a0", " ") for par in paragraphes if par.style.name == "Caption"]
    assert legendes == ["Tableau S1. Chiffres clés (en %)", "Graphique 1. Fiches collectées par ville", "Graphique 2. Marque la plus achetée",
                        "Tableau 1. Marque selon la ville", "Tableau 2. Récapitulatif"]
    assert "78,6 % des répondants ont déjà été déçus (tableau S1)." in textes
    assert "Schneider est citée par 75,0 % des 12 acheteurs (graphique 2)." in textes
    assert textes.count("Source : plan : 3") == 1 and "Base : 12 répondants (2 sans réponse valide exclus)." in textes
    assert not any("[[" in t for t in textes), "aucune référence non résolue dans le rapport"
    diapos = Presentation(str(ecrits["pptx"])).slides
    assert len(diapos) == 1 + 4 + 4, "couverture + 4 sections + S1, 2 graphiques, tableau croisé"


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("powershell")), reason="ni LibreOffice ni Word")
def test_pdf(declaration: Path, jeux: dict[str, JeuDonnees], tmp_path: Path) -> None:
    rapport, interdits = charger(declaration, jeux, "Base : 16 fiches reçues ; 14 retenues.")
    try:
        ecrits, _ = produire(rapport, jeux, interdits, tmp_path, "RAPPORT_TEST", ["pdf"])
    except RuntimeError as e:
        pytest.skip(f"conversion PDF indisponible sur ce poste : {e}")
    assert ecrits["pdf"].read_bytes()[:5] == b"%PDF-" and not (tmp_path / "RAPPORT_TEST.docx").exists()


def test_rapport_bloque_rien_n_est_ecrit(declaration: Path, jeux: dict[str, JeuDonnees], tmp_path: Path) -> None:
    declaration.write_text(RAPPORT.replace("Deux marques dominent le marché.", "Comme pour les boissons, 2 marques dominent."), encoding="utf-8")
    rapport, interdits = charger(declaration, jeux, "Base.")
    with pytest.raises(RapportBloque) as err:
        produire(rapport, jeux, interdits, tmp_path / "sortie", "RAPPORT_TEST", ["docx"])
    assert {a.code for a in err.value.anomalies} == {"TERME_ETRANGER", "CHIFFRE_SANS_LIEN"}
    assert not (tmp_path / "sortie").exists()


def test_tableau_recapitulatif_incomplet(declaration: Path, jeux: dict[str, JeuDonnees]) -> None:
    declaration.write_text(RAPPORT.replace("attendu: 2", "attendu: 3"), encoding="utf-8")
    rapport, interdits = charger(declaration, jeux, "Base.")
    assert [a.code for a in bloquantes(controler(rapport, jeux, interdits))] == ["TABLEAU_INCOMPLET"]


@pytest.mark.parametrize(("remplace", "par"), [("id: croise", "id: marques"), ("base: collecte", "base: autre"),
                                               ("{section: \"Synthèse\", niveau: 1}", "{chapitre: x}")])
def test_declarations_refusees(declaration: Path, jeux: dict[str, JeuDonnees], remplace: str, par: str) -> None:
    declaration.write_text(RAPPORT.replace(remplace, par), encoding="utf-8")
    with pytest.raises(RapportInvalide):
        charger(declaration, jeux, "Base.")


def test_tableaux_composes(jeux: dict[str, JeuDonnees]) -> None:
    jd = jeux["analyse"]
    s = indicateurs(jd, ["P_SOL", "P_DECU"])
    assert [(v, round(m, 1), n) for v, _, m, n in s.lignes] == [("P_SOL", 50.0, 14.0), ("P_DECU", 78.6, 14.0)]
    k = synthese_khi2(jd, ["MARQ"], ["VILLE"])
    assert k.lignes[0][:3] == ("MARQ", "VILLE", 12) and k.avertissements, "petits effectifs théoriques signalés"
    r = recapitulatif(jd, ["VILLE", "MARQ"], attendu=3)
    assert r.lignes[1][2:4] == (12, "Schneider") and r.avertissements == ["incomplet : 2 variables sur 3 attendues"]


def test_charte(tmp_path: Path, jeux: dict[str, JeuDonnees]) -> None:
    from plateforme.analyse.procedures import frequences

    (tmp_path / "charte.yaml").write_text("nom: essai\npolice: Arial\ncouleur_titres: 'AA0000'\npalette: ['AA0000', '333333']\n", encoding="utf-8")
    c = charger_charte(tmp_path / "charte.yaml")
    assert (c.nom, c.police, c.palette[0], c.logo) == ("essai", "Arial", "AA0000", None)
    (tmp_path / "charte.yaml").write_text("logo: absent.png\n", encoding="utf-8")
    with pytest.raises(FileNotFoundError):
        charger_charte(tmp_path / "charte.yaml")
    assert graphique(frequences(jeux["analyse"], "MARQ"), jeux["analyse"], Charte())[:8] == b"\x89PNG\r\n\x1a\n"


def test_base_et_filtre_de_chaque_element_controles(declaration: Path, jeux: dict[str, JeuDonnees]) -> None:
    """Le recalcul de contrôle se fait sur la base de l'élément (collecte) avec son filtre ; une base erronée est vue."""
    declaration.write_text(RAPPORT.replace("procedure: frequences, arguments: {variable: MARQ}}",
                                           "procedure: frequences, arguments: {variable: MARQ}, filtre: \"VILLE == 1\"}"), encoding="utf-8")
    rapport, interdits = charger(declaration, jeux, "Base.")
    assert controler(rapport, jeux, interdits) == []
    rapport.numerotes()[1].base = "analyse"                     # graphique 1 calculé sur la collecte, contrôlé sur l'analyse
    assert "BASE_DIFFERENTE" in {a.code for a in controler(rapport, jeux, interdits)}


def _jeu_marques() -> JeuDonnees:
    """60 acheteurs : 3 villes, 3 marques nettement réparties par ville, dépense positive."""
    villes = [1] * 20 + [2] * 20 + [3] * 20
    marques = [1] * 14 + [2] * 3 + [3] * 3 + [1] * 4 + [2] * 14 + [3] * 2 + [1] * 2 + [2] * 3 + [3] * 15
    depense = [float(10 + k % 7) for k in range(60)]
    table = pd.DataFrame({"VILLE": villes, "MARQ": marques, "DEP": depense})
    return JeuDonnees(table, {"VILLE": Variable("VILLE", "Ville", {1.0: "Douala", 2.0: "Yaoundé", 3.0: "Bafoussam"}),
                              "MARQ": Variable("MARQ", "Marque", {1.0: "Schneider", 2.0: "Legrand", 3.0: "Chint"}),
                              "DEP": Variable("DEP", "Dépense", mesure="scale")})


def test_repartitions_carte_et_parts_en_valeur(tmp_path: Path) -> None:
    from plateforme.analyse.procedures import correspondances, frequences, parts_bootstrap, repartitions
    from plateforme.analyse.rapport.modele import resoudre

    jd = _jeu_marques()
    r = repartitions(jd, "MARQ", {"Douala": "VILLE == 1", "Ensemble": ""})
    assert (r.parts["Douala"].base, r.parts["Ensemble"].base, round(r.parts["Douala"].pct(1), 1)) == (20, 60, 70.0)
    c = correspondances(jd, "MARQ", "VILLE")
    inverse = correspondances(jd, "MARQ", "VILLE", inverser_axes=[1])
    assert all(abs(c.points_colonnes[k][0] + inverse.points_colonnes[k][0]) < 1e-9 for k in c.points_colonnes)    # axe 1 retourné
    assert all(abs(c.points_colonnes[k][1] - inverse.points_colonnes[k][1]) < 1e-9 for k in c.points_colonnes)
    assert len(c.part_axes) == 2 and abs(sum(c.part_axes) - 100) < 1e-6
    v = parts_bootstrap(jd, "MARQ", "DEP", {"Schneider": [1], "Autres": [2, 3]}, tirages=200)
    point, bas, haut = v.parts["Schneider"]
    assert bas < point < haut and abs(point + v.parts["Autres"][0] - 100) < 1e-9
    assert parts_bootstrap(jd, "MARQ", "DEP", {"Schneider": [1]}, tirages=200).parts["Schneider"] == (point, bas, haut)   # graine fixe
    pondere = frequences(jd, "MARQ", "DEP")
    assert abs(point - pondere.pct(1)) < 1e-9
    res = {"r": r, "c": c, "c2": correspondances(jd, "MARQ", "VILLE"), "v": v}
    assert resoudre("[[r.pct(Douala|1)]] % ; [[r.n(Ensemble|)]] ; [[v.bas(Schneider)]]–[[v.haut(Schneider)]] % ; [[c.inertie_sur(c2)]]", res) \
        == f"70,0 % ; 60 ; {round(bas)}–{round(haut)} % ; 1,0"
    for forme, resultat in (("barres", r), ("secteurs", r), ("carte", c)):
        assert graphique(resultat, jd, Charte(), forme)[:4] == b"\x89PNG", forme
    with pytest.raises(TypeError):
        graphique(c, jd, Charte(), "barres")


def test_base_des_repartitions_controlee(tmp_path: Path) -> None:
    from plateforme.analyse.rapport.chargeur import note_base

    jd = _jeu_marques()
    chemin = tmp_path / "rapport.yaml"
    chemin.write_text("""
titre: "Marques"
etude: TEST
elements:
  - {section: "Résultats", niveau: 1}
  - {figure: "1", id: g1, titre: "Marques par ville", source: "plan : 1", procedure: repartitions,
     arguments: {variable: MARQ, groupes: {Douala: "VILLE == 1", Ensemble: ""}}}
  - {section: "Conclusions", niveau: 1, role: conclusions}
  - {texte: "Schneider domine à Douala ([[g1.pct(Douala|1)]] %)."}
  - {section: "Recommandations", niveau: 1, role: recommandations}
  - {texte: "Nous recommandons de suivre Douala."}
""", encoding="utf-8")
    jeux = {"analyse": jd, "collecte": jd}
    rapport, interdits = charger(chemin, jeux, "Base.")
    assert note_base(rapport.numerotes()[0].resultat, "analyse", jd) == "Base : n = 20 (Douala) et 60 (Ensemble)."
    assert controler(rapport, jeux, interdits) == []
    rapport.numerotes()[0].resultat.conditions["Douala"] = "VILLE == 2"       # sous-population déclarée ≠ calculée
    assert "BASE_DIFFERENTE" in {a.code for a in controler(rapport, jeux, interdits)}


def test_hosmer_lemeshow_ex_aequo_dans_le_meme_groupe() -> None:
    """Déciles des probabilités prédites ; des fiches de même probabilité ne sont jamais séparées ; calcul à la main."""
    from scipy.stats import chi2

    from plateforme.analyse.procedures import hosmer_lemeshow

    prob = [0.1] * 10 + [0.2] * 10 + [0.3] * 10 + [0.4] * 10 + [0.5] * 10        # 5 valeurs distinctes : 5 groupes de 10
    y = ([1] + [0] * 9) + ([1] * 4 + [0] * 6) + ([1] * 3 + [0] * 7) + ([1] * 4 + [0] * 6) + ([1] * 5 + [0] * 5)
    attendu = sum((o - 10 * q) ** 2 / (10 * q) + (o - 10 * q) ** 2 / (10 - 10 * q) for o, q in ((1, 0.1), (4, 0.2), (3, 0.3), (4, 0.4), (5, 0.5)))
    stat, ddl, p = hosmer_lemeshow(y, prob)
    assert ddl == 3 and abs(stat - attendu) < 1e-9 and abs(p - chi2.sf(attendu, 3)) < 1e-12


def test_espaces_insecables() -> None:
    from plateforme.analyse.rapport.rendus import insecables

    assert insecables("porte 69,3 % ; le « mot » : oui ?") == "porte 69,3\u00a0%\u00a0; le «\u00a0mot\u00a0»\u00a0: oui\u00a0?"


def test_annexe_questionnaire_generee_depuis_le_dictionnaire(tmp_path: Path) -> None:
    """L'annexe est produite depuis le dictionnaire (jamais ressaisie) ; ses tranches (« <25 ») ne sont pas des chiffres
    d'enquête : pas d'anomalie CHIFFRE_SANS_LIEN ; un texte rédigé avec un chiffre en dur reste bloqué."""
    from plateforme.analyse.rapport.modele import Paragraphe, Titre
    from plateforme.dictionnaire.modele import Dictionnaire

    dico = Dictionnaire.model_validate({
        "etude": "TEST", "version": "01", "sections": [{"id": "S1", "libelles": {"fr": "1. Profil"}}],
        "questions": [{"numero": 1, "code": "AGEREPON", "section": "S1", "libelles": {"fr": "Âge"}, "type": "choix_unique", "mesure": "ORDINAL",
                       "modalites": [{"code": 1, "libelles": {"fr": "<25"}}, {"code": 2, "libelles": {"fr": "25–34"}}]},
                      {"numero": 2, "code": "UNESUGGE", "section": "S1", "libelles": {"fr": "Une suggestion ?"}, "type": "texte", "obligatoire": False}]})
    jd = _jeu_marques()
    chemin = tmp_path / "rapport.yaml"
    chemin.write_text("""
titre: "Annexe"
etude: TEST
elements:
  - {section: "Conclusions", niveau: 1, role: conclusions}
  - {texte: "Le questionnaire figure en annexe."}
  - {section: "Recommandations", niveau: 1, role: recommandations}
  - {texte: "Nous recommandons de le relire."}
  - {section: "Annexe 1. Questionnaire", niveau: 1}
  - {questionnaire: "@dictionnaire", niveau: 2}
""", encoding="utf-8")
    jeux = {"analyse": jd, "collecte": jd}
    rapport, interdits = charger(chemin, jeux, "Base.", contexte={"dictionnaire": dico})
    genere = [e for e in rapport.elements if isinstance(e, Paragraphe) and e.origine == "dictionnaire"]
    assert [e.texte for e in genere] == ["Q1. Âge : ☐ <25 ☐ 25–34", "Q2. Une suggestion ? ……………………"]
    assert any(isinstance(e, Titre) and e.texte == "1. Profil" for e in rapport.elements)
    assert controler(rapport, jeux, interdits) == []
    rapport.elements.append(Paragraphe("Les moins de 25 ans sont 40 %."))
    assert "CHIFFRE_SANS_LIEN" in {a.code for a in controler(rapport, jeux, interdits)}
    with pytest.raises(RapportInvalide):
        charger(chemin, jeux, "Base.")                                         # dictionnaire non fourni


def test_import_d_un_resultat_enregistre(tmp_path: Path) -> None:
    """Variable lue dans un .sav livré et rattachée par N° de fiche (groupes d'un TwoStep non reproductible) ; clé en double
    refusée ; fichier absent : variable simplement non disponible."""
    import pyreadstat

    from plateforme.analyse.plan import Import, charger, importer

    table = pd.DataFrame({"NUMFICHE": ["DLA-ENQ1-001", "DLA-ENQ1-002", "YDE-ENQ4-001"], "VILLE": [1, 1, 2]})
    jd = JeuDonnees(table, {"NUMFICHE": Variable("NUMFICHE"), "VILLE": Variable("VILLE")})
    sav = tmp_path / "livre.sav"
    pyreadstat.write_sav(pd.DataFrame({"NUMFICHE": ["YDE-ENQ4-001", "DLA-ENQ1-001"], "SEG": [2.0, 1.0]}), str(sav),
                         column_labels={"NUMFICHE": "N° de fiche", "SEG": "Groupe d'acheteurs"},
                         variable_value_labels={"SEG": {1.0: "Chasseurs de prix", 2.0: "Exigeants premium"}})
    (tmp_path / "plan.yaml").write_text("imports:\n  - {variable: SEG, fichier: livre.sav, cle: NUMFICHE, source: \"essai\"}\n", encoding="utf-8")
    plan = charger(tmp_path / "plan.yaml")
    assert plan.imports == [Import("SEG", sav.resolve(), "NUMFICHE", "essai")]
    avec = importer(jd, plan.imports[0])
    assert avec.table["SEG"].tolist()[0] == 1.0 and pd.isna(avec.table["SEG"].tolist()[1]) and avec.table["SEG"].tolist()[2] == 2.0
    assert avec.variables["SEG"].valeurs == {1.0: "Chasseurs de prix", 2.0: "Exigeants premium"}
    assert avec.variables["SEG"].label == "Groupe d'acheteurs"
    pyreadstat.write_sav(pd.DataFrame({"NUMFICHE": ["DLA-ENQ1-001", "DLA-ENQ1-001"], "SEG": [1.0, 2.0]}), str(sav))
    with pytest.raises(ValueError, match="double"):
        importer(jd, plan.imports[0])
    (tmp_path / "plan.yaml").write_text("imports:\n  - {variable: SEG, fichier: livre.sav}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="champs attendus"):
        charger(tmp_path / "plan.yaml")


def test_decimales_imposees_et_source_affichee(declaration: Path, jeux: dict[str, JeuDonnees], tmp_path: Path) -> None:
    """« [[x~0]] » arrondit à l'unité ; sous chaque élément, la source de l'étude (ou celle propre à l'élément) ; la requête
    d'origine reste listée dans l'annexe des sources."""
    from docx import Document

    from plateforme.analyse.rapport.modele import resoudre

    texte = RAPPORT.replace("termes_interdits: [boissons]", "termes_interdits: [boissons]\nsource_etude: \"enquête test, oct. 2026\"")
    texte = texte.replace('source: "plan : 2",', 'source: "plan : 2", source_affichee: "relevé des marques",')
    declaration.write_text(texte, encoding="utf-8")
    rapport, interdits = charger(declaration, jeux, "Base.")
    res = rapport.resultats()
    assert resoudre("[[marques.pct(1)]] ; [[marques.pct(1)~0]]", res) == "75,0 ; 75"
    ecrits, _ = produire(rapport, jeux, interdits, tmp_path, "RAPPORT_TEST", ["docx"])
    textes = [par.text.replace("\u00a0", " ") for par in Document(str(ecrits["docx"])).paragraphs]
    assert "Source : enquête test, oct. 2026" in textes and "Source : relevé des marques" in textes
    assert "Source : plan : 1" not in textes and "Graphique 1 : plan : 1." in textes       # requête d'origine en annexe
