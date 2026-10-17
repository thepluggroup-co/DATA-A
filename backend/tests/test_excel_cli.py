"""Classeurs Excel générés, aller-retour du livre de codes, YAML et ligne de commande."""

from __future__ import annotations

from pathlib import Path

from conftest import PROJET_EDM
from openpyxl import load_workbook

from plateforme.cli import main
from plateforme.dictionnaire import fichiers
from plateforme.dictionnaire.diff import TypeEcart, comparer
from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs import excel
from plateforme.generateurs.masque import colonnes_masque
from plateforme.generateurs.tri_a_plat import gabarit
from plateforme.imports.livre_de_codes import depuis_livre_de_codes, lire_xlsx
from plateforme.terrain.plan import PlanTerrain


def test_yaml_aller_retour(dico_edm: Dictionnaire) -> None:
    assert fichiers.depuis_texte(fichiers.vers_texte(dico_edm)) == dico_edm


def test_livre_de_codes_aller_retour(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, tmp_path: Path) -> None:
    """Exporter le livre de codes puis le réimporter ne perd aucune question, modalité ni variable."""
    chemin = tmp_path / "livre.xlsx"
    excel.livre_de_codes(colonnes_masque(dico_edm, terrain_edm.villes_masque()), chemin)
    res = depuis_livre_de_codes(lire_xlsx(chemin), dico_edm.etude, "relu", "xlsx", reference=dico_edm)
    assert res.avertissements == []
    ecarts = comparer(dico_edm, res.dictionnaire)
    # Le livre de codes ne contient ni l'anglais ni les règles de cohérence : seuls ces écarts sont admis.
    assert {e.type for e in ecarts} <= {TypeEcart.LANGUE_SUPPRIMEE, TypeEcart.REGLE_SUPPRIMEE}, ecarts


def test_parametrage_au_format_v001(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, tmp_path: Path) -> None:
    chemin = tmp_path / "parametrage.xlsx"
    excel.parametrage(dico_edm, colonnes_masque(dico_edm, terrain_edm.villes_masque()), chemin)
    wb = load_workbook(chemin, read_only=True)
    assert {"00_PARAMETRES", "03_QUESTIONS", "04_MODALITES", "05_VARIABLES_SPS", "07_CONTROLES"} <= set(wb.sheetnames)
    entete = next(wb["03_QUESTIONS"].iter_rows(values_only=True))
    assert list(entete[:7]) == ["Q_ID", "SECTION", "QUESTION_COURTE", "TYPE", "OBLIGATOIRE", "CONDITION", "VARIABLE_PRINCIPALE"]
    assert len(list(wb["05_VARIABLES_SPS"].iter_rows())) == 68     # en-tête + 67 variables
    wb.close()


def test_gabarit_tri_a_plat(dico_edm: Dictionnaire, terrain_edm: PlanTerrain) -> None:
    tables = gabarit(dico_edm, colonnes_masque(dico_edm, terrain_edm.villes_masque()))
    par_code = {t.variable: t for t in tables}
    assert "CODEZONE" not in par_code                  # la ville sert de colonnes
    assert par_code["VOUSETES"].ligne_total and par_code["VOUSETES"].lignes[-1] == (99, "Non-réponse (question laissée vide)")
    assert par_code["PRACCABL"].lignes == ((1, "Cité"),) and not par_code["PRACCABL"].ligne_total


def test_cli_generer(tmp_path: Path) -> None:
    code = main(["dictionnaire", "generer", str(PROJET_EDM / "dictionnaire.yaml"),
                 "--plan-terrain", str(PROJET_EDM / "plan_terrain.yaml"), "--sortie", str(tmp_path)])
    assert code == 0
    for f in ("colonnes_masque.json", "etiquetage.sps", "livre_de_codes.xlsx", "parametrage.xlsx", "tri_a_plat_gabarit.xlsx",
              "google_forms_requetes.json", "google_forms_actions_apps_script.txt", "diagnostic_formulaire.md"):
        assert (tmp_path / f).stat().st_size > 0, f


def test_cli_valider_et_diff(tmp_path: Path, capsys: object) -> None:
    yaml = str(PROJET_EDM / "dictionnaire.yaml")
    assert main(["dictionnaire", "valider", yaml]) == 0
    assert main(["dictionnaire", "diff", yaml, yaml, "--sortie", str(tmp_path / "d.md")]) == 0
    assert "Aucune différence" in (tmp_path / "d.md").read_text(encoding="utf-8")
