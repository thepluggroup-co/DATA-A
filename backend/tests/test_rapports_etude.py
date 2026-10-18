"""CRITÈRE D'ACCEPTATION (phase 3) : rapports journaliers ÉGAUX à ceux produits pendant l'étude.

Comparaison aux onglets RAPPORT_J1 à RAPPORT_J4 (générés le 08/10/2026 vers 19 h 57) et RAPPORT_09-10 (provisoire,
09/10 à 07 h 06) du masque du 09/10 : synthèse, recommandations, production par ville, classement des enquêteurs
et cibles. Seule normalisation : les noms des enquêteurs (que V11 ajoutait entre parenthèses et que la plateforme
ne stocke pas dans le dépôt) sont retirés des deux côtés.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from test_rapport_oracle import QUOTAS

from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs.masque import colonnes_masque
from plateforme.suivi.base import Enqueteur, Ville, base_suivi
from plateforme.suivi.rapport import ParametresRapport, Rapport, calculer_rapport
from plateforme.terrain.plan import PlanTerrain

MASQUE = Path(__file__).resolve().parents[2] / "reference" / "prive" / "donnees" / "EDM_ELECTRIQUE_CAMEROUN_MASQUE_TEMPS_REEL__7_.xlsx"
pytestmark = pytest.mark.skipif(not MASQUE.exists(), reason="données réelles absentes (reference/prive/, hors Git)")
JOURS = ["2026-10-05", "2026-10-06", "2026-10-07", "2026-10-08"]
SANS_NOM = re.compile(r"(ENQ\d+) \([^)]*\)")


def _sections(onglet: str) -> dict[str, list[list[Any]]]:
    from openpyxl import load_workbook
    wb = load_workbook(MASQUE, read_only=True, data_only=True)
    lignes = [list(li) for li in wb[onglet].iter_rows(values_only=True)]
    wb.close()
    sections: dict[str, list[list[Any]]] = {}
    courante = ""
    for li in lignes:
        t = str(li[0]) if li and li[0] is not None else ""
        m = re.match(r"^(\d)\. ", t)
        if m and (len(li) < 2 or li[1] in (None, "")):
            courante = m.group(1)
            sections[courante] = []
        elif courante and any(v not in (None, "") for v in li):
            sections[courante].append(li)
    return sections


def _egal(a: Any, b: Any) -> bool:
    if isinstance(a, int | float) and isinstance(b, int | float) and not isinstance(a, bool):
        return abs(float(a) - float(b)) < 1e-9
    return SANS_NOM.sub(r"\1", "" if a is None else str(a)) == SANS_NOM.sub(r"\1", "" if b is None else str(b))


def _tableau(lignes: list[list[Any]], n: int) -> list[list[Any]]:
    """Lignes d'un tableau de section (sans l'en-tête), tronquées aux n premières colonnes."""
    return [li[:n] for li in lignes[1:]]


def _rapport(dico: Dictionnaire, terrain: PlanTerrain, jour: str, maintenant: tuple[str, float] | None) -> Rapport:
    from donnees_reelles import charger
    reel = charger(dico, terrain)
    p = ParametresRapport(jour=jour, jours=JOURS, nb_jours=3, enqueteurs=[Enqueteur(e.id, "", e.ville or "") for e in terrain.enqueteurs],
                          villes=[Ville(v.code, v.id, v.libelle) for v in terrain.villes], quotas=QUOTAS, maintenant=maintenant)
    return calculer_rapport(dico, colonnes_masque(dico, terrain.villes_masque()), base_suivi(reel.fiches), p)


@pytest.mark.parametrize(("onglet", "jour", "maintenant"), [
    ("RAPPORT_J1", "2026-10-05", None), ("RAPPORT_J2", "2026-10-06", None), ("RAPPORT_J3", "2026-10-07", None),
    ("RAPPORT_J4", "2026-10-08", None), ("RAPPORT_09-10", "2026-10-09", ("2026-10-09", 7 + 6 / 60)),
])
def test_rapport_egal_a_celui_de_l_etude(onglet: str, jour: str, maintenant: tuple[str, float] | None,
                                         dico_edm: Dictionnaire, terrain_edm: PlanTerrain) -> None:
    s = _sections(onglet)
    r = _rapport(dico_edm, terrain_edm, jour, maintenant)
    ecarts: list[str] = []

    def verifier(nom: str, attendu: list[list[Any]], obtenu: list[list[Any]]) -> None:
        if len(attendu) != len(obtenu):
            ecarts.append(f"{nom} : {len(obtenu)} lignes au lieu de {len(attendu)}")
            return
        for i, (a, o) in enumerate(zip(attendu, obtenu, strict=True)):
            for j, (x, y) in enumerate(zip(a, o, strict=False)):
                if not _egal(x, y):
                    ecarts.append(f"{nom} ligne {i + 1} colonne {j + 1}")

    verifier("synthèse", _tableau(s["1"], 3), [[li[0], li[1], li[2]] for li in r.synthese])
    verifier("recommandations", _tableau(s["2"], 2), r.recos)
    verifier("villes", _tableau(s["3"], 9), [[v["z"].libelle, v["nJ"], v["cumul"], v["quota"], v["cumul"] / v["quota"] if v["quota"] else "",
                                             v["reste"], v["nEnq"], "—" if v["requis"] is None else round(v["requis"] * 10) / 10, v["etat"]]
                                            for v in r.villes])
    # Classement : toutes les colonnes sauf le nom (4e colonne)
    attendu_cl = [li[:3] + li[4:14] for li in _tableau(s["4"], 14) if isinstance(li[0], int | float)]
    obtenu_cl = [[c["rang"], c["id"], c["n"], c["zone"] or "(à saisir)", "" if c["pObj"] is None else c["pObj"], c["cumul"], c["pb"],
                  c["rej"], c["dc"], c["tel"], "" if c["nr"] is None else c["nr"], c["score"], c["avis"]] for c in r.classement]
    verifier("classement", attendu_cl, obtenu_cl)
    verifier("cibles", [li for li in _tableau(s["5"], 7) if li[0] not in (None, "")],
             [[c["lib"], c["nJ"], "" if c["pJ"] is None else c["pJ"], c["cumul"], c["quota"], "" if c["pQ"] is None else c["pQ"], c["reste"]]
              for c in r.cibles])
    assert ecarts == [], ecarts[:15]
