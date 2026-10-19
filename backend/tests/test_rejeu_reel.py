"""CRITÈRE D'ACCEPTATION DE LA PHASE 2 sur les RÉPONSES RÉELLES de l'étude EDM.

- mêmes N° de fiche, mêmes codes, mêmes anomalies et mêmes statuts que V11, pour les 446 fiches ;
- 446 collectées / 427 retenues retrouvées, fiche par fiche, comme le filtre FILTRE_QUAL du .sav livré.

Les données réelles contiennent des données personnelles : elles restent dans reference/prive/ (hors Git).
Sur un poste qui ne les a pas, ces tests sont sautés avec un message explicite ; aucune valeur n'est
affichée en cas d'échec (seulement des N° de fiche et des noms de variables).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from donnees_reelles import MASQUE_0910, SAV_446, DonneesReelles, charger

from plateforme.codification.qualite import COPIE, SUSPECT, ParametresBase, base_analyse, quasi_doublons, taux_remplissage
from plateforme.codification.rejeu import HORS_COMPTAGE
from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.generateurs.masque import colonnes_masque
from plateforme.terrain.plan import PlanTerrain

pytestmark = pytest.mark.skipif(not MASQUE_0910.exists(), reason="données réelles absentes (reference/prive/, hors Git)")


@pytest.fixture(scope="module")
def reel(dico_edm: Dictionnaire, terrain_edm: PlanTerrain) -> DonneesReelles:
    return charger(dico_edm, terrain_edm)


def test_historique_reconstitue(reel: DonneesReelles) -> None:
    assert reel.nb_reponses == 451
    assert len(reel.masque) == 446
    assert len(reel.rejetees) == 4                       # fiches vides du 06/10, numérotées puis retirées
    assert reel.ecartees == [("ENQ6", "2026-10-03", "18:56")]   # fiche de test supprimée avant le terrain
    assert sum(f.statut == "ANNULEE" for f in reel.fiches) == 4


def test_fiches_identiques_a_v11(reel: DonneesReelles, dico_edm: Dictionnaire, terrain_edm: PlanTerrain) -> None:
    cols = colonnes_masque(dico_edm, terrain_edm.villes_masque())
    textes = {c.code for c in cols if not c.est_numerique()}
    par_num = {r["NUMFICHE"]: r for r in reel.masque}
    valides = [f for f in reel.fiches if f.statut not in HORS_COMPTAGE]
    assert sorted(f.numfiche for f in valides) == sorted(par_num), "N° de fiche"
    ecarts: dict[str, list[str]] = {}
    for f in valides:
        attendu = par_num[f.numfiche]
        for c in cols:
            a, o = attendu.get(c.code, ""), f.valeurs.get(c.code, "")
            if c.code in textes:
                a, o = str(a), str(o)                    # le tableur a converti en nombre un quartier saisi en chiffres
            if a != o:
                ecarts.setdefault(c.code, []).append(f.numfiche)
    assert ecarts == {}, {k: v[:5] for k, v in ecarts.items()}


def test_remplissage_et_quasi_doublons_comme_le_sav(reel: DonneesReelles, dico_edm: Dictionnaire) -> None:
    pyreadstat = pytest.importorskip("pyreadstat")
    df, _ = pyreadstat.read_sav(str(SAV_446), usecols=["NUMFICHE", "TX_REMPL", "DOUBLON_SUSP", "FILTRE_QUAL"])
    sav = {r.NUMFICHE.strip(): r for r in df.itertuples()}
    valides = [f for f in reel.fiches if f.statut not in HORS_COMPTAGE]
    assert [f.numfiche for f in valides if abs(taux_remplissage(dico_edm, f) - sav[f.numfiche].TX_REMPL) > 0.05] == []
    qd = quasi_doublons(dico_edm, valides, ParametresBase())
    copies = {f.numfiche for f in valides if qd[f.id_reponse] == COPIE}
    assert copies == {n for n, r in sav.items() if r.DOUBLON_SUSP == 2}, "copies probables (écartées)"
    # « Ressemblance forte » : la liste SPSS, faite à la main, a omis deux fiches situées comme 032 et 037
    # (qu'elle signale) à 6 écarts d'un membre du groupe A. Ce statut n'écarte aucune fiche.
    suspects = {f.numfiche for f in valides if qd[f.id_reponse] == SUSPECT}
    assert suspects - {n for n, r in sav.items() if r.DOUBLON_SUSP == 1} == {"YDE-ENQ5-041", "YDE-ENQ5-042"}
    assert {n for n, r in sav.items() if r.DOUBLON_SUSP == 1} <= suspects


DECISIONS = MASQUE_0910.parents[1] / "decisions_qc_edm.yaml"


@pytest.mark.skipif(not DECISIONS.exists(), reason="décisions de l'étude absentes (reference/prive/, hors Git)")
def test_ligne_de_commande_base_coder(tmp_path: Path) -> None:
    """Du tableur brut aux livrables, avec les décisions réelles (1 réponse de test EXCLUE, 4 fiches vides ANNULEE)."""
    from conftest import PROJET_EDM

    from plateforme.cli import main
    code = main(["base", "coder", str(MASQUE_0910), "--dictionnaire", str(PROJET_EDM / "dictionnaire.yaml"),
                 "--plan-terrain", str(PROJET_EDM / "plan_terrain.yaml"), "--decisions", str(DECISIONS), "--sortie", str(tmp_path)])
    assert code == 0
    assert "446 fiches valides collectées" in (tmp_path / "NOTE_DE_BASE.txt").read_text(encoding="utf-8")
    assert "427 fiches retenues" in (tmp_path / "NOTE_DE_BASE.txt").read_text(encoding="utf-8")
    for f in ("MASQUE_DEPOUILLEMENT.xlsx", "BASE_ANALYSE.sav", "BASE_ANALYSE_ENSEMBLES_MULTIPLES.sps", "BASE_ANALYSE.csv", "CONTACTS.xlsx"):
        assert (tmp_path / f).stat().st_size > 0, f


def test_base_446_collectees_427_retenues(reel: DonneesReelles, dico_edm: Dictionnaire) -> None:
    pyreadstat = pytest.importorskip("pyreadstat")
    df, _ = pyreadstat.read_sav(str(SAV_446), usecols=["NUMFICHE", "FILTRE_QUAL"])
    retenues_sav = {r.NUMFICHE.strip() for r in df.itertuples() if r.FILTRE_QUAL == 1}
    retenues, note = base_analyse(dico_edm, reel.fiches, ParametresBase(), reel.decisions)
    assert {f.numfiche for f in retenues} == retenues_sav
    assert note.retenues == 427
    assert note.valides == 446
    assert note.texte() == ("Base : 450 fiches reçues, dont 4 annulées (incident technique, non imputées), soit 446 fiches valides collectées ; "
                            "19 écartées de l'analyse (remplissage < 70 % : 11 ; copies probables : 8) ; "
                            "427 fiches retenues pour l'analyse.")


@pytest.mark.skipif(not (DECISIONS.exists() and SAV_446.exists()), reason="données de l'étude absentes (reference/prive/, hors Git)")
def test_ligne_de_commande_analyse_base(tmp_path: Path) -> None:
    """Du tableur brut aux deux bases SPSS livrées (collecte 446, analyse 427), dérivées comprises."""
    from conftest import PROJET_EDM

    from plateforme.analyse.donnees import depuis_sav
    from plateforme.cli import main
    code = main(["analyse", "base", str(MASQUE_0910), "--dictionnaire", str(PROJET_EDM / "dictionnaire.yaml"),
                 "--plan-terrain", str(PROJET_EDM / "plan_terrain.yaml"), "--plan-analyse", str(PROJET_EDM / "plan_analyse.yaml"),
                 "--decisions", str(DECISIONS), "--sortie", str(tmp_path)])
    assert code == 0
    collecte, analyse = depuis_sav(tmp_path / "BASE_COLLECTE.sav"), depuis_sav(tmp_path / "BASE_ANALYSE_DERIVEES.sav")
    assert (len(collecte.table), len(analyse.table)) == (446, 427)
    assert int((collecte.table["FILTRE_QUAL"] == 1).sum()) == 427
    for v in ("MARQUE_G", "IDX_SAV", "DEPENSE_AN", "TX_REMPL", "JOURTERR"):
        assert v in analyse.table.columns


@pytest.mark.skipif(not (DECISIONS.exists() and SAV_446.exists()), reason="données de l'étude absentes (reference/prive/, hors Git)")
def test_ligne_de_commande_analyse_resultats(tmp_path: Path) -> None:
    from conftest import PROJET_EDM
    from openpyxl import load_workbook

    from plateforme.cli import main
    code = main(["analyse", "resultats", str(MASQUE_0910), "--dictionnaire", str(PROJET_EDM / "dictionnaire.yaml"),
                 "--plan-terrain", str(PROJET_EDM / "plan_terrain.yaml"), "--plan-analyse", str(PROJET_EDM / "plan_analyse.yaml"),
                 "--decisions", str(DECISIONS), "--sortie", str(tmp_path)])
    assert code == 0
    sps = (tmp_path / "RESULTATS.sps").read_text(encoding="utf-8")
    assert "COMPUTE CTRL_BASE = (CTRL_N = 427)." in sps and "GET FILE" not in sps
    assert len(list(load_workbook(tmp_path / "RESULTATS.xlsx")["Syntaxe SPSS"].iter_rows())) == 21
