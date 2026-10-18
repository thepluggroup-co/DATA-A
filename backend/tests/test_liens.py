"""Liens personnels pré-remplis : régénérés depuis UN modèle, identiques à ceux de V11 (onglet ENQUETEURS)."""

from __future__ import annotations

from pathlib import Path

import pytest

from plateforme.deploiement.liens import ModeleInvalide, analyser_modele, lien, qr_code
from plateforme.terrain.plan import PlanTerrain

MODELE = ("https://docs.google.com/forms/d/e/ABC/viewform?usp=pp_url&entry.111=ENQ1&entry.222=Douala&entry.333=Douala"
          "&entry.444=Texte%20fixe")


def test_substitution_et_encodage() -> None:
    m = analyser_modele(MODELE, "ENQ1", ["Douala"])
    assert (m.champs_enqueteur, m.champs_ville) == (["entry.111"], ["entry.222", "entry.333"])
    assert lien(m, "ENQ4", "Yaoundé") == ("https://docs.google.com/forms/d/e/ABC/viewform?usp=pp_url&entry.111=ENQ4&entry.222=Yaound%C3%A9"
                                         "&entry.333=Yaound%C3%A9&entry.444=Texte%20fixe")
    assert "entry.222" not in lien(m, "ENQ9", None)


def test_modele_invalide() -> None:
    with pytest.raises(ModeleInvalide):
        analyser_modele(MODELE, "ENQ7", ["Douala"])
    with pytest.raises(ModeleInvalide):
        analyser_modele(MODELE, "ENQ1", ["Bafoussam"])


def test_qr_code(tmp_path: Path) -> None:
    assert qr_code(MODELE, tmp_path / "ENQ1.png").read_bytes()[:4] == b"\x89PNG"


MASQUE = Path(__file__).resolve().parents[2] / "reference" / "prive" / "donnees" / "EDM_ELECTRIQUE_CAMEROUN_MASQUE_TEMPS_REEL__7_.xlsx"


@pytest.mark.skipif(not MASQUE.exists(), reason="données réelles absentes (reference/prive/, hors Git)")
def test_liens_identiques_a_v11(terrain_edm: PlanTerrain) -> None:
    from openpyxl import load_workbook
    wb = load_workbook(MASQUE, read_only=True, data_only=True)
    v11 = {str(li[0]): str(li[4]) for li in list(wb["ENQUETEURS"].iter_rows(values_only=True))[1:] if li and li[0] and len(li) > 4 and li[4]}
    wb.close()
    m = analyser_modele(v11["ENQ1"], "ENQ1", ["Douala"])
    villes = {v.id: v.libelle for v in terrain_edm.villes}
    produits = {e.id: lien(m, e.id, villes.get(e.ville or "")) for e in terrain_edm.enqueteurs}
    assert produits == v11
