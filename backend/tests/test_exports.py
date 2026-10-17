"""Export .sav vérifié, CSV, CONTACTS — sur données synthétiques (tout poste) et réelles (poste de l'étude)."""

from __future__ import annotations

import math
import random
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from conftest import REFERENCE
from test_codification_oracle import _items_aleatoires, _params

from plateforme.codification.fiche import EtatCodification, FicheCodee, construire_fiche
from plateforme.codification.reponses import ItemReponse, ReponseBrute
from plateforme.dictionnaire.modele import Dictionnaire
from plateforme.exports.sav import ExportInvalide, exporter_sav, syntaxe_mrsets, verifier_sav
from plateforme.exports.tableaux import NomDeFichierInvalide, exporter_contacts, exporter_csv
from plateforme.generateurs.masque import colonnes_masque
from plateforme.terrain.plan import PlanTerrain

pyreadstat = pytest.importorskip("pyreadstat")


@pytest.fixture(scope="module")
def fiches(dico_edm: Dictionnaire, terrain_edm: PlanTerrain) -> list[FicheCodee]:
    rnd, etat, params = random.Random(7), EtatCodification(), _params(terrain_edm)
    t = datetime(2026, 10, 5, 8, tzinfo=UTC)
    out = []
    for i in range(80):
        t += timedelta(minutes=rnd.choice([3, 8, 20]))
        items = tuple(ItemReponse(x["titre"], x["rep"], tuple(x["lignes"]) if x["lignes"] else None)
                      for x in _items_aleatoires(dico_edm, rnd, ["", "694388296"]))
        if i == 0:      # marque « Autre » de 120 caractères, placée en tête (la 1re réponse non vide l'emporte)
            items = (ItemReponse("Q12. Marque achetée le plus souvent", "Marque locale " + "x" * 106), *items)
        out.append(construire_fiche(dico_edm, ReponseBrute(f"S{i}", t, items), params, etat))
    return out


def test_sav_ecrit_relu_et_verifie(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, fiches: list[FicheCodee],
                                   tmp_path: Path) -> None:
    cols = colonnes_masque(dico_edm, terrain_edm.villes_masque())
    res = exporter_sav(dico_edm, cols, fiches, tmp_path / "EDM.sav", libelle_fichier="Test", note="Base : " + "x" * 300)
    df, meta = pyreadstat.read_sav(str(res.chemin), user_missing=True)
    assert (len(df), len(df.columns)) == (80, 67)
    assert meta.column_names_to_labels["VOUSETES"] == "Q1. Vous êtes"
    assert meta.variable_value_labels["MAACPLSO"][3.0] == "Vectron"
    assert sorted(r["lo"] for r in meta.missing_ranges["VOUSETES"]) == [97.0, 98.0, 99.0]
    assert meta.variable_measure["AGEREPON"] == "ordinal" and meta.variable_measure["DELAIPRE"] == "scale"
    assert meta.original_variable_types["DATEENQU"] == "SDATE10"
    assert df.MAACAUTR[0].startswith("Marque locale") and len(df.MAACAUTR[0]) == 120, "texte « Autre » tronqué"
    assert int(meta.original_variable_types["MAACAUTR"][1:]) >= 120
    assert res.syntaxe_mrsets is not None and "$PRODACHE" in res.syntaxe_mrsets.read_text(encoding="utf-8")


def test_syntaxe_mrsets_sans_chemin(dico_edm: Dictionnaire) -> None:
    sps = syntaxe_mrsets(dico_edm)
    assert "GET" not in sps and ":\\" not in sps and "CD " not in sps
    assert "/MDGROUP NAME=$RENFFIDE" in sps


def test_verification_detecte_une_alteration(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, fiches: list[FicheCodee],
                                             tmp_path: Path) -> None:
    cols = colonnes_masque(dico_edm, terrain_edm.villes_masque())
    res = exporter_sav(dico_edm, cols, fiches, tmp_path / "EDM.sav")
    donnees = {c.code: [f.valeurs.get(c.code, "") for f in fiches] for c in cols}
    donnees["VOUSETES"] = [float(v) for v in donnees["VOUSETES"]]
    donnees["VOUSETES"][5] = 4.5                          # valeur attendue qui n'a pas été écrite
    with pytest.raises(ExportInvalide, match="VOUSETES"):
        verifier_sav(res.chemin, cols, donnees, {c.code: c.spss if not c.spss.startswith("A") else "A1" for c in cols}, {}, [97, 98, 99])


@pytest.mark.parametrize("nom", ["EXPORT_SPSS.csv.sav", "EDM.sav.sav", "EDM.csv", "EDM"])
def test_noms_de_fichier_refuses(nom: str, dico_edm: Dictionnaire, terrain_edm: PlanTerrain, fiches: list[FicheCodee],
                                 tmp_path: Path) -> None:
    with pytest.raises(ExportInvalide):
        exporter_sav(dico_edm, colonnes_masque(dico_edm, terrain_edm.villes_masque()), fiches, tmp_path / nom)


def test_csv_et_contacts(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, fiches: list[FicheCodee], tmp_path: Path) -> None:
    from openpyxl import load_workbook
    cols = colonnes_masque(dico_edm, terrain_edm.villes_masque())
    csv = exporter_csv(cols, fiches, tmp_path / "EXPORT_SPSS.csv").read_text(encoding="utf-8").splitlines()
    assert csv[0].split(",")[:3] == ["NUMFICHE", "CODEENQU", "CODEZONE"] and len(csv) == 81
    with pytest.raises(NomDeFichierInvalide):
        exporter_csv(cols, fiches, tmp_path / "EXPORT_SPSS.csv.csv")
    wb = load_workbook(exporter_contacts(fiches, tmp_path / "CONTACTS.xlsx"), read_only=True)
    lignes = list(wb["CONTACTS"].iter_rows(values_only=True))
    assert lignes[0][9] == "TÉLÉPHONE" and len(lignes) == 81
    assert all(isinstance(li[9], str) or li[9] is None for li in lignes[1:]), "téléphone converti en nombre"
    wb.close()


LIVRE_427 = REFERENCE / "prive" / "sav" / "SPSS_EDM_EE_V6_ANALYSE_427.sav"


@pytest.mark.skipif(not LIVRE_427.exists(), reason="données réelles absentes (reference/prive/, hors Git)")
def test_export_427_comme_le_livrable_sans_sa_troncature(dico_edm: Dictionnaire, terrain_edm: PlanTerrain, tmp_path: Path) -> None:
    """Base de 427 exportée = .sav livré au client sur les 60 variables communes, sauf MAACAUTR que le livrable
    avait TRONQUÉ à 8 caractères (défaut documenté dans EDM_EE_RECONFIGURATION_V4_1.sps, ligne 35)."""
    from donnees_reelles import charger

    from plateforme.codification.qualite import ParametresBase, base_analyse
    reel = charger(dico_edm, terrain_edm)
    retenues, note = base_analyse(dico_edm, reel.fiches, ParametresBase(), reel.decisions)
    cols = colonnes_masque(dico_edm, terrain_edm.villes_masque())
    res = exporter_sav(dico_edm, cols, retenues, tmp_path / "EDM_427.sav", note=note.texte())
    nous, _ = pyreadstat.read_sav(str(res.chemin), user_missing=True)       # 97/98/99 lus comme valeurs, pas NaN
    livre, _ = pyreadstat.read_sav(str(LIVRE_427), user_missing=True)
    n_idx, l_idx = nous.set_index(nous.NUMFICHE.str.strip()), livre.set_index(livre.NUMFICHE.str.strip())
    assert set(n_idx.index) == set(l_idx.index)
    communes = [c.code for c in cols if c.code in livre.columns]
    assert len(communes) == 60

    def egal(a: object, b: object) -> bool:
        vide = [x in ("", None) or (isinstance(x, float) and math.isnan(x)) for x in (a, b)]
        return all(vide) or a == b or str(a).strip() == str(b).strip()

    ecarts = {c: [n for n in n_idx.index if not egal(n_idx.at[n, c], l_idx.at[n, c])] for c in communes}
    assert {c for c, v in ecarts.items() if v} == {"MAACAUTR"}
    for n in ecarts["MAACAUTR"]:
        assert str(n_idx.at[n, "MAACAUTR"]).encode("utf-8")[:8].decode("utf-8", "ignore").strip() == str(l_idx.at[n, "MAACAUTR"]).strip()
