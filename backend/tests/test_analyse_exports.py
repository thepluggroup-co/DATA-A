"""Exports des résultats d'analyse : classeur Excel (bases, pondération, avertissements, TOTAL contrôlé) et syntaxe SPSS."""

from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from plateforme.analyse.donnees import JeuDonnees, Variable
from plateforme.analyse.exports import ResultatTitre, TotalIncoherent, exporter_excel, exporter_syntaxe, syntaxe
from plateforme.analyse.procedures import correspondances, frequences, kruskal_wallis, moyennes, reponses_multiples, tri_croise
from plateforme.exports.tableaux import NomDeFichierInvalide


@pytest.fixture
def jd() -> JeuDonnees:
    villes = [1] * 6 + [2] * 6 + [3] * 6
    profil = [1, 1, 2, 2, 3, 99, 1, 2, 2, 3, 3, 3, 1, 1, 1, 2, 3, 3]
    table = pd.DataFrame({"VILLE": villes, "PROFIL": profil, "MONTANT": [float(i * 10) for i in range(18)],
                          "SOL": [1, 0] * 9, "LUM": [0, 1, 1] * 6, "W": [1.0, 2.0] * 9})
    variables = {"VILLE": Variable("VILLE", "Ville", {1.0: "Douala", 2.0: "Yaoundé", 3.0: "Bafoussam"}),
                 "PROFIL": Variable("PROFIL", "Profil", {1.0: "Particulier", 2.0: "Entreprise", 3.0: "Revendeur", 99.0: "NR"}, {99.0}),
                 "MONTANT": Variable("MONTANT", "Montant", mesure="scale"), "SOL": Variable("SOL", "Solaire"), "LUM": Variable("LUM", "Luminaires"),
                 "W": Variable("W", "Poids", mesure="scale")}
    return JeuDonnees(table, variables)


def _resultats(jd: JeuDonnees) -> list[ResultatTitre]:
    return [ResultatTitre("Profil des répondants", frequences(jd, "PROFIL")),
            ResultatTitre("Profil selon la ville", tri_croise(jd, "PROFIL", "VILLE")),
            ResultatTitre("Montant moyen par ville", moyennes(jd, "MONTANT", par="VILLE")),
            ResultatTitre("Produits achetés", reponses_multiples(jd, ["SOL", "LUM"])),
            ResultatTitre("Montant selon la ville (Kruskal-Wallis)", kruskal_wallis(jd, "MONTANT", "VILLE")),
            ResultatTitre("Carte profils × villes", correspondances(jd, "PROFIL", "VILLE"))]


def test_classeur(jd: JeuDonnees, tmp_path: Path) -> None:
    chemin = exporter_excel(_resultats(jd), jd, tmp_path / "RESULTATS.xlsx")
    wb = load_workbook(chemin)
    assert wb.sheetnames == ["Résultats", "Syntaxe SPSS"]
    lignes = list(wb["Résultats"].iter_rows(values_only=True))
    assert lignes[0][0] == "Tableau 1. Profil des répondants"
    assert lignes[1][:4] == ("Procédure", "FREQUENCIES", "Base (fiches)", 17)            # le NR (99) est hors base
    total = next(li for li in lignes if li[0] == "TOTAL")
    assert total[2] == 17 and total[3] == 100
    assert tuple(next(li for li in lignes if li[0] == "Manquants")[:2]) == ("Manquants", 1)
    titres = [li[0] for li in lignes if isinstance(li[0], str) and li[0].startswith("Tableau ")]
    assert len(titres) == 6
    assert any(isinstance(li[0], str) and li[0].startswith("Attention : ") for li in lignes), "effectifs théoriques < 5 signalés"
    syn = list(wb["Syntaxe SPSS"].iter_rows(values_only=True))
    assert syn[1][2] == "FREQUENCIES VARIABLES=PROFIL." and len(syn) == 7


def test_total_controle(jd: JeuDonnees, tmp_path: Path) -> None:
    r = frequences(jd, "PROFIL")
    r.total_pondere += 1
    with pytest.raises(TotalIncoherent):
        exporter_excel([ResultatTitre("faux", r)], jd, tmp_path / "FAUX.xlsx")
    assert not (tmp_path / "FAUX.xlsx").exists(), "rien n'est écrit quand un contrôle échoue"


def test_syntaxe(jd: JeuDonnees, tmp_path: Path) -> None:
    res = [*_resultats(jd), ResultatTitre("Un titre très long " * 6, frequences(jd, "VILLE", poids="W"))]
    texte = syntaxe(res, base_attendue=18, nom_base="BASE.sav")
    assert texte.startswith("* Encoding: UTF-8.")
    assert "GET FILE" not in texte and "CD " not in texte, "aucun chemin de fichier dans la syntaxe (CLAUDE.md §5)"
    assert "COMPUTE CTRL_BASE = (CTRL_N = 18)." in texte, "base active contrôlée en tête"
    assert "NPAR TESTS /K-W=MONTANT BY VILLE(1 3)" in texte
    assert "CORRESPONDENCE TABLE=PROFIL(1 3) BY VILLE(1 3)" in texte
    assert "WEIGHT BY W.\nFREQUENCIES VARIABLES=VILLE." in texte
    assert texte.count("WEIGHT OFF.") == len(res) + 1, "en tête, puis chaque bloc repart sans pondération"
    assert all(len(li) <= len('TITLE "".') + 60 for li in texte.splitlines() if li.startswith("TITLE"))
    chemin = exporter_syntaxe(res, tmp_path / "RESULTATS.sps")
    assert chemin.read_text(encoding="utf-8") == syntaxe(res)
    assert "CTRL_BASE" not in syntaxe(res), "sans effectif attendu, pas de contrôle"
    with pytest.raises(NomDeFichierInvalide):
        exporter_syntaxe(res, tmp_path / "RESULTATS.sps.sps")


def test_mediane_ponderee_non_donnee(jd: JeuDonnees) -> None:
    m = moyennes(jd, "MONTANT", poids="W")
    assert math.isnan(m.mediane) and any("médiane" in a for a in m.avertissements)
    assert m.moyenne == pytest.approx(sum(i * 10 * (1 if i % 2 == 0 else 2) for i in range(18)) / 27)
    assert not math.isnan(moyennes(jd, "MONTANT").mediane)


def test_sav_du_jeu(jd: JeuDonnees, tmp_path: Path) -> None:
    from plateforme.analyse.donnees import depuis_sav
    from plateforme.analyse.exports import exporter_sav_jeu

    jd = jd.avec_variable("DATEENQU", ["2026-10-05"] * 18).avec_variable("TXT", ["Schneider, ingelec"] + [""] * 17)
    chemin = exporter_sav_jeu(jd, tmp_path / "ANALYSE.sav", formats={"DATEENQU": "SDATE10"}, libelle_fichier="Essai", note="Base d'essai.")
    relu = depuis_sav(chemin)
    assert list(relu.table.columns) == list(jd.table.columns)
    assert relu.variables["PROFIL"].manquants == {99.0} and relu.variables["PROFIL"].valeurs[2.0] == "Entreprise"
    assert relu.valides("PROFIL").isna().sum() == 1
    assert relu.table["TXT"].iloc[0] == "Schneider, ingelec", "texte complet, jamais tronqué"
    assert str(relu.table["DATEENQU"].iloc[0]) == "2026-10-05"


def test_sav_plus_de_trois_manquants(jd: JeuDonnees, tmp_path: Path) -> None:
    from plateforme.analyse.exports import SavInvalide, exporter_sav_jeu

    jd.variables["PROFIL"] = Variable("PROFIL", "Profil", {}, {96.0, 97.0, 98.0, 99.0})
    with pytest.raises(SavInvalide, match="PROFIL"):
        exporter_sav_jeu(jd, tmp_path / "ANALYSE.sav")


def test_regression_et_anova_exportees(jd: JeuDonnees, tmp_path: Path) -> None:
    from plateforme.analyse.procedures import anova_un_facteur, regression_lineaire

    res = [ResultatTitre("Montant selon solaire et luminaires", regression_lineaire(jd, "MONTANT", ["SOL", "LUM"])),
           ResultatTitre("Montant selon la ville (ANOVA)", anova_un_facteur(jd, "MONTANT", "VILLE"))]
    lignes = list(load_workbook(exporter_excel(res, jd, tmp_path / "R.xlsx"))["Résultats"].iter_rows(values_only=True))
    assert any(li[0] == "const" for li in lignes) and any(li[0] == "R²" for li in lignes)
    assert any(li[1] == "Yaoundé" for li in lignes), "groupes d'ANOVA libellés"
    assert "ONEWAY MONTANT BY VILLE" in syntaxe(res) and "REGRESSION /MISSING LISTWISE" in syntaxe(res)
